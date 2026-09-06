"""Bounded physical document reads from the monitoring candidate copy store."""
from __future__ import annotations

import hashlib
import io
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from packages.medical_monitoring.admission.source_tools import SourceToolError
from packages.medical_monitoring.intelligence.primitives import canonical_json, content_hash
from .protocol_text_extractor import parse_protocol_docx

_SUFFIX = {
    "application/pdf": ".pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
}


def _range(value, minimum, maximum):
    if type(value) is not int or not minimum <= value <= maximum:
        raise SourceToolError("invalid_document_read_range")
    return value


class FrozenDocumentEvidenceTools:
    def __init__(self, *, candidate_root: Path, project_id: str, input_revision: str, resolver):
        self.candidate_root = Path(candidate_root)
        self.project_id = project_id
        self.input_revision = input_revision
        self.resolver = resolver

    def read_document_units(self, *, binding, offset: int = 0, limit: int = 4,
                            text_offset: int = 0, text_limit: int = 16000,
                            column_start: int = 0, column_count: int = 12) -> dict[str, Any]:
        _range(offset, 0, 10_000_000)
        _range(limit, 1, 8)
        _range(text_offset, 0, 10_000_000)
        _range(text_limit, 1, 24000)
        _range(column_start, 0, 16383)
        _range(column_count, 1, 24)
        self.resolver.assert_current_binding(project_id=self.project_id, binding=binding)
        suffix = _SUFFIX.get(binding.media_type)
        if suffix is None:
            raise SourceToolError("document_format_not_supported")
        path = self.candidate_root / "files" / (binding.content_sha256 + suffix)
        try:
            content = path.read_bytes()
        except OSError as exc:
            raise SourceToolError("frozen_document_unavailable") from exc
        if hashlib.sha256(content).hexdigest() != binding.content_sha256:
            raise SourceToolError("frozen_document_content_changed")
        limitations = []
        inventory = {}
        units = []
        if suffix == ".pdf":
            import pymupdf
            with pymupdf.open(stream=content, filetype="pdf") as document:
                total = len(document)
                for number in range(offset, min(total, offset + limit)):
                    page = document[number]
                    images = len(page.get_image_info())
                    drawings = len(page.get_drawings())
                    units.append({"locator": f"pdf:page:{number + 1}", "text": page.get_text("text"),
                                  "image_regions": images, "vector_objects": drawings,
                                  "visual_content_assessed": False})
            limitations.append("native_text_only_visual_content_not_assessed")
        elif suffix == ".docx":
            document = parse_protocol_docx("frozen.docx", content)
            spans = document.spans
            all_units = [{"locator": span.source_locator, "text": span.text, "kind": span.kind,
                          "body_order": span.body_order} for span in spans]
            all_units.extend(self._word_notes_and_revisions(content))
            total = len(all_units)
            units = all_units[offset:offset + limit]
            inventory = {"embedded_image_count": len(document.embedded_images),
                         "embedded_images": [{"locator": image.source_locator,
                                              "content_sha256": image.image_sha256,
                                              "media_type": image.media_type}
                                             for image in document.embedded_images],
                         "revision_fragments": sum(u.get("kind") in {"inserted_revision", "deleted_revision"}
                                                   for u in all_units)}
            # Preserve the shared extractor; the monitoring tool makes no
            # assertion that drawing semantics or revision intent were read.
            limitations.append("embedded_images_and_revision_intent_not_assessed")
        else:
            import openpyxl
            workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            cached = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                total = sum(sheet.max_row or 0 for sheet in workbook.worksheets)
                base = 0
                for sheet in workbook.worksheets:
                    row_total = sheet.max_row or 0
                    start = max(0, offset - base)
                    end = min(row_total, offset + limit - base)
                    if start < end:
                        maximum = min(sheet.max_column or 0, column_start + column_count)
                        if maximum <= column_start:
                            raise SourceToolError("document_column_out_of_range")
                        rows = sheet.iter_rows(min_row=start + 1, max_row=end,
                                               min_col=column_start + 1, max_col=maximum)
                        cached_rows = cached[sheet.title].iter_rows(min_row=start + 1, max_row=end,
                                               min_col=column_start + 1, max_col=maximum)
                        for row_number, (row, cached_row) in enumerate(zip(rows, cached_rows), start=start + 1):
                            cells = [{"coordinate": f"{openpyxl.utils.get_column_letter(column_start + cell_index + 1)}{row_number}", "value": self._value(cell.value),
                                      "data_type": cell.data_type,
                                      "cached_value": self._value(cache.value) if cell.data_type == "f" else None}
                                     for cell_index, (cell, cache) in enumerate(zip(row, cached_row))]
                            units.append({"locator": f"xlsx:sheet:{sheet.title}:row:{row_number}",
                                          "sheet": sheet.title, "row_number": row_number, "cells": cells,
                                          "total_columns": sheet.max_column,
                                          "next_column_start": maximum if maximum < sheet.max_column else None})
                    base += row_total
            finally:
                workbook.close()
                cached.close()
            limitations.append("spreadsheet_drawings_comments_and_formula_evaluation_not_assessed")
        if offset > total:
            raise SourceToolError("document_unit_out_of_range")
        for index, unit in enumerate(units, start=offset):
            unit["unit_index"] = index
            if "text" in unit:
                original = unit["text"]
                unit.update(text=original[text_offset:text_offset + text_limit], text_offset=text_offset,
                            total_text_characters=len(original), text_sha256=content_hash(original),
                            next_text_offset=(text_offset + text_limit if text_offset + text_limit < len(original) else None))
        self.resolver.assert_current_binding(project_id=self.project_id, binding=binding)
        result = {"schema_version": "mm-frozen-document-units-v1", "project_id": self.project_id,
                  "input_revision_sha256": self.input_revision, "source_entry_id": binding.source_entry_id,
                  "source_content_sha256": binding.content_sha256, "locator_index_sha256": binding.locator_index_sha256,
                  "coverage": "partial", "coverage_scope": "physical_document_units",
                  "absence_claim_supported": False, "units": units, "total_units": total,
                  "physical_inventory": inventory,
                  "next_offset": offset + len(units) if offset + len(units) < total else None,
                  "limitation_codes": limitations}
        if len(canonical_json(result).encode("utf-8")) > 256000:
            raise SourceToolError("document_result_too_large_request_smaller_region")
        return {**result, "result_sha256": content_hash(result)}

    @staticmethod
    def _value(value):
        return value.isoformat() if hasattr(value, "isoformat") else value

    @staticmethod
    def _word_notes_and_revisions(content):
        """Return OOXML text with explicit provenance, never accept/reject edits."""
        ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        def text_of(element):
            if element.tag in {ns + "t", ns + "delText"}:
                return element.text or ""
            if element.tag == ns + "tab":
                return "\t"
            if element.tag == ns + "br":
                return "\n"
            return "".join(text_of(child) for child in element) + ("\n" if element.tag == ns + "p" else "")
        units = []
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            for part, tag in (("word/footnotes.xml", "footnote"), ("word/endnotes.xml", "endnote")):
                if part not in archive.namelist():
                    continue
                for note in ElementTree.fromstring(archive.read(part)).findall(ns + tag):
                    if note.get(ns + "type") in {"separator", "continuationSeparator", "continuationNotice"}:
                        continue
                    units.append({"locator": f"docx:{tag}:{note.get(ns + 'id')}", "kind": tag,
                                  "text": text_of(note)})
            document = ElementTree.fromstring(archive.read("word/document.xml"))
            for index, element in enumerate(document.iter()):
                if element.tag not in {ns + "ins", ns + "del"}:
                    continue
                units.append({"locator": f"docx:revision-element:{index}",
                              "kind": "inserted_revision" if element.tag == ns + "ins" else "deleted_revision",
                              "text": text_of(element),
                              "accepted_as_current_text": False})
        return units
