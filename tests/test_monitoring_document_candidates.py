from __future__ import annotations

import io
import json
from pathlib import Path

import docx
import openpyxl
import pymupdf
import pytest

import services.api.app.monitoring_document_candidates as candidates_module
from services.api.app.monitoring_document_candidates import (
    CandidateOcrUnavailableError,
    MonitoringDocumentCandidateDecomposer,
)


def _docx_bytes(*paragraphs: str) -> bytes:
    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _xlsx_bytes() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "不良事件"
    sheet.append(["受试者编号", "事件名称", "严重程度"])
    sheet.append(["SUBJECT-SECRET", "HEADACHE", "MILD"])
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def _pdf_bytes(*, text: str = "") -> bytes:
    document = pymupdf.open()
    page = document.new_page()
    if text:
        page.insert_text((72, 72), text)
    payload = document.tobytes()
    document.close()
    return payload


def _image_only_pdf_bytes() -> bytes:
    source = pymupdf.open()
    source_page = source.new_page(width=32, height=32)
    source_page.draw_rect(source_page.rect, fill=(0, 0, 0))
    image = source_page.get_pixmap().tobytes("png")
    source.close()
    document = pymupdf.open()
    page = document.new_page()
    page.insert_image(page.rect, stream=image)
    payload = document.tobytes()
    document.close()
    return payload


def _mixed_pdf_bytes() -> bytes:
    document = pymupdf.open()
    document.new_page().insert_text((72, 72), "Current protocol")
    document.new_page()
    payload = document.tobytes()
    document.close()
    return payload


def _blank_pdf_bytes(page_count: int) -> bytes:
    document = pymupdf.open()
    for _ in range(page_count):
        document.new_page()
    payload = document.tobytes()
    document.close()
    return payload


def _dense_native_and_blank_pdf_bytes() -> bytes:
    document = pymupdf.open()
    for page_number in range(1, 13):
        document.new_page().insert_text((72, 72), f"Native page {page_number}")
    document.new_page()
    payload = document.tobytes()
    document.close()
    return payload


def test_xlsx_candidate_exposes_structure_without_row_values(tmp_path: Path) -> None:
    decomposer = MonitoringDocumentCandidateDecomposer(tmp_path / "candidates")

    candidate = decomposer.decompose("/private/incoming/eCRF.xlsx", _xlsx_bytes())
    serialized = json.dumps(candidate.to_dict(), ensure_ascii=False)

    assert candidate.filename == "eCRF.xlsx"
    assert candidate.role_hypotheses == ("ecrf",)
    assert candidate.authority_status == "not_promoted"
    assert candidate.sheets[0].sheet_name == "不良事件"
    assert candidate.sheets[0].headers == ("受试者编号", "事件名称", "严重程度")
    assert "SUBJECT-SECRET" not in serialized
    assert "HEADACHE" not in serialized
    assert "/private/incoming" not in serialized


def test_candidate_metadata_redacts_paths_from_excel_structure(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "路径检查"
    sheet.append(["/Users/example/source", r"C:\private\column"])
    sheet.append(["SECRET", "VALUE"])
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()

    candidate = MonitoringDocumentCandidateDecomposer(
        tmp_path / "candidates"
    ).decompose(r"C:\private\ecrf.xlsx", stream.getvalue())
    serialized = json.dumps(candidate.to_dict(), ensure_ascii=False)

    assert candidate.filename == "ecrf.xlsx"
    assert "/Users/example" not in serialized
    assert r"C:\private" not in serialized
    assert "SECRET" not in serialized


def test_headerless_xlsx_does_not_expose_first_row_as_headers(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["SUBJECT-SECRET", "HEADACHE", "MILD"])
    sheet.append(["SUBJECT-SECOND", "NAUSEA", "MODERATE"])
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()

    candidate = MonitoringDocumentCandidateDecomposer(
        tmp_path / "candidates"
    ).decompose("headerless.xlsx", stream.getvalue())
    serialized = json.dumps(candidate.to_dict(), ensure_ascii=False)

    assert candidate.sheets[0].headers == ()
    assert "SUBJECT-SECRET" not in serialized
    assert "HEADACHE" not in serialized


def test_marker_like_headerless_values_are_not_exposed(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["SUBJID", "SITEID", "VISIT"])
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()

    candidate = MonitoringDocumentCandidateDecomposer(
        tmp_path / "candidates"
    ).decompose("headerless-markers.xlsx", stream.getvalue())

    assert candidate.sheets[0].headers == ()


def test_xlsx_content_profile_binds_rows_without_exposing_values(tmp_path: Path) -> None:
    def workbook_bytes(value: str) -> bytes:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["SUBJID", "AETERM"])
        sheet.append(["001", value])
        stream = io.BytesIO()
        workbook.save(stream)
        workbook.close()
        return stream.getvalue()

    left = MonitoringDocumentCandidateDecomposer(tmp_path / "left-xlsx").decompose(
        "listing.xlsx", workbook_bytes("Headache")
    )
    right = MonitoringDocumentCandidateDecomposer(tmp_path / "right-xlsx").decompose(
        "listing.xlsx", workbook_bytes("Nausea")
    )

    assert left.sheets[0].headers == right.sheets[0].headers
    assert left.content_profile is not None
    assert right.content_profile is not None
    assert left.content_profile.normalized_text_sha256 != (
        right.content_profile.normalized_text_sha256
    )
    serialized = json.dumps(left.to_dict(), ensure_ascii=False)
    assert "Headache" not in serialized


def test_docx_candidate_is_deterministic_and_redacts_local_paths(tmp_path: Path) -> None:
    content = _docx_bytes(
        "研究方案摘要",
        "路径 /Users/example/a.docx /private/x/b.pdf /home/u/c.pdf ",
        r"/Volumes/data/d.pdf C:\private\e.docx \\server\share\f.pdf ",
        "file:///tmp/g.pdf 均仅用于定位。",
        r"/secret C:\Users\Alice Smith\h.docx \\server\shared folder\i.pdf",
        "file:///Users/Alice Smith/j.docx",
    )
    decomposer = MonitoringDocumentCandidateDecomposer(tmp_path / "candidates")

    first = decomposer.decompose("protocol.docx", content)
    second = decomposer.decompose("protocol.docx", content)

    assert first == second
    assert first.role_hypotheses == (
        "protocol",
        "investigator_brochure",
        "ecrf",
        "sap",
    )
    assert first.extraction_status == "parsed"
    assert first.locator_count == 6
    assert "[local_path_redacted]" in json.dumps(first.to_dict(), ensure_ascii=False)
    assert "/Users/example" not in json.dumps(first.to_dict(), ensure_ascii=False)
    assert "/private/x" not in json.dumps(first.to_dict(), ensure_ascii=False)
    assert "/home/u" not in json.dumps(first.to_dict(), ensure_ascii=False)
    assert "/Volumes/data" not in json.dumps(first.to_dict(), ensure_ascii=False)
    assert r"C:\private" not in json.dumps(first.to_dict(), ensure_ascii=False)
    assert r"\\server\share" not in json.dumps(first.to_dict(), ensure_ascii=False)
    assert "file:///tmp" not in json.dumps(first.to_dict(), ensure_ascii=False)
    assert "/secret" not in json.dumps(first.to_dict(), ensure_ascii=False)
    assert "Alice Smith" not in json.dumps(first.to_dict(), ensure_ascii=False)
    assert "shared folder" not in json.dumps(first.to_dict(), ensure_ascii=False)
    assert len(list((tmp_path / "candidates" / "files").iterdir())) == 1
    assert len(list((tmp_path / "candidates" / "manifests").iterdir())) == 1


def test_pdf_candidate_distinguishes_native_text_from_ocr_need(tmp_path: Path) -> None:
    decomposer = MonitoringDocumentCandidateDecomposer(tmp_path / "candidates")

    native = decomposer.decompose("ib.pdf", _pdf_bytes(text="Investigator Brochure"))
    scanned = decomposer.decompose("sap.pdf", _pdf_bytes())

    assert native.extraction_status == "parsed"
    assert native.role_hypotheses == (
        "protocol",
        "investigator_brochure",
        "ecrf",
        "sap",
    )
    assert native.locator_count == 1
    assert scanned.extraction_status == "needs_ocr"
    assert scanned.technical_status == "ready"
    assert scanned.zero_text_page_count == 1
    assert scanned.limitation_codes == ("native_text_absent_ocr_required",)

    image_only = decomposer.decompose("scan.pdf", _image_only_pdf_bytes())
    assert image_only.extraction_status == "needs_ocr"
    assert image_only.locator_count == 0


def test_pdf_candidate_recovers_zero_text_pages_with_injected_ocr(
    tmp_path: Path,
) -> None:
    calls: list[tuple[int, int, str, bytes]] = []

    class OcrText(str):
        model = "actual-ocr"
        provider = "local-test"
        fell_back = True

        @property
        def text(self) -> str:
            return str(self)

    def ocr(page: int, dpi: int, model: str, image: bytes) -> OcrText:
        calls.append((page, dpi, model, image))
        return OcrText("签署勘误：第 3 条更正")

    content = _image_only_pdf_bytes()
    without_ocr = MonitoringDocumentCandidateDecomposer(
        tmp_path / "without-ocr"
    ).decompose("erratum.pdf", content)
    recovered = MonitoringDocumentCandidateDecomposer(
        tmp_path / "with-ocr",
        ocr_runner=ocr,
        ocr_model="requested-ocr",
        ocr_dpi=200,
    ).decompose("erratum.pdf", content)

    assert recovered.candidate_id == without_ocr.candidate_id
    assert recovered.extraction_status == "parsed"
    assert recovered.limitation_codes == ()
    assert recovered.locator_count == 1
    assert recovered.excerpts[0].locator.endswith(":p1:ocr")
    assert recovered.ocr_recovery_pages[0].status == "recovered"
    assert recovered.ocr_recovery_pages[0].requested_model == "requested-ocr"
    assert recovered.ocr_recovery_pages[0].actual_model == "actual-ocr"
    assert recovered.ocr_recovery_pages[0].provider == "local-test"
    assert recovered.ocr_recovery_pages[0].fell_back is True
    assert calls[0][:3] == (1, 200, "requested-ocr")
    assert calls[0][3].startswith(b"\x89PNG\r\n\x1a\n")


@pytest.mark.parametrize(
    ("runner", "expected_limitation", "expected_status"),
    (
        (lambda *_args: "", "ocr_recovery_empty", "empty"),
        (
            lambda *_args: (_ for _ in ()).throw(
                CandidateOcrUnavailableError("gateway down")
            ),
            "ocr_recovery_failed",
            "failed",
        ),
    ),
)
def test_pdf_candidate_ocr_failure_remains_non_promotable(
    tmp_path: Path,
    runner: object,
    expected_limitation: str,
    expected_status: str,
) -> None:
    candidate = MonitoringDocumentCandidateDecomposer(
        tmp_path / expected_status,
        ocr_runner=runner,  # type: ignore[arg-type]
    ).decompose("scan.pdf", _image_only_pdf_bytes())

    assert candidate.extraction_status == "needs_ocr"
    assert candidate.locator_count == 0
    assert candidate.limitation_codes == (expected_limitation,)
    assert candidate.ocr_recovery_pages[0].status == expected_status


def test_pdf_candidate_path_only_ocr_is_not_admissible_evidence(
    tmp_path: Path,
) -> None:
    candidate = MonitoringDocumentCandidateDecomposer(
        tmp_path / "path-only",
        ocr_runner=lambda *_args: "/private/tmp/page-result.json",
    ).decompose("scan.pdf", _image_only_pdf_bytes())

    assert candidate.extraction_status == "needs_ocr"
    assert candidate.locator_count == 0
    assert candidate.ocr_recovery_pages[0].status == "empty"


def test_pdf_candidate_preserves_meaning_before_redacted_ocr_path(
    tmp_path: Path,
) -> None:
    candidate = MonitoringDocumentCandidateDecomposer(
        tmp_path / "text-plus-path",
        ocr_runner=lambda *_args: "签署勘误：第 3 条更正 /private/tmp/result.json",
    ).decompose("scan.pdf", _image_only_pdf_bytes())

    assert candidate.extraction_status == "parsed"
    assert candidate.excerpts[0].text == "签署勘误：第 3 条更正 [local_path_redacted]"


def test_docx_content_profile_uses_text_beyond_excerpt_budget(tmp_path: Path) -> None:
    shared = tuple(f"共同正文段落 {index}" for index in range(12))
    left = MonitoringDocumentCandidateDecomposer(tmp_path / "left").decompose(
        "left.docx", _docx_bytes(*shared, "第十三段有效修订 A")
    )
    right = MonitoringDocumentCandidateDecomposer(tmp_path / "right").decompose(
        "right.docx", _docx_bytes(*shared, "第十三段有效修订 B")
    )

    assert [item.text for item in left.excerpts] == [
        item.text for item in right.excerpts
    ]
    assert left.content_profile is not None
    assert right.content_profile is not None
    assert left.content_profile.normalized_text_sha256 != (
        right.content_profile.normalized_text_sha256
    )


def test_pdf_candidate_failed_ocr_records_requested_model_without_fake_actual(
    tmp_path: Path,
) -> None:
    candidate = MonitoringDocumentCandidateDecomposer(
        tmp_path / "failed-model",
        ocr_runner=lambda *_args: (_ for _ in ()).throw(
            CandidateOcrUnavailableError(
                "gateway down",
                requested_model="PaddleOCR-VL-1.6",
                provider="paddle_official",
            )
        ),
    ).decompose("scan.pdf", _image_only_pdf_bytes())

    evidence = candidate.ocr_recovery_pages[0]
    assert evidence.requested_model == "PaddleOCR-VL-1.6"
    assert evidence.actual_model == ""
    assert evidence.provider == "paddle_official"


def test_pdf_candidate_requires_recovery_for_every_native_zero_page(
    tmp_path: Path,
) -> None:
    candidate = MonitoringDocumentCandidateDecomposer(
        tmp_path / "mixed",
        ocr_runner=lambda page, *_args: "" if page == 2 else "unexpected",
    ).decompose("mixed.pdf", _mixed_pdf_bytes())

    assert candidate.extraction_status == "needs_ocr"
    assert candidate.locator_count == 1
    assert candidate.zero_text_page_samples == (2,)
    assert candidate.limitation_codes == ("ocr_recovery_empty",)


def test_native_excerpt_budget_reserves_space_for_zero_text_page_ocr(
    tmp_path: Path,
) -> None:
    candidate = MonitoringDocumentCandidateDecomposer(
        tmp_path / "dense-mixed",
        ocr_runner=lambda page, *_args: f"Recovered page {page}",
    ).decompose("dense-mixed.pdf", _dense_native_and_blank_pdf_bytes())

    assert candidate.extraction_status == "parsed"
    assert candidate.locator_count == 12
    assert len(candidate.ocr_recovery_pages) == 1
    assert candidate.ocr_recovery_pages[0].page_number == 13
    assert candidate.ocr_recovery_pages[0].status == "recovered"
    assert sum(excerpt.locator.endswith(":ocr") for excerpt in candidate.excerpts) == 1


def test_pdf_candidate_does_not_mark_unexposed_ocr_pages_as_parsed(
    tmp_path: Path,
) -> None:
    candidate = MonitoringDocumentCandidateDecomposer(
        tmp_path / "long-scan", ocr_runner=lambda page, *_args: f"page {page}"
    ).decompose("long-scan.pdf", _blank_pdf_bytes(13))

    assert candidate.extraction_status == "needs_ocr"
    assert candidate.locator_count == 12
    assert len(candidate.excerpts) == len(candidate.ocr_recovery_pages) == 12
    assert candidate.limitation_codes == ("ocr_evidence_budget_exhausted",)
    for excerpt, evidence in zip(candidate.excerpts, candidate.ocr_recovery_pages):
        assert evidence.locator == excerpt.locator
        assert evidence.text_sha256 == excerpt.text_sha256
        assert evidence.character_count == len(excerpt.text)


def test_pdf_candidate_ocr_retry_creates_new_immutable_evidence_revision(
    tmp_path: Path,
) -> None:
    root = tmp_path / "retry"
    content = _image_only_pdf_bytes()
    first = MonitoringDocumentCandidateDecomposer(root).decompose("scan.pdf", content)
    second = MonitoringDocumentCandidateDecomposer(
        root, ocr_runner=lambda *_args: "recovered"
    ).decompose("scan.pdf", content)

    assert first.candidate_id == second.candidate_id
    assert first.evidence_revision_sha256 != second.evidence_revision_sha256
    manifests = list((root / "manifests" / first.candidate_id).glob("*.json"))
    assert len(manifests) == 2


def test_unexpected_ocr_runner_defect_is_not_hidden(tmp_path: Path) -> None:
    def broken_runner(*_args: object) -> str:
        raise RuntimeError("programming defect")

    with pytest.raises(RuntimeError, match="programming defect"):
        MonitoringDocumentCandidateDecomposer(
            tmp_path / "broken", ocr_runner=broken_runner
        ).decompose("scan.pdf", _image_only_pdf_bytes())
    assert not (tmp_path / "broken" / "manifests").exists()


def test_pdf_candidate_rejects_unsafe_ocr_configuration(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="OCR configuration"):
        MonitoringDocumentCandidateDecomposer(
            tmp_path / "candidates", ocr_runner=lambda *_args: "text", ocr_dpi=199
        )


def test_malformed_supported_file_is_retained_as_unreadable_candidate(
    tmp_path: Path,
) -> None:
    candidate = MonitoringDocumentCandidateDecomposer(
        tmp_path / "candidates"
    ).decompose("unknown.pdf", b"not-a-pdf")

    assert candidate.extraction_status == "unreadable"
    assert candidate.technical_status == "failed"
    assert candidate.use_status == "candidate_only"
    assert candidate.limitation_codes == ("candidate_parse_failed",)


def test_malformed_xlsx_is_retained_as_unreadable_candidate(tmp_path: Path) -> None:
    candidate = MonitoringDocumentCandidateDecomposer(
        tmp_path / "candidates"
    ).decompose("unknown.xlsx", b"not-an-xlsx")

    assert candidate.extraction_status == "unreadable"
    assert candidate.technical_status == "failed"
    manifest = (
        tmp_path
        / "candidates"
        / "manifests"
        / candidate.candidate_id
        / f"{candidate.evidence_revision_sha256}.json"
    )
    assert manifest.exists()


def test_unexpected_parser_failure_is_not_persisted_as_file_damage(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    decomposer = MonitoringDocumentCandidateDecomposer(tmp_path / "candidates")

    def fail_unexpectedly(*_args: object) -> None:
        raise RuntimeError("programming defect")

    monkeypatch.setattr(decomposer, "_decompose_xlsx", fail_unexpectedly)
    with pytest.raises(RuntimeError, match="programming defect"):
        decomposer.decompose("valid.xlsx", _xlsx_bytes())

    assert not (tmp_path / "candidates" / "manifests").exists()


def test_crc_damaged_xlsx_is_normalized_as_malformed_input(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def fail_with_crc_error(*_args: object) -> None:
        raise candidates_module.zipfile.BadZipFile("bad CRC")

    monkeypatch.setattr(candidates_module, "parse_listing_file", fail_with_crc_error)
    candidate = MonitoringDocumentCandidateDecomposer(
        tmp_path / "candidates"
    ).decompose("damaged.xlsx", _xlsx_bytes())

    assert candidate.extraction_status == "unreadable"
    assert candidate.limitation_codes == ("candidate_parse_failed",)


def test_candidate_batch_identity_does_not_depend_on_upload_order(tmp_path: Path) -> None:
    files = [
        ("protocol.docx", _docx_bytes("Protocol")),
        ("ecrf.xlsx", _xlsx_bytes()),
    ]
    decomposer = MonitoringDocumentCandidateDecomposer(tmp_path / "candidates")

    first = decomposer.decompose_many(files)
    second = decomposer.decompose_many(list(reversed(files)))

    assert first == second
    assert first.authority_status == "not_adjudicated"
    assert [candidate.candidate_id for candidate in first.candidates] == sorted(
        candidate.candidate_id for candidate in first.candidates
    )
    persisted = json.loads(
        (
            tmp_path
            / "candidates"
            / "batches"
            / f"{first.batch_id}.json"
        ).read_text(encoding="utf-8")
    )
    assert persisted == json.loads(json.dumps(first.to_dict(), ensure_ascii=False))
