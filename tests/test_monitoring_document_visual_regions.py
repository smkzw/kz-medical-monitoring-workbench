"""Synthetic regressions for bounded frozen-document visual input extraction."""
from __future__ import annotations

import io
import json
import math
import struct
import zlib
from hashlib import sha256
from types import SimpleNamespace

import pymupdf
import pytest

from packages.medical_monitoring.admission.source_tools import SourceToolError
from packages.medical_monitoring.intelligence.primitives import content_hash
from services.api.app.monitoring_document_visual_regions import (
    FrozenDocumentVisualRegionTools,
)
from services.api.app.protocol_text_extractor import parse_protocol_docx

_PDF_MEDIA = "application/pdf"
_DOCX_MEDIA = (
    "application/vnd.openxmlformats-officedocument."
    "wordprocessingml.document"
)
_XLSX_MEDIA = (
    "application/vnd.openxmlformats-officedocument."
    "spreadsheetml.sheet"
)


def color_png(rgb, size=4):
    """Minimal solid-color PNG so exact embedded bytes can be compared."""
    def chunk(tag, payload):
        block = struct.pack(">I", len(payload)) + tag + payload
        return block + struct.pack(">I", zlib.crc32(tag + payload) & 0xFFFFFFFF)

    header = struct.pack(">IIBBBBB", size, size, 8, 2, 0, 0, 0)
    row = b"\x00" + bytes(rgb) * size
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", header)
        + chunk(b"IDAT", zlib.compress(row * size))
        + chunk(b"IEND", b"")
    )


def freeze(tmp_path, content, suffix, media_type):
    digest = sha256(content).hexdigest()
    root = tmp_path / "copies"
    (root / "files").mkdir(parents=True, exist_ok=True)
    path = root / "files" / (digest + suffix)
    path.write_bytes(content)
    binding = SimpleNamespace(
        content_sha256=digest,
        media_type=media_type,
        source_entry_id="frozen-visual",
        locator_index_sha256="c" * 64,
    )
    checked = []
    resolver = SimpleNamespace(
        assert_current_binding=lambda **kw: checked.append(kw)
    )
    tools = FrozenDocumentVisualRegionTools(
        candidate_root=root,
        project_id="synthetic",
        input_revision="a" * 64,
        resolver=resolver,
    )
    return tools, binding, checked, path


def header_and_two_image_pages_pdf():
    """Page one: header text + green figure. Page two: red figure, no text."""
    green, red = color_png((0, 200, 0)), color_png((220, 0, 0))
    with pymupdf.open() as document:
        first = document.new_page()
        first.insert_text((40, 40), "HEADER frozen visual source")
        first.insert_image(pymupdf.Rect(200, 300, 300, 380), stream=green)
        second = document.new_page()
        second.insert_image(pymupdf.Rect(200, 300, 300, 380), stream=red)
        # new_page() invalidates earlier page handles, so re-grab page one.
        image_bbox = tuple(document[0].get_image_info()[0]["bbox"])
        content = document.tobytes()
    return content, image_bbox, green, red


def dark_pixel_count(png_bytes):
    with pymupdf.open(stream=png_bytes, filetype="png") as image:
        pixmap = image[0].get_pixmap()
    samples = pixmap.samples
    stride = pixmap.n
    return sum(
        1
        for offset in range(0, len(samples), stride)
        if min(samples[offset:offset + 3]) < 200
    )


def colored_pixel_count(png_bytes, predicate):
    with pymupdf.open(stream=png_bytes, filetype="png") as image:
        pixmap = image[0].get_pixmap()
    samples = pixmap.samples
    stride = pixmap.n
    return sum(
        1
        for offset in range(0, len(samples), stride)
        if predicate(samples[offset], samples[offset + 1], samples[offset + 2])
    )


def test_pdf_page_and_region_render_returns_bounded_facts(tmp_path):
    content, image_bbox, _, _ = header_and_two_image_pages_pdf()
    tools, binding, checked, _ = freeze(tmp_path, content, ".pdf", _PDF_MEDIA)
    page = tools.render_pdf_region(binding=binding, page_index=0, dpi=100)
    assert page.extraction_kind == "pdf_page_render"
    assert page.locator == "pdf:page:1"
    assert page.image.image_bytes.startswith(b"\x89PNG\r\n\x1a\n")
    assert page.image.image_sha256 == sha256(page.image.image_bytes).hexdigest()
    assert page.image.media_type == "image/png"
    assert page.page_size_points == (595.0, 842.0)
    assert page.requested_bbox is None
    assert page.render_dpi == 100
    assert page.pixel_width == math.ceil(595.0 * 100 / 72)
    assert page.pixel_height == math.ceil(842.0 * 100 / 72)
    assert page.source_entry_id == "frozen-visual"
    assert page.source_content_sha256 == binding.content_sha256
    assert page.semantic_interpretation == "not_performed"
    assert page.limitation_codes == (
        "visual_input_extract_only_no_ocr_or_semantics",
    )
    assert [call["project_id"] for call in checked] == ["synthetic", "synthetic"]
    assert [call["binding"] for call in checked] == [binding, binding]

    double = tools.render_pdf_region(binding=binding, page_index=0, dpi=200)
    assert double.pixel_width == math.ceil(595.0 * 200 / 72)
    assert double.render_dpi == 200

    region = tools.render_pdf_region(
        binding=binding,
        page_index=0,
        bbox=image_bbox,
        dpi=150,
    )
    assert region.extraction_kind == "pdf_region_render"
    assert region.locator == "pdf:page:1:region:200.00,300.00,300.00,380.00"
    assert region.requested_bbox == image_bbox
    assert region.image.image_sha256 != page.image.image_sha256
    assert region.pixel_width == math.ceil(100 * 150 / 72)
    assert region.pixel_height == math.ceil(80 * 150 / 72)
    # The requested region actually carries the figure: strong green content.
    assert colored_pixel_count(
        region.image.image_bytes,
        lambda r, g, b: g > 120 and r < 120,
    ) > 100

    blank = tools.render_pdf_region(
        binding=binding,
        page_index=0,
        bbox=(400.0, 600.0, 500.0, 700.0),
        dpi=150,
    )
    assert blank.locator == "pdf:page:1:region:400.00,600.00,500.00,700.00"
    assert blank.image.image_sha256 != region.image.image_sha256
    assert dark_pixel_count(blank.image.image_bytes) == 0

    repeat = tools.render_pdf_region(
        binding=binding,
        page_index=0,
        bbox=image_bbox,
        dpi=150,
    )
    assert repeat.image.image_bytes == region.image.image_bytes
    assert repeat.extraction_sha256 == region.extraction_sha256


def test_pdf_zero_based_page_index_renders_distinct_pages(tmp_path):
    content, image_bbox, _, _ = header_and_two_image_pages_pdf()
    tools, binding, _, _ = freeze(tmp_path, content, ".pdf", _PDF_MEDIA)
    first = tools.render_pdf_region(
        binding=binding, page_index=0, bbox=image_bbox, dpi=150
    )
    second = tools.render_pdf_region(
        binding=binding, page_index=1, bbox=image_bbox, dpi=150
    )
    assert second.locator == "pdf:page:2:region:200.00,300.00,300.00,380.00"
    assert first.image.image_sha256 != second.image.image_sha256
    assert colored_pixel_count(
        first.image.image_bytes, lambda r, g, b: g > 120 and r < 120
    ) > 100
    assert colored_pixel_count(
        second.image.image_bytes, lambda r, g, b: r > 120 and g < 120
    ) > 100
    with pytest.raises(SourceToolError, match="visual_page_index_out_of_range"):
        tools.render_pdf_region(binding=binding, page_index=2)
    with pytest.raises(SourceToolError, match="visual_page_index_out_of_range"):
        tools.render_pdf_region(binding=binding, page_index=-1)
    with pytest.raises(SourceToolError, match="visual_page_index_out_of_range"):
        tools.render_pdf_region(binding=binding, page_index=True)


def test_pdf_rejects_out_of_bounds_invalid_regions_and_budget_excess(tmp_path):
    content, _, _, _ = header_and_two_image_pages_pdf()
    tools, binding, _, _ = freeze(tmp_path, content, ".pdf", _PDF_MEDIA)
    with pytest.raises(SourceToolError, match="visual_region_out_of_page_bounds"):
        tools.render_pdf_region(
            binding=binding, page_index=0, bbox=(500.0, 700.0, 700.0, 900.0)
        )
    with pytest.raises(SourceToolError, match="visual_region_invalid"):
        tools.render_pdf_region(
            binding=binding, page_index=0, bbox=(100.0, 100.0, 100.0, 200.0)
        )
    with pytest.raises(SourceToolError, match="visual_region_invalid"):
        tools.render_pdf_region(
            binding=binding, page_index=0, bbox=(10.0, 10.0, 20.0)
        )
    with pytest.raises(SourceToolError, match="visual_region_invalid"):
        tools.render_pdf_region(
            binding=binding,
            page_index=0,
            bbox=(10.0, 10.0, float("nan"), 20.0),
        )
    with pytest.raises(SourceToolError, match="visual_render_dpi_out_of_range"):
        tools.render_pdf_region(binding=binding, page_index=0, dpi=600)
    with pytest.raises(SourceToolError, match="visual_render_dpi_out_of_range"):
        tools.render_pdf_region(binding=binding, page_index=0, dpi=True)

    # The pixel budget applies to the requested region, not the whole page:
    # the small corner of the oversized page renders, the full page cannot.
    with pymupdf.open() as document:
        document.new_page(width=2000, height=2000)
        oversized = document.tobytes()
    big_tools, big_binding, _, _ = freeze(
        tmp_path, oversized, ".pdf", _PDF_MEDIA
    )
    corner = big_tools.render_pdf_region(
        binding=big_binding, page_index=0, bbox=(0.0, 0.0, 100.0, 100.0), dpi=300
    )
    assert corner.extraction_kind == "pdf_region_render"
    with pytest.raises(
        SourceToolError, match="visual_region_pixel_budget_exceeded"
    ):
        big_tools.render_pdf_region(binding=big_binding, page_index=0, dpi=300)


def test_changed_missing_copy_and_stale_binding_cannot_render(tmp_path):
    content, _, _, _ = header_and_two_image_pages_pdf()
    tools, binding, checked, path = freeze(tmp_path, content, ".pdf", _PDF_MEDIA)
    path.write_bytes(b"changed")
    with pytest.raises(SourceToolError, match="frozen_document_content_changed"):
        tools.render_pdf_region(binding=binding, page_index=0)
    assert len(checked) == 1

    missing_binding = SimpleNamespace(
        content_sha256="0" * 64,
        media_type=_PDF_MEDIA,
        source_entry_id="frozen-visual",
        locator_index_sha256="c" * 64,
    )
    with pytest.raises(SourceToolError, match="frozen_document_unavailable"):
        tools.render_pdf_region(binding=missing_binding, page_index=0)

    path.write_bytes(content)

    def stale(**_):
        raise ValueError("monitoring document binding is no longer current")

    tools.resolver.assert_current_binding = stale
    with pytest.raises(ValueError, match="no longer current"):
        tools.render_pdf_region(binding=binding, page_index=0)


def test_word_locator_selects_exact_embedded_image(tmp_path):
    green, red = color_png((0, 200, 0)), color_png((220, 0, 0))
    assert sha256(green).hexdigest() != sha256(red).hexdigest()
    import docx as python_docx

    document = python_docx.Document()
    document.add_picture(io.BytesIO(green), width=python_docx.shared.Inches(1))
    document.add_picture(io.BytesIO(red), width=python_docx.shared.Inches(1))
    buffer = io.BytesIO()
    document.save(buffer)
    tools, binding, checked, _ = freeze(
        tmp_path, buffer.getvalue(), ".docx", _DOCX_MEDIA
    )
    parsed = parse_protocol_docx("frozen.docx", buffer.getvalue())
    locators = [image.source_locator for image in parsed.embedded_images]
    assert len(locators) == 2 and len(set(locators)) == 2

    # Request the second image first: an unknown locator must never fall
    # back to any other embedded image, and order must not matter.
    second = tools.extract_word_embedded_image(
        binding=binding, source_locator=locators[1]
    )
    assert second.extraction_kind == "word_embedded_image"
    assert second.locator == locators[1]
    assert second.image.image_bytes == red
    assert second.image.image_sha256 == sha256(red).hexdigest()
    assert second.image.media_type == "image/png"
    assert second.word_part_name == "word/media/image2.png"
    assert second.page_size_points is None
    assert second.render_dpi is None
    assert second.semantic_interpretation == "not_performed"

    first = tools.extract_word_embedded_image(
        binding=binding, source_locator=locators[0]
    )
    assert first.image.image_bytes == green
    assert first.word_part_name == "word/media/image1.png"
    assert first.image.image_sha256 != second.image.image_sha256
    assert len(checked) == 4  # two successful extractions, re-verified each

    with pytest.raises(
        SourceToolError, match="word_embedded_image_locator_unknown"
    ):
        tools.extract_word_embedded_image(
            binding=binding, source_locator="docx:drawing:99:0"
        )
    with pytest.raises(
        SourceToolError, match="word_embedded_image_locator_unknown"
    ):
        tools.extract_word_embedded_image(binding=binding, source_locator="  ")


def test_public_view_and_repr_exclude_image_bytes(tmp_path):
    content, image_bbox, _, _ = header_and_two_image_pages_pdf()
    tools, binding, _, _ = freeze(tmp_path, content, ".pdf", _PDF_MEDIA)
    region = tools.render_pdf_region(
        binding=binding, page_index=0, bbox=image_bbox, dpi=150
    )
    public = region.public_dict()
    dumped = json.dumps(public)
    assert "image_bytes" not in dumped
    assert public["image_sha256"] == region.image.image_sha256
    assert public["extraction_kind"] == "pdf_region_render"
    assert b"\x89PNG" not in repr(region).encode("utf-8", "replace")
    assert "image_bytes=" not in repr(region)

    payload = dict(public)
    payload.pop("extraction_sha256")
    assert region.extraction_sha256 == content_hash(payload)


def test_unsupported_media_types_and_cross_format_calls_rejected(tmp_path):
    content, _, _, _ = header_and_two_image_pages_pdf()
    tools, binding, _, _ = freeze(tmp_path, content, ".pdf", _PDF_MEDIA)
    docx_binding = SimpleNamespace(
        content_sha256=binding.content_sha256,
        media_type=_DOCX_MEDIA,
        source_entry_id="frozen-visual",
        locator_index_sha256="c" * 64,
    )
    xlsx_binding = SimpleNamespace(
        content_sha256=binding.content_sha256,
        media_type=_XLSX_MEDIA,
        source_entry_id="frozen-visual",
        locator_index_sha256="c" * 64,
    )
    with pytest.raises(SourceToolError, match="document_format_not_supported"):
        tools.render_pdf_region(binding=docx_binding, page_index=0)
    with pytest.raises(SourceToolError, match="document_format_not_supported"):
        tools.render_pdf_region(binding=xlsx_binding, page_index=0)
    with pytest.raises(SourceToolError, match="document_format_not_supported"):
        tools.extract_word_embedded_image(
            binding=binding, source_locator="docx:drawing:0:0"
        )
    with pytest.raises(SourceToolError, match="document_format_not_supported"):
        tools.extract_word_embedded_image(
            binding=xlsx_binding, source_locator="docx:drawing:0:0"
        )


def test_rotated_pdf_region_uses_displayed_page_coordinates(tmp_path):
    with pymupdf.open() as document:
        page = document.new_page(width=100, height=200)
        page.draw_rect(pymupdf.Rect(10, 10, 40, 40), color=(1, 0, 0), fill=(1, 0, 0))
        page.set_rotation(90)
        content = document.tobytes()
    tools, binding, _, _ = freeze(tmp_path, content, ".pdf", _PDF_MEDIA)
    region = tools.render_pdf_region(binding=binding, page_index=0, bbox=(160, 0, 200, 40), dpi=72)
    pixmap = pymupdf.Pixmap(region.image.image_bytes)
    assert (pixmap.width, pixmap.height) == (40, 40)
    assert pixmap.pixel(20, 20) == (255, 0, 0)


def test_word_visual_output_obeys_the_same_byte_budget(tmp_path, monkeypatch):
    import docx as python_docx
    import services.api.app.monitoring_document_visual_regions as module
    picture = color_png((0, 200, 0))
    document = python_docx.Document()
    document.add_picture(io.BytesIO(picture), width=python_docx.shared.Inches(1))
    buffer = io.BytesIO()
    document.save(buffer)
    tools, binding, _, _ = freeze(tmp_path, buffer.getvalue(), '.docx', _DOCX_MEDIA)
    locator = parse_protocol_docx('frozen.docx', buffer.getvalue()).embedded_images[0].source_locator
    monkeypatch.setattr(module, '_MAX_IMAGE_BYTES', len(picture) - 1)
    with pytest.raises(SourceToolError, match='visual_region_bytes_budget_exceeded'):
        tools.extract_word_embedded_image(binding=binding, source_locator=locator)
