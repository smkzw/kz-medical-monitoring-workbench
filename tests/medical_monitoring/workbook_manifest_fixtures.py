"""Pure-synthetic workbook builders for the C3 physical-integrity matrix.

Work item 3 of ``mm-c3-workbook-manifest-p0-20260903``.  Every workbook is
built purely in memory (openpyxl plus targeted OOXML zip patches for cached
formula results, which openpyxl never writes) and declares the physical facts
it embeds as ``ground_truth``.  No real project file is read, no service is
started, and no model route is touched anywhere in this module.

The independent OOXML oracle lives in :mod:`tests.medical_monitoring.workbook_manifest_oracle`
and the manifest contract / fail-closed gate in
:mod:`tests.medical_monitoring.workbook_manifest_contract`.
"""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any, Dict, List, Mapping

import openpyxl
from openpyxl.worksheet.table import Table

@dataclass(frozen=True)
class WorkbookCase:
    """One synthetic workbook plus the physical facts it was built to carry."""

    case_id: str
    filename: str
    content: bytes
    ground_truth: Dict[str, Any]


def _save_workbook(workbook: openpyxl.Workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def _patch_sheet_xml(content: bytes, replacements: Mapping[str, str]) -> bytes:
    """Replace exact serialized fragments inside worksheet XML parts.

    Used to embed cached formula results, which openpyxl never writes: a
    native openpyxl formula cell serializes as ``<c r="B2"><f>SUM(1,2)</f>
    <v></v></c>`` (an empty ``<v>``).  Every replacement must hit exactly
    once, so an openpyxl serialization change fails loudly instead of
    silently weakening a case.
    """
    with zipfile.ZipFile(io.BytesIO(content)) as source:
        entries = [(info.filename, source.read(info.filename)) for info in source.infolist()]
    patched_any = False
    output_entries: List[Tuple[str, bytes]] = []
    for name, payload in entries:
        if name.startswith("xl/worksheets/") and name.endswith(".xml"):
            text = payload.decode("utf-8")
            for old, new in replacements.items():
                hits = text.count(old)
                if hits:
                    if hits != 1:
                        raise AssertionError(
                            f"expected exactly one occurrence of {old!r}, found {hits}"
                        )
                    text = text.replace(old, new)
                    patched_any = True
            payload = text.encode("utf-8")
        output_entries.append((name, payload))
    if not patched_any:
        raise AssertionError(
            "no replacement matched; openpyxl serialization changed and the "
            "cached-formula patch must be updated"
        )
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for name, payload in output_entries:
            target.writestr(name, payload)
    return output.getvalue()


# ---------------------------------------------------------------------------
# Case builders (pure synthetic, in-memory)
# ---------------------------------------------------------------------------

def build_sheet_visibility_case() -> WorkbookCase:
    workbook = openpyxl.Workbook()
    plain = workbook.active
    plain.title = "PLAIN"
    plain.append(["SUBJID", "VISIT"])
    plain.append(["S01", "V1"])
    sealed = workbook.create_sheet("SEALED")
    sealed.sheet_state = "hidden"
    sealed.append(["SUBJID", "VISIT"])
    sealed.append(["S02", "V2"])
    ghost = workbook.create_sheet("GHOST")
    ghost.sheet_state = "veryHidden"
    ghost.append(["SUBJID", "VISIT"])
    ghost.append(["S03", "V3"])
    return WorkbookCase(
        case_id="sheet_visibility",
        filename="mm_c3_matrix_visibility.xlsx",
        content=_save_workbook(workbook),
        ground_truth={
            "sheet_order": ["PLAIN", "SEALED", "GHOST"],
            "sheet_visibility": {"PLAIN": "visible", "SEALED": "hidden", "GHOST": "veryHidden"},
            "sheets": {
                name: {
                    "used_range": "A1:B2",
                    "physical_row_count": 2,
                    "physical_column_count": 2,
                    "merged_cell_ranges": [],
                    "hidden_row_numbers": [],
                    "hidden_column_letters": [],
                    "named_tables": [],
                    "autofilter_ref": None,
                    "formula_cells": [],
                    "number_format_cells": {},
                    "leading_zero_text_cells": [],
                    "duplicate_source_headers": [],
                }
                for name in ("PLAIN", "SEALED", "GHOST")
            },
        },
    )


def build_empty_sheet_case() -> WorkbookCase:
    workbook = openpyxl.Workbook()
    data = workbook.active
    data.title = "DATA"
    data.append(["SUBJID", "VISIT"])
    data.append(["S01", "V1"])
    workbook.create_sheet("EMPTY")
    return WorkbookCase(
        case_id="empty_sheet",
        filename="mm_c3_matrix_empty_sheet.xlsx",
        content=_save_workbook(workbook),
        ground_truth={
            "sheet_order": ["DATA", "EMPTY"],
            "sheet_visibility": {"DATA": "visible", "EMPTY": "visible"},
            "sheets": {
                "DATA": {
                    "used_range": "A1:B2",
                    "physical_row_count": 2,
                    "physical_column_count": 2,
                    "merged_cell_ranges": [],
                    "hidden_row_numbers": [],
                    "hidden_column_letters": [],
                    "named_tables": [],
                    "autofilter_ref": None,
                    "formula_cells": [],
                    "number_format_cells": {},
                    "leading_zero_text_cells": [],
                    "duplicate_source_headers": [],
                },
                "EMPTY": {
                    "used_range": None,
                    "physical_row_count": 0,
                    "physical_column_count": 0,
                    "merged_cell_ranges": [],
                    "hidden_row_numbers": [],
                    "hidden_column_letters": [],
                    "named_tables": [],
                    "autofilter_ref": None,
                    "formula_cells": [],
                    "number_format_cells": {},
                    "leading_zero_text_cells": [],
                    "duplicate_source_headers": [],
                },
            },
        },
    )


def build_header_only_case() -> WorkbookCase:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "HEADONLY"
    sheet.append(["SUBJID", "AETERM", "AESTDTC"])
    return WorkbookCase(
        case_id="header_only",
        filename="mm_c3_matrix_header_only.xlsx",
        content=_save_workbook(workbook),
        ground_truth={
            "sheet_order": ["HEADONLY"],
            "sheet_visibility": {"HEADONLY": "visible"},
            "sheets": {
                "HEADONLY": {
                    "used_range": "A1:C1",
                    "physical_row_count": 1,
                    "physical_column_count": 3,
                    "merged_cell_ranges": [],
                    "hidden_row_numbers": [],
                    "hidden_column_letters": [],
                    "named_tables": [],
                    "autofilter_ref": None,
                    "formula_cells": [],
                    "number_format_cells": {},
                    "leading_zero_text_cells": [],
                    "duplicate_source_headers": [],
                },
            },
            # Declared intent (semantic, not oracle-derivable).
            "header_row_numbers": {"HEADONLY": [1]},
            "is_header_only": {"HEADONLY": True},
        },
    )


def build_multilayer_merged_headers_case() -> WorkbookCase:
    workbook = openpyxl.Workbook()
    main = workbook.active
    main.title = "MAIN"
    main["A1"] = "受试者信息"
    main.merge_cells("A1:B1")
    main["C1"] = "不良事件"
    main.merge_cells("C1:E1")
    main["F1"] = "备注"
    main.merge_cells("F1:F2")
    for column, header in zip("ABCDE", ["SITEID", "SUBJID", "AETERM", "AESTDTC", "AESEV"]):
        main[f"{column}2"] = header
    main["A3"] = "0101"
    main["B3"] = "S01001"
    main["C3"] = "头痛"
    # Text-stored dates on purpose: assigning a datetime.date would make
    # openpyxl silently attach its default date number format and the case
    # would no longer be format-free.
    main["D3"] = "2026-07-13"
    main["E3"] = "1级"
    main["F3"] = "与研究药物可能有关"
    main["A4"] = "0102"
    main["B4"] = "S01002"
    main["C4"] = "恶心"
    main["D4"] = "2026-07-14"
    main["E4"] = "2级"
    return WorkbookCase(
        case_id="multilayer_merged_headers",
        filename="mm_c3_matrix_merged_headers.xlsx",
        content=_save_workbook(workbook),
        ground_truth={
            "sheet_order": ["MAIN"],
            "sheet_visibility": {"MAIN": "visible"},
            "sheets": {
                "MAIN": {
                    "used_range": "A1:F4",
                    "physical_row_count": 4,
                    "physical_column_count": 6,
                    "merged_cell_ranges": ["A1:B1", "C1:E1", "F1:F2"],
                    "hidden_row_numbers": [],
                    "hidden_column_letters": [],
                    "named_tables": [],
                    "autofilter_ref": None,
                    "formula_cells": [],
                    "number_format_cells": {},
                    "leading_zero_text_cells": [
                        {"coordinate": "A3", "value": "0101"},
                        {"coordinate": "A4", "value": "0102"},
                    ],
                    "duplicate_source_headers": [],
                },
            },
            "header_row_numbers": {"MAIN": [1, 2]},
            "is_header_only": {"MAIN": False},
        },
    )


def build_hidden_rows_columns_case() -> WorkbookCase:
    workbook = openpyxl.Workbook()
    main = workbook.active
    main.title = "MAIN"
    main.append(["SUBJID", "VISIT", "AVAL"])
    main.append(["S01", "V1", 120])
    main.append(["S02", "V1", 118])
    main.append(["S03", "V2", 96])
    main.row_dimensions[3].hidden = True
    main.column_dimensions["B"].hidden = True
    return WorkbookCase(
        case_id="hidden_rows_columns",
        filename="mm_c3_matrix_hidden_rows_columns.xlsx",
        content=_save_workbook(workbook),
        ground_truth={
            "sheet_order": ["MAIN"],
            "sheet_visibility": {"MAIN": "visible"},
            "sheets": {
                "MAIN": {
                    "used_range": "A1:C4",
                    "physical_row_count": 4,
                    "physical_column_count": 3,
                    "merged_cell_ranges": [],
                    "hidden_row_numbers": [3],
                    "hidden_column_letters": ["B"],
                    "named_tables": [],
                    "autofilter_ref": None,
                    "formula_cells": [],
                    "number_format_cells": {},
                    "leading_zero_text_cells": [],
                    "duplicate_source_headers": [],
                },
            },
            "header_row_numbers": {"MAIN": [1]},
            "is_header_only": {"MAIN": False},
        },
    )


def build_named_table_case() -> WorkbookCase:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "TBL"
    sheet.append(["SUBJID", "VISIT"])
    sheet.append(["S01", "V1"])
    sheet.append(["S02", "V2"])
    sheet.add_table(Table(displayName="AEData", ref="A1:B3"))
    return WorkbookCase(
        case_id="named_table",
        filename="mm_c3_matrix_named_table.xlsx",
        content=_save_workbook(workbook),
        ground_truth={
            "sheet_order": ["TBL"],
            "sheet_visibility": {"TBL": "visible"},
            "sheets": {
                "TBL": {
                    "used_range": "A1:B3",
                    "physical_row_count": 3,
                    "physical_column_count": 2,
                    "merged_cell_ranges": [],
                    "hidden_row_numbers": [],
                    "hidden_column_letters": [],
                    "named_tables": [{"name": "AEData", "ref": "A1:B3"}],
                    # A named table carries its own autoFilter inside
                    # xl/tables/tableN.xml; the worksheet-level autoFilter is
                    # absent, and the manifest must not confuse the two.
                    "autofilter_ref": None,
                    "formula_cells": [],
                    "number_format_cells": {},
                    "leading_zero_text_cells": [],
                    "duplicate_source_headers": [],
                },
            },
            "header_row_numbers": {"TBL": [1]},
            "is_header_only": {"TBL": False},
        },
    )


def build_formula_cache_case() -> WorkbookCase:
    workbook = openpyxl.Workbook()
    calc = workbook.active
    calc.title = "CALC"
    calc["A1"] = "SUBJID"
    calc["B1"] = "AVAL"
    calc["C1"] = "NOTE"
    calc["A2"] = "S01"
    calc["B2"] = "=SUM(1,2)"   # native openpyxl: formula without cached value
    calc["B3"] = "=1+2"        # patched below to carry cached numeric value 3
    calc["C4"] = "=A2"         # patched below to carry cached string value S01
    content = _save_workbook(workbook)
    content = _patch_sheet_xml(content, {
        '<c r="B3"><f>1+2</f><v></v></c>': '<c r="B3"><f>1+2</f><v>3</v></c>',
        '<c r="C4"><f>A2</f><v></v></c>': '<c r="C4" t="str"><f>A2</f><v>S01</v></c>',
    })
    return WorkbookCase(
        case_id="formula_cache",
        filename="mm_c3_matrix_formula_cache.xlsx",
        content=content,
        ground_truth={
            "sheet_order": ["CALC"],
            "sheet_visibility": {"CALC": "visible"},
            "sheets": {
                "CALC": {
                    "used_range": "A1:C4",
                    "physical_row_count": 4,
                    "physical_column_count": 3,
                    "merged_cell_ranges": [],
                    "hidden_row_numbers": [],
                    "hidden_column_letters": [],
                    "named_tables": [],
                    "autofilter_ref": None,
                    "formula_cells": [
                        {
                            "coordinate": "B2",
                            "formula_text": "=SUM(1,2)",
                            "cached_value_present": False,
                            "cached_value": None,
                        },
                        {
                            "coordinate": "B3",
                            "formula_text": "=1+2",
                            "cached_value_present": True,
                            "cached_value": "3",
                        },
                        {
                            "coordinate": "C4",
                            "formula_text": "=A2",
                            "cached_value_present": True,
                            "cached_value": "S01",
                        },
                    ],
                    "number_format_cells": {},
                    "leading_zero_text_cells": [],
                    "duplicate_source_headers": [],
                },
            },
            "header_row_numbers": {"CALC": [1]},
            "is_header_only": {"CALC": False},
        },
    )


def build_date_formats_case() -> WorkbookCase:
    workbook = openpyxl.Workbook()
    dates = workbook.active
    dates.title = "DATES"
    dates.append(["SUBJID", "D_ISO", "D_SLASH", "D_CN", "D_TIME"])
    dates["A2"] = "S01"
    dates["B2"] = date(2026, 7, 13)
    dates["B2"].number_format = "yyyy-mm-dd"
    dates["C2"] = date(2026, 7, 14)
    dates["C2"].number_format = "yyyy/m/d"
    dates["D2"] = date(2026, 7, 15)
    dates["D2"].number_format = 'yyyy"年"m"月"d"日"'
    dates["E2"] = datetime(2026, 7, 16, 8, 30)
    dates["E2"].number_format = "yyyy-mm-dd hh:mm"
    return WorkbookCase(
        case_id="date_formats",
        filename="mm_c3_matrix_date_formats.xlsx",
        content=_save_workbook(workbook),
        ground_truth={
            "sheet_order": ["DATES"],
            "sheet_visibility": {"DATES": "visible"},
            "sheets": {
                "DATES": {
                    "used_range": "A1:E2",
                    "physical_row_count": 2,
                    "physical_column_count": 5,
                    "merged_cell_ranges": [],
                    "hidden_row_numbers": [],
                    "hidden_column_letters": [],
                    "named_tables": [],
                    "autofilter_ref": None,
                    "formula_cells": [],
                    "number_format_cells": {
                        "B2": "yyyy-mm-dd",
                        "C2": "yyyy/m/d",
                        "D2": 'yyyy"年"m"月"d"日"',
                        "E2": "yyyy-mm-dd hh:mm",
                    },
                    "leading_zero_text_cells": [],
                    "duplicate_source_headers": [],
                },
            },
            "header_row_numbers": {"DATES": [1]},
            "is_header_only": {"DATES": False},
        },
    )


def build_leading_zeros_case() -> WorkbookCase:
    workbook = openpyxl.Workbook()
    ids = workbook.active
    ids.title = "IDS"
    ids.append(["SUBJID", "SITEID", "VISIT"])
    ids.append(["0012", "010-10008", "W1"])
    ids.append(["0013", "020-10009", "W2"])
    return WorkbookCase(
        case_id="leading_zeros",
        filename="mm_c3_matrix_leading_zeros.xlsx",
        content=_save_workbook(workbook),
        ground_truth={
            "sheet_order": ["IDS"],
            "sheet_visibility": {"IDS": "visible"},
            "sheets": {
                "IDS": {
                    "used_range": "A1:C3",
                    "physical_row_count": 3,
                    "physical_column_count": 3,
                    "merged_cell_ranges": [],
                    "hidden_row_numbers": [],
                    "hidden_column_letters": [],
                    "named_tables": [],
                    "autofilter_ref": None,
                    "formula_cells": [],
                    "number_format_cells": {},
                    "leading_zero_text_cells": [
                        {"coordinate": "A2", "value": "0012"},
                        {"coordinate": "A3", "value": "0013"},
                        {"coordinate": "B2", "value": "010-10008"},
                        {"coordinate": "B3", "value": "020-10009"},
                    ],
                    "duplicate_source_headers": [],
                },
            },
            "header_row_numbers": {"IDS": [1]},
            "is_header_only": {"IDS": False},
        },
    )


def build_duplicate_columns_case() -> WorkbookCase:
    workbook = openpyxl.Workbook()
    dupl = workbook.active
    dupl.title = "DUPL"
    dupl.append(["SUBJID", "VISIT", "SUBJID"])
    dupl.append(["0012", "W1", "010-10008"])
    dupl.append(["0013", "W2", "010-10009"])
    return WorkbookCase(
        case_id="duplicate_columns",
        filename="mm_c3_matrix_duplicate_columns.xlsx",
        content=_save_workbook(workbook),
        ground_truth={
            "sheet_order": ["DUPL"],
            "sheet_visibility": {"DUPL": "visible"},
            "sheets": {
                "DUPL": {
                    "used_range": "A1:C3",
                    "physical_row_count": 3,
                    "physical_column_count": 3,
                    "merged_cell_ranges": [],
                    "hidden_row_numbers": [],
                    "hidden_column_letters": [],
                    "named_tables": [],
                    "autofilter_ref": None,
                    "formula_cells": [],
                    "number_format_cells": {},
                    "leading_zero_text_cells": [
                        {"coordinate": "A2", "value": "0012"},
                        {"coordinate": "A3", "value": "0013"},
                        {"coordinate": "C2", "value": "010-10008"},
                        {"coordinate": "C3", "value": "010-10009"},
                    ],
                    "duplicate_source_headers": [
                        {"header_text": "SUBJID", "column_numbers": [1, 3]},
                    ],
                },
            },
            "header_row_numbers": {"DUPL": [1]},
            "is_header_only": {"DUPL": False},
        },
    )


def build_worksheet_autofilter_case() -> WorkbookCase:
    workbook = openpyxl.Workbook()
    filtered = workbook.active
    filtered.title = "FILTERED"
    filtered.append(["SUBJID", "VISIT"])
    filtered.append(["S01", "V1"])
    filtered.append(["S02", "V2"])
    filtered.auto_filter.ref = "A1:B3"
    return WorkbookCase(
        case_id="worksheet_autofilter",
        filename="mm_c3_matrix_autofilter.xlsx",
        content=_save_workbook(workbook),
        ground_truth={
            "sheet_order": ["FILTERED"],
            "sheet_visibility": {"FILTERED": "visible"},
            "sheets": {
                "FILTERED": {
                    "used_range": "A1:B3",
                    "physical_row_count": 3,
                    "physical_column_count": 2,
                    "merged_cell_ranges": [],
                    "hidden_row_numbers": [],
                    "hidden_column_letters": [],
                    "named_tables": [],
                    "autofilter_ref": "A1:B3",
                    "formula_cells": [],
                    "number_format_cells": {},
                    "leading_zero_text_cells": [],
                    "duplicate_source_headers": [],
                },
            },
            "header_row_numbers": {"FILTERED": [1]},
            "is_header_only": {"FILTERED": False},
        },
    )


def build_combined_matrix_case() -> WorkbookCase:
    """One workbook embedding every matrix feature, for whole-manifest
    reconciliation and for the omission gate tests."""
    workbook = openpyxl.Workbook()
    main = workbook.active
    main.title = "MAIN"
    # Two-layer merged header (rows 1-2).
    main["A1"] = "受试者信息"
    main.merge_cells("A1:B1")
    main["C1"] = "检查结果"
    main.merge_cells("C1:D1")
    main["E1"] = "标识"
    main.merge_cells("E1:F1")
    for column, header in zip("ABCDEF", ["SITEID", "SUBJID", "AESTDTC", "AVAL", "VISIT", "LBTEST"]):
        main[f"{column}2"] = header
    main["A3"] = "0101"
    main["B3"] = "0012"
    main["C3"] = date(2026, 7, 13)
    main["C3"].number_format = "yyyy-mm-dd"
    main["D3"] = 3.14
    main["D3"].number_format = "0.00"
    main["E3"] = "W1"
    main["F3"] = "WBC"
    main["A4"] = "0101"
    main["B4"] = "0013"
    main["C4"] = date(2026, 7, 14)
    main["C4"].number_format = "yyyy/m/d"
    main["D4"] = 2.5
    main["D4"].number_format = "0.00"
    main["E4"] = "W2"
    main["F4"] = "PLT"
    main["A5"] = "0102"
    main["B5"] = "=B4"       # patched below to cached string "0013"
    main["E5"] = "W3"
    main["F5"] = "=LEN(B4)"  # stays without cached value
    main.row_dimensions[4].hidden = True
    main.column_dimensions["D"].hidden = True
    main.auto_filter.ref = "A2:F5"
    dupl = workbook.create_sheet("DUPL")
    dupl.append(["SUBJID", "VISIT", "SUBJID"])
    dupl.append(["0012", "W1", "010-10008"])
    headonly = workbook.create_sheet("HEADONLY")
    headonly.append(["SUBJID", "AETERM"])
    workbook.create_sheet("EMPTY")
    hidden = workbook.create_sheet("HIDDEN")
    hidden.sheet_state = "hidden"
    hidden.append(["SUBJID"])
    hidden.append(["S09"])
    ghost = workbook.create_sheet("GHOST")
    ghost.sheet_state = "veryHidden"
    ghost.append(["SUBJID"])
    ghost.append(["S10"])
    tbl = workbook.create_sheet("TBL")
    tbl.append(["SUBJID", "VISIT"])
    tbl.append(["S01", "V1"])
    tbl.add_table(Table(displayName="LBData", ref="A1:B2"))
    content = _save_workbook(workbook)
    content = _patch_sheet_xml(content, {
        '<c r="B5"><f>B4</f><v></v></c>': '<c r="B5" t="str"><f>B4</f><v>0013</v></c>',
    })
    return WorkbookCase(
        case_id="combined_matrix",
        filename="mm_c3_matrix_combined.xlsx",
        content=content,
        ground_truth={
            "sheet_order": ["MAIN", "DUPL", "HEADONLY", "EMPTY", "HIDDEN", "GHOST", "TBL"],
            "sheet_visibility": {
                "MAIN": "visible", "DUPL": "visible", "HEADONLY": "visible",
                "EMPTY": "visible", "HIDDEN": "hidden", "GHOST": "veryHidden",
                "TBL": "visible",
            },
            "sheets": {
                "MAIN": {
                    "used_range": "A1:F5",
                    "physical_row_count": 5,
                    "physical_column_count": 6,
                    "merged_cell_ranges": ["A1:B1", "C1:D1", "E1:F1"],
                    "hidden_row_numbers": [4],
                    "hidden_column_letters": ["D"],
                    "named_tables": [],
                    "autofilter_ref": "A2:F5",
                    "formula_cells": [
                        {
                            "coordinate": "B5",
                            "formula_text": "=B4",
                            "cached_value_present": True,
                            "cached_value": "0013",
                        },
                        {
                            "coordinate": "F5",
                            "formula_text": "=LEN(B4)",
                            "cached_value_present": False,
                            "cached_value": None,
                        },
                    ],
                    "number_format_cells": {
                        "C3": "yyyy-mm-dd",
                        "C4": "yyyy/m/d",
                        "D3": "0.00",
                        "D4": "0.00",
                    },
                    "leading_zero_text_cells": [
                        {"coordinate": "A3", "value": "0101"},
                        {"coordinate": "A4", "value": "0101"},
                        {"coordinate": "A5", "value": "0102"},
                        {"coordinate": "B3", "value": "0012"},
                        {"coordinate": "B4", "value": "0013"},
                    ],
                    "duplicate_source_headers": [],
                },
                "DUPL": {
                    "used_range": "A1:C2",
                    "physical_row_count": 2,
                    "physical_column_count": 3,
                    "merged_cell_ranges": [],
                    "hidden_row_numbers": [],
                    "hidden_column_letters": [],
                    "named_tables": [],
                    "autofilter_ref": None,
                    "formula_cells": [],
                    "number_format_cells": {},
                    "leading_zero_text_cells": [
                        {"coordinate": "A2", "value": "0012"},
                        {"coordinate": "C2", "value": "010-10008"},
                    ],
                    "duplicate_source_headers": [
                        {"header_text": "SUBJID", "column_numbers": [1, 3]},
                    ],
                },
                "HEADONLY": {
                    "used_range": "A1:B1",
                    "physical_row_count": 1,
                    "physical_column_count": 2,
                    "merged_cell_ranges": [],
                    "hidden_row_numbers": [],
                    "hidden_column_letters": [],
                    "named_tables": [],
                    "autofilter_ref": None,
                    "formula_cells": [],
                    "number_format_cells": {},
                    "leading_zero_text_cells": [],
                    "duplicate_source_headers": [],
                },
                "EMPTY": {
                    "used_range": None,
                    "physical_row_count": 0,
                    "physical_column_count": 0,
                    "merged_cell_ranges": [],
                    "hidden_row_numbers": [],
                    "hidden_column_letters": [],
                    "named_tables": [],
                    "autofilter_ref": None,
                    "formula_cells": [],
                    "number_format_cells": {},
                    "leading_zero_text_cells": [],
                    "duplicate_source_headers": [],
                },
                "HIDDEN": {
                    "used_range": "A1:A2",
                    "physical_row_count": 2,
                    "physical_column_count": 1,
                    "merged_cell_ranges": [],
                    "hidden_row_numbers": [],
                    "hidden_column_letters": [],
                    "named_tables": [],
                    "autofilter_ref": None,
                    "formula_cells": [],
                    "number_format_cells": {},
                    "leading_zero_text_cells": [],
                    "duplicate_source_headers": [],
                },
                "GHOST": {
                    "used_range": "A1:A2",
                    "physical_row_count": 2,
                    "physical_column_count": 1,
                    "merged_cell_ranges": [],
                    "hidden_row_numbers": [],
                    "hidden_column_letters": [],
                    "named_tables": [],
                    "autofilter_ref": None,
                    "formula_cells": [],
                    "number_format_cells": {},
                    "leading_zero_text_cells": [],
                    "duplicate_source_headers": [],
                },
                "TBL": {
                    "used_range": "A1:B2",
                    "physical_row_count": 2,
                    "physical_column_count": 2,
                    "merged_cell_ranges": [],
                    "hidden_row_numbers": [],
                    "hidden_column_letters": [],
                    "named_tables": [{"name": "LBData", "ref": "A1:B2"}],
                    "autofilter_ref": None,
                    "formula_cells": [],
                    "number_format_cells": {},
                    "leading_zero_text_cells": [],
                    "duplicate_source_headers": [],
                },
            },
            "header_row_numbers": {
                "MAIN": [1, 2], "DUPL": [1], "HEADONLY": [1],
                "HIDDEN": [1], "GHOST": [1], "TBL": [1],
            },
            "is_header_only": {
                "MAIN": False, "DUPL": False, "HEADONLY": True,
                "HIDDEN": False, "GHOST": False, "TBL": False,
            },
        },
    )


def build_all_cases() -> List[WorkbookCase]:
    return [
        build_sheet_visibility_case(),
        build_empty_sheet_case(),
        build_header_only_case(),
        build_multilayer_merged_headers_case(),
        build_hidden_rows_columns_case(),
        build_named_table_case(),
        build_formula_cache_case(),
        build_date_formats_case(),
        build_leading_zeros_case(),
        build_duplicate_columns_case(),
        build_worksheet_autofilter_case(),
        build_combined_matrix_case(),
    ]


# ---------------------------------------------------------------------------
# Independent OOXML oracle (stdlib only, no openpyxl on the read path)
# ---------------------------------------------------------------------------
