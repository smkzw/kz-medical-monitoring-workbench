from __future__ import annotations

from copy import deepcopy
from dataclasses import replace
from datetime import datetime, timezone
import hashlib
import io
import json
from pathlib import Path
from types import SimpleNamespace
from typing import Sequence

import pytest
import docx

from packages.medical_monitoring.admission.document_evidence import (
    DOCUMENT_ROLES,
    CurrentDocumentBinding,
    DocumentEvidenceError,
    DocumentRoleEvidence,
    MonitoringDocumentEvidencePacket,
    validate_document_evidence_packet,
)
from packages.medical_monitoring.admission import (
    AdmissionMappingPipeline,
    DataAdmissionPipeline,
    MappingBridgeError,
    admission_record_to_harness_input,
)
from packages.medical_monitoring.admission.mapping_gate import (
    MONITORING_C3_MAPPING_MODEL,
    MONITORING_C3_MAPPING_PROFILE_ID,
    MONITORING_C3_MAPPING_PROVIDER,
    MONITORING_C3_VERIFIER_MODEL,
    MONITORING_C3_VERIFIER_PROVIDER,
)
from packages.medical_monitoring.admission.mapping_pipeline import (
    AdmissionMappingPipelineError,
)
from packages.medical_monitoring.graph.store import Store
from packages.medical_monitoring.runtime.runtime_progress import (
    ARTIFACT_DIR_NAME,
    RUNTIME_DB_NAME,
    RUNTIME_DIR_NAME,
)
from packages.medical_monitoring.intelligence.primitives import content_hash
from services.api.app.listing_file_parser import parse_listing_file
from services.api.app.protocol_text_extractor import (
    ProtocolTextDocument,
    ProtocolTextSpan,
)
from services.api.app.monitoring_document_evidence import (
    MonitoringDocumentEvidenceResolver,
)
from services.api.app.source_intake import (
    SourceRegistryService,
    SourceRegistryStore,
)
from services.api.app.source_content_validation import (
    SourceContentValidationService,
    SourceContentValidationStore,
    SourceExpectedContext,
)
from services.api.app.monitoring_ai_contracts import (
    MonitoringAiInputRevision,
    MonitoringAiTaskType,
)
from services.api.app.monitoring_ai_repository import MonitoringAiRepository
from services.api.app.monitoring_ai_service import (
    MonitoringAiRuntimeBinding,
    MonitoringAiService,
)
from tests.medical_monitoring.relationship_profiler_stub import (
    build_relationship_profile,
)


PROJECT_ID = "document-evidence-demo"


def _binding(role: str, token: str) -> CurrentDocumentBinding:
    return CurrentDocumentBinding(
        role=role,
        source_entry_id=f"source-{token}",
        source_revision=f"revision-{token}",
        content_sha256=token * 64,
        media_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            if role == "ecrf"
            else "application/pdf"
        ),
        parser_name="document-parser",
        parser_version="v1",
        validation_id=f"validation-{token}",
        validation_revision=1,
        validator_version="source-content-consistency-v2",
        validation_context_sha256="8" * 64,
        locator_index_sha256=("f" if token != "f" else "e") * 64,
        locator_count=2,
        locator_samples=(f"{role}:page:1", f"{role}:page:2"),
        selection_basis="registry_current",
    )


def _packet(
    *,
    ecrf_status: str = "current",
) -> MonitoringDocumentEvidencePacket:
    records = []
    for role, token in zip(DOCUMENT_ROLES, "abcd"):
        if role in {"investigator_brochure", "sap"}:
            records.append(
                DocumentRoleEvidence(
                    role=role,
                    status="missing",
                    limitation_codes=(f"{role}_not_registered",),
                )
            )
        elif role == "ecrf" and ecrf_status != "current":
            records.append(
                DocumentRoleEvidence(
                    role=role,
                    status=ecrf_status,
                    limitation_codes=("ecrf_not_current",),
                )
            )
        else:
            records.append(
                DocumentRoleEvidence(
                    role=role,
                    status="current",
                    binding=_binding(role, token),
                )
            )
    return MonitoringDocumentEvidencePacket(
        project_id=PROJECT_ID,
        listing_admission_date="2026-09-03",
        registry_revision_sha256="9" * 64,
        roles=tuple(records),
    )


def test_packet_accounts_for_all_roles_without_forcing_optional_user_work() -> None:
    packet = _packet()
    payload = packet.to_dict()

    assert packet.mapping_context_ready is True
    assert [item["role"] for item in payload["roles"]] == list(DOCUMENT_ROLES)
    assert payload["roles"][1]["status"] == "missing"
    assert payload["roles"][3]["status"] == "missing"
    assert "no_ctcae_grade_risk_query_or_clinical_join" in (
        payload["clinical_boundary"]
    )
    assert MonitoringDocumentEvidencePacket.from_dict(payload).to_dict() == payload


def test_protocol_and_ecrf_are_required_for_mapping_context() -> None:
    packet = _packet(ecrf_status="incomplete")

    assert packet.mapping_context_ready is False
    with pytest.raises(
        DocumentEvidenceError,
        match="mapping_document_evidence_incomplete",
    ):
        validate_document_evidence_packet(
            packet.to_dict(),
            project_id=PROJECT_ID,
            require_mapping_context=True,
        )


def test_project_and_packet_digest_are_fail_closed() -> None:
    payload = _packet().to_dict()
    with pytest.raises(
        DocumentEvidenceError,
        match="document_packet_project_mismatch",
    ):
        validate_document_evidence_packet(
            payload,
            project_id="another-project",
            require_mapping_context=True,
        )

    tampered = deepcopy(payload)
    tampered["roles"][0]["binding"]["locator_samples"].append(
        "protocol:page:3"
    )
    tampered["roles"][0]["binding"]["locator_count"] = 3
    with pytest.raises(
        DocumentEvidenceError,
        match="document_packet_sha256_mismatch",
    ):
        MonitoringDocumentEvidencePacket.from_dict(tampered)


def test_registry_current_protocol_cannot_claim_event_applicability() -> None:
    with pytest.raises(
        DocumentEvidenceError,
        match="clinical_applicability_unproven",
    ):
        CurrentDocumentBinding(
            **{
                **_binding("protocol", "a").__dict__,
                "clinical_applicability_resolved": True,
            }
        )


def test_current_binding_requires_complete_unique_locators() -> None:
    with pytest.raises(
        DocumentEvidenceError,
        match="document_locators_duplicated",
    ):
        CurrentDocumentBinding(
            **{
                **_binding("ecrf", "c").__dict__,
                "locator_samples": (
                    "ecrf:sheet:1",
                    "ecrf:sheet:1",
                ),
            }
        )


def test_non_current_role_cannot_smuggle_source_identity() -> None:
    with pytest.raises(
        DocumentEvidenceError,
        match="limited_document_role_invalid",
    ):
        DocumentRoleEvidence(
            role="sap",
            status="missing",
            binding=_binding("sap", "d"),
            limitation_codes=("sap_not_registered",),
        )


def _supplementary(
    role: str,
    token: str,
    main: CurrentDocumentBinding,
) -> CurrentDocumentBinding:
    return replace(
        _binding(role, token),
        main_source_entry_id=main.source_entry_id,
    )


def _composite_packet() -> MonitoringDocumentEvidencePacket:
    main = _binding("protocol", "a")
    roles = list(_packet().roles)
    roles[0] = DocumentRoleEvidence(
        role="protocol",
        status="current",
        binding=main,
        supplementary_bindings=(_supplementary("protocol", "e", main),),
    )
    return MonitoringDocumentEvidencePacket(
        project_id=PROJECT_ID,
        listing_admission_date="2026-09-03",
        registry_revision_sha256="9" * 64,
        roles=tuple(roles),
    )


def test_packet_round_trips_main_and_supplementary_bindings() -> None:
    packet = _composite_packet()
    payload = packet.to_dict()
    protocol = payload["roles"][0]

    assert protocol["binding"]["main_source_entry_id"] == ""
    assert protocol["supplementary_bindings"][0]["source_entry_id"] == (
        "source-e"
    )
    assert protocol["supplementary_bindings"][0]["main_source_entry_id"] == (
        "source-a"
    )
    restored = MonitoringDocumentEvidencePacket.from_dict(payload)

    assert restored.to_dict() == payload
    assert restored.roles[0].supplementary_bindings == (
        packet.roles[0].supplementary_bindings
    )
    assert restored.mapping_context_ready is True


def test_supplementary_bindings_fail_closed_on_broken_relations() -> None:
    main = _binding("protocol", "a")

    def evidence_with(
        supplementary: CurrentDocumentBinding,
    ) -> DocumentRoleEvidence:
        return DocumentRoleEvidence(
            role="protocol",
            status="current",
            binding=main,
            supplementary_bindings=(supplementary,),
        )

    with pytest.raises(
        DocumentEvidenceError,
        match="document_supplementary_role_invalid",
    ):
        evidence_with(_binding("sap", "e"))
    with pytest.raises(
        DocumentEvidenceError,
        match="document_supplementary_relation_invalid",
    ):
        evidence_with(
            replace(_binding("protocol", "e"), main_source_entry_id="source-z")
        )
    with pytest.raises(
        DocumentEvidenceError,
        match="document_supplementary_relation_invalid",
    ):
        evidence_with(_binding("protocol", "e"))
    with pytest.raises(
        DocumentEvidenceError,
        match="document_supplementary_duplicated",
    ):
        DocumentRoleEvidence(
            role="protocol",
            status="current",
            binding=main,
            supplementary_bindings=(
                _supplementary("protocol", "e", main),
                _supplementary("protocol", "e", main),
            ),
        )
    with pytest.raises(
        DocumentEvidenceError,
        match="document_supplementary_relation_invalid",
    ):
        CurrentDocumentBinding(
            **{
                **_binding("protocol", "e").__dict__,
                "main_source_entry_id": "source-e",
            }
        )
    with pytest.raises(
        DocumentEvidenceError,
        match="document_supplementary_relation_invalid",
    ):
        DocumentRoleEvidence(
            role="protocol",
            status="current",
            binding=replace(main, main_source_entry_id="source-b"),
        )
    payload = _composite_packet().to_dict()
    payload["roles"][0]["supplementary_bindings"] = "erratum"
    with pytest.raises(
        DocumentEvidenceError,
        match="document_supplementary_binding_invalid",
    ):
        MonitoringDocumentEvidencePacket.from_dict(payload)


def test_non_current_role_rejects_supplementary_bindings() -> None:
    main = _binding("sap", "d")
    with pytest.raises(
        DocumentEvidenceError,
        match="limited_document_role_invalid",
    ):
        DocumentRoleEvidence(
            role="sap",
            status="missing",
            supplementary_bindings=(_supplementary("sap", "e", main),),
            limitation_codes=("sap_not_registered",),
        )


def test_legacy_v1_document_evidence_snapshot_is_rejected() -> None:
    payload = _packet().to_dict()
    payload["schema_version"] = "mm-c3-document-evidence-v1"
    with pytest.raises(
        DocumentEvidenceError,
        match="document_schema_version_invalid",
    ):
        MonitoringDocumentEvidencePacket.from_dict(payload)


def _admission_record(tmp_path: Path) -> tuple[dict, Path]:
    openpyxl = pytest.importorskip("openpyxl")
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "不良事件"
    sheet.append(["受试者编号(SUBJID)", "事件名称(AETERM)"])
    sheet.append(["S001", "头痛"])
    stream = io.BytesIO()
    workbook.save(stream)
    source = tmp_path / "source"
    source.mkdir()
    (source / "listing.xlsx").write_bytes(stream.getvalue())
    workspace = tmp_path / "workspace"
    result = DataAdmissionPipeline(parse_listing_file).create_attempt(
        project_id=PROJECT_ID,
        source_dir=source,
        workspace_dir=workspace,
    )
    runtime = workspace / RUNTIME_DIR_NAME
    store = Store(runtime / RUNTIME_DB_NAME, runtime / ARTIFACT_DIR_NAME)
    try:
        persisted = store.get_domain_object(
            "data_admission", result["attempt_id"]
        )
    finally:
        store.close()
    assert persisted is not None
    return persisted[1], workspace


def test_mapping_bridge_requires_bound_current_documents_when_enabled(
    tmp_path: Path,
) -> None:
    record, _ = _admission_record(tmp_path)
    with pytest.raises(
        MappingBridgeError,
        match="mapping_document_evidence_missing",
    ):
        admission_record_to_harness_input(
            project_id=PROJECT_ID,
            attempt_id=record["attempt_id"],
            record=record,
            require_document_evidence=True,
        )

    bound = deepcopy(record)
    bound["technical_details"]["monitoring_document_evidence"] = (
        _packet().to_dict()
    )
    bridged = admission_record_to_harness_input(
        project_id=PROJECT_ID,
        attempt_id=record["attempt_id"],
        record=bound,
        require_document_evidence=True,
    )
    assert bridged.field_profile["document_evidence"]["packet_sha256"] == (
        _packet().packet_sha256
    )
    assert bridged.field_profile["bridge_schema_version"].endswith("-v8")


def test_missing_required_document_creates_no_mapping_job(
    tmp_path: Path,
) -> None:
    record, workspace = _admission_record(tmp_path)
    repository = MonitoringAiRepository(tmp_path / "missing-doc-ai.sqlite3")

    def runtime(provider: str, model: str, profile_id: str):
        return MonitoringAiRuntimeBinding(
            profile_id=profile_id,
            provider=provider,
            model=model,
            env={},
            available=True,
        )

    pipeline = AdmissionMappingPipeline(
        ai_service=MonitoringAiService(
            repository,
            runtime_resolver=lambda: runtime(
                MONITORING_C3_MAPPING_PROVIDER,
                MONITORING_C3_MAPPING_MODEL,
                MONITORING_C3_MAPPING_PROFILE_ID,
            ),
        ),
        verifier_ai_service=MonitoringAiService(
            repository,
            runtime_resolver=lambda: runtime(
                MONITORING_C3_VERIFIER_PROVIDER,
                MONITORING_C3_VERIFIER_MODEL,
                "independent_ai__zhipu_coding_plan_glm_flash",
            ),
        ),
        ai_repository=repository,
        input_revision_factory=MonitoringAiInputRevision.model_validate,
        task_type=MonitoringAiTaskType.LISTING_FIELD_MAPPING,
        relationship_profiler=build_relationship_profile,
        require_document_evidence=True,
        document_evidence_resolver=lambda **_kwargs: _packet(
            ecrf_status="missing"
        ),
    )

    with pytest.raises(
        AdmissionMappingPipelineError,
        match="mapping_document_evidence_incomplete",
    ):
        pipeline.generate_dual_candidates(
            project_id=PROJECT_ID,
            attempt_id=record["attempt_id"],
            workspace_dir=workspace,
        )
    assert repository.list_jobs(PROJECT_ID) == ()


def test_mapping_pipeline_freezes_resolved_documents_at_attempt_date(
    tmp_path: Path,
) -> None:
    record, workspace = _admission_record(tmp_path)
    calls = []

    def resolve_documents(
        *,
        project_id: str,
        listing_admission_date: str,
        selected_entry_ids: dict[str, str],
    ):
        assert selected_entry_ids == {}
        calls.append((project_id, listing_admission_date))
        return _packet()

    bridged = AdmissionMappingPipeline(
        require_document_evidence=True,
        document_evidence_resolver=resolve_documents,
    )._frozen_harness_input(
        project_id=PROJECT_ID,
        attempt_id=record["attempt_id"],
        record=record,
        workspace_dir=workspace,
    )

    assert calls == [(PROJECT_ID, calls[0][1])]
    assert len(calls[0][1]) == 10
    assert bridged.field_profile["document_evidence"]["packet_sha256"] == (
        _packet().packet_sha256
    )


def test_dual_submission_resolves_one_shared_document_snapshot(
    tmp_path: Path,
) -> None:
    record, workspace = _admission_record(tmp_path)
    repository = MonitoringAiRepository(tmp_path / "monitoring-ai.sqlite3")
    calls = []
    wakes = []

    def resolve_documents(**kwargs):
        calls.append(kwargs)
        return _packet()

    def runtime(provider: str, model: str, profile_id: str):
        return MonitoringAiRuntimeBinding(
            profile_id=profile_id,
            provider=provider,
            model=model,
            env={},
            available=True,
        )

    pipeline = AdmissionMappingPipeline(
        ai_service=MonitoringAiService(
            repository,
            runtime_resolver=lambda: runtime(
                MONITORING_C3_MAPPING_PROVIDER,
                MONITORING_C3_MAPPING_MODEL,
                MONITORING_C3_MAPPING_PROFILE_ID,
            ),
        ),
        verifier_ai_service=MonitoringAiService(
            repository,
            runtime_resolver=lambda: runtime(
                MONITORING_C3_VERIFIER_PROVIDER,
                MONITORING_C3_VERIFIER_MODEL,
                "independent_ai__zhipu_coding_plan_glm_flash",
            ),
        ),
        ai_repository=repository,
        input_revision_factory=MonitoringAiInputRevision.model_validate,
        task_type=MonitoringAiTaskType.LISTING_FIELD_MAPPING,
        relationship_profiler=build_relationship_profile,
        require_document_evidence=True,
        document_evidence_resolver=resolve_documents,
        worker_wake=lambda: wakes.append(True),
    )

    result = pipeline.generate_dual_candidates(
        project_id=PROJECT_ID,
        attempt_id=record["attempt_id"],
        workspace_dir=workspace,
    )
    jobs = repository.list_jobs(PROJECT_ID)
    profiles = [
        repository.input_payload(PROJECT_ID, job.job_id)["field_profile"]
        for job in jobs
    ]

    assert result["verification"]["state"] == "generating"
    assert len(calls) == 1
    assert wakes == [True]
    assert len({
        profile["document_evidence"]["packet_sha256"]
        for profile in profiles
    }) == 1
    assert len({profile["full_input_sha256"] for profile in profiles}) == 1
    assert pipeline._document_selection_ids(
        workspace,
        record["attempt_id"],
    ) == {
        "protocol": "source-a",
        "ecrf": "source-c",
    }


def test_document_selection_is_validated_before_it_is_versioned(
    tmp_path: Path,
) -> None:
    record, workspace = _admission_record(tmp_path)

    def resolve_documents(**kwargs):
        selected = kwargs["selected_entry_ids"].get("ecrf", "")
        if selected == "source-c":
            return _packet()
        roles = list(_packet().roles)
        roles[2] = DocumentRoleEvidence(
            role="ecrf",
            status="incomplete",
            limitation_codes=("ecrf_selection_stale",),
        )
        return MonitoringDocumentEvidencePacket(
            project_id=PROJECT_ID,
            listing_admission_date="2026-09-03",
            registry_revision_sha256="9" * 64,
            roles=tuple(roles),
        )

    pipeline = AdmissionMappingPipeline(
        require_document_evidence=True,
        document_evidence_resolver=resolve_documents,
    )
    with pytest.raises(
        AdmissionMappingPipelineError,
        match="mapping_document_selection_not_found_or_role_mismatch",
    ):
        pipeline.select_document(
            project_id=PROJECT_ID,
            attempt_id=record["attempt_id"],
            workspace_dir=workspace,
            role="ecrf",
            source_entry_id="source-wrong-role",
        )
    assert pipeline._document_selection_ids(workspace, record["attempt_id"]) == {}

    saved = pipeline.select_document(
        project_id=PROJECT_ID,
        attempt_id=record["attempt_id"],
        workspace_dir=workspace,
        role="ecrf",
        source_entry_id="source-c",
    )
    assert saved["version"] == 1
    assert pipeline._document_selection_ids(
        workspace,
        record["attempt_id"],
    ) == {
        "ecrf": "source-c"
    }


def test_document_selection_is_isolated_between_admission_attempts(
    tmp_path: Path,
) -> None:
    first, workspace = _admission_record(tmp_path)
    second = DataAdmissionPipeline(parse_listing_file).create_attempt(
        project_id=PROJECT_ID,
        source_dir=tmp_path / "source",
        workspace_dir=workspace,
    )

    def resolve_documents(**kwargs):
        selected = kwargs["selected_entry_ids"].get("ecrf", "")
        token = {"source-c": "c", "source-d": "d"}.get(selected, "c")
        roles = list(_packet().roles)
        roles[2] = DocumentRoleEvidence(
            role="ecrf",
            status="current",
            binding=_binding("ecrf", token),
        )
        return MonitoringDocumentEvidencePacket(
            project_id=PROJECT_ID,
            listing_admission_date="2026-09-03",
            registry_revision_sha256="9" * 64,
            roles=tuple(roles),
        )

    pipeline = AdmissionMappingPipeline(
        require_document_evidence=True,
        document_evidence_resolver=resolve_documents,
    )
    pipeline.select_document(
        project_id=PROJECT_ID,
        attempt_id=first["attempt_id"],
        workspace_dir=workspace,
        role="ecrf",
        source_entry_id="source-c",
    )
    pipeline.select_document(
        project_id=PROJECT_ID,
        attempt_id=second["attempt_id"],
        workspace_dir=workspace,
        role="ecrf",
        source_entry_id="source-d",
    )

    assert pipeline._document_selection_ids(
        workspace,
        first["attempt_id"],
    ) == {"ecrf": "source-c"}
    assert pipeline._document_selection_ids(
        workspace,
        second["attempt_id"],
    ) == {"ecrf": "source-d"}


def test_document_readiness_is_plain_chinese_and_hides_registry_identity(
    tmp_path: Path,
) -> None:
    record, workspace = _admission_record(tmp_path)
    pipeline = AdmissionMappingPipeline(
        require_document_evidence=True,
        document_evidence_resolver=lambda **_kwargs: _packet(),
    )

    readiness = pipeline.document_readiness(
        project_id=PROJECT_ID,
        attempt_id=record["attempt_id"],
        workspace_dir=workspace,
    )
    blob = json.dumps(readiness, ensure_ascii=False)

    assert readiness["ready"] is True
    assert "无需逐项确认" in readiness["guidance"]
    assert [item["label"] for item in readiness["roles"]] == [
        "研究方案",
        "研究者手册",
        "电子病例报告表",
        "统计分析计划",
    ]
    assert "source_entry_id" not in blob
    assert "packet_sha256" not in blob


class _Registry:
    def __init__(self, entries: list, spans: list) -> None:
        self._entries = entries
        self._spans = spans

    def list_entries(self, project_id: str) -> list:
        return [
            item for item in self._entries if item.project_id == project_id
        ]

    def list_spans(self, project_id: str) -> list:
        return [
            item for item in self._spans if item.project_id == project_id
        ]

    def current_content_validation(
        self, project_id: str, source_entry_id: str
    ):
        return SimpleNamespace(
            validation_id=f"validation-{source_entry_id}",
            revision=1,
            expected_context_hash="8" * 64,
            use_status="allowed",
            technical_status="ready",
            validator_version="source_content_consistency_v2",
            source_entry_id=source_entry_id,
            project_id=project_id,
            module="medical_monitoring",
            file_sha256=next(
                item.content_hash
                for item in self._entries
                if item.entry_id == source_entry_id
            ),
        )

    def assert_operational_sources_usable(
        self, project_id: str, source_ids: list[str]
    ) -> None:
        assert project_id == PROJECT_ID
        assert source_ids
        target = next(
            item
            for item in self._entries
            if source_ids[0].startswith(item.entry_id + "-span-")
        )
        latest = max(
            [
                item
                for item in self._entries
                if item.project_id == project_id
                and item.module == "medical_monitoring"
                and item.source_kind == target.source_kind
            ],
            key=lambda item: item.created_at,
        )
        promoted = [
            item
            for item in self._entries
            if item.project_id == project_id
            and item.module == "medical_monitoring"
            and item.source_kind == target.source_kind
            and item.metadata.get("monitoring_authority_status") == "promoted"
        ]
        if promoted:
            latest = max(promoted, key=lambda item: item.created_at)
        if latest.entry_id != target.entry_id:
            raise ValueError("superseded")


def _entry(
    role: str,
    order: int,
    *,
    module: str = "medical_monitoring",
    authority_status: str = "current_effective",
):
    source_kind = {
        "protocol": "protocol_docx",
        "investigator_brochure": "investigator_brochure",
        "ecrf": "ecrf_xlsx",
        "sap": "statistical_analysis_plan",
    }[role]
    suffix = ".docx" if role == "protocol" else (
        ".xlsx" if role == "ecrf" else ".pdf"
    )
    entry_id = f"source-{role}-{order}"
    locator_manifest = [
        {
            "source_id": f"{entry_id}-span-{index}",
            "locator": f"document:page:{index}",
        }
        for index in (1, 2)
    ]
    return SimpleNamespace(
        entry_id=entry_id,
        project_id=PROJECT_ID,
        module=module,
        source_kind=source_kind,
        public_title=f"{role}{suffix}",
        content_hash=str(order) * 64,
        parser_status="parsed",
        parser_version="parser-v1",
        span_count=2,
        metadata={
            "filename": f"{role}{suffix}",
            "monitoring_authority_status": authority_status,
            "document_revision": f"{role}-revision-{order}",
            "expected_locator_count": 2,
            "expected_locator_index_sha256": hashlib.sha256(
                json.dumps(
                    locator_manifest,
                    ensure_ascii=False,
                    sort_keys=True,
                    separators=(",", ":"),
                ).encode("utf-8")
            ).hexdigest(),
            "locator_manifest_complete": True,
        },
        created_at=datetime(
            2026, 9, order, tzinfo=timezone.utc
        ),
    )


def _span(entry, index: int):
    return SimpleNamespace(
        source_id=f"{entry.entry_id}-span-{index}",
        entry_id=entry.entry_id,
        project_id=entry.project_id,
        module=entry.module,
        locator=f"document:page:{index}",
        text_preview=f"bounded extract {index}",
    )


def _promote_entries(*entries) -> None:
    registrations = [
        {
            "role": {
                "protocol_docx": "protocol",
                "investigator_brochure": "investigator_brochure",
                "ecrf_xlsx": "ecrf",
                "statistical_analysis_plan": "sap",
            }[entry.source_kind],
            "candidate_id": f"candidate-{entry.entry_id}",
            "source_entry_id": entry.entry_id,
            "content_sha256": entry.content_hash,
        }
        for entry in entries
    ]
    receipt = {
        "schema_version": "monitoring-document-authority-promotion-v1",
        "batch_id": f"mmbatch_{'b' * 24}",
        "input_sha256": "a" * 64,
        "analysis_job_ids": ["primary-job", "verifier-job"],
        "review_job_ids": [],
        "analysis_run_ids": ["primary-run", "verifier-run"],
        "review_run_ids": [],
        "document_identities": [
            {
                "role": item["role"],
                "candidate_id": item["candidate_id"],
                "document_version": "",
                "document_date": "",
            }
            for item in registrations
        ],
        "registrations": registrations,
    }
    receipt_sha256 = content_hash(receipt)
    for entry in entries:
        entry.metadata.update({
            "monitoring_authority_status": "promoted",
            "document_authority_receipt_sha256": receipt_sha256,
            "document_authority_receipt": receipt,
        })


def test_registry_resolver_rejects_incomplete_promoted_authority_set() -> None:
    protocol = _entry("protocol", 1)
    ecrf = _entry("ecrf", 2)
    _promote_entries(protocol, ecrf)
    registry = _Registry(
        [protocol],
        [_span(protocol, 1), _span(protocol, 2)],
    )

    packet = MonitoringDocumentEvidenceResolver(registry).resolve(
        project_id=PROJECT_ID,
    )

    assert packet.roles[0].status == "missing"


def test_registry_resolver_rejects_tampered_promotion_receipt() -> None:
    protocol = _entry("protocol", 1)
    _promote_entries(protocol)
    protocol.metadata["document_authority_receipt"]["batch_id"] = "tampered"
    registry = _Registry(
        [protocol],
        [_span(protocol, 1), _span(protocol, 2)],
    )

    packet = MonitoringDocumentEvidenceResolver(registry).resolve(
        project_id=PROJECT_ID,
    )

    assert packet.roles[0].status == "missing"


def test_registry_resolver_rejects_self_hashed_receipt_missing_core_evidence() -> None:
    protocol = _entry("protocol", 1)
    _promote_entries(protocol)
    receipt = protocol.metadata["document_authority_receipt"]
    receipt.pop("analysis_job_ids")
    protocol.metadata["document_authority_receipt_sha256"] = content_hash(receipt)
    registry = _Registry(
        [protocol],
        [_span(protocol, 1), _span(protocol, 2)],
    )

    packet = MonitoringDocumentEvidenceResolver(registry).resolve(
        project_id=PROJECT_ID,
    )

    assert packet.roles[0].status == "missing"


def test_registry_resolver_accepts_complete_promoted_authority_set() -> None:
    protocol = _entry("protocol", 1)
    ecrf = _entry("ecrf", 2)
    _promote_entries(protocol, ecrf)
    entries = [protocol, ecrf]
    registry = _Registry(
        entries,
        [_span(entry, index) for entry in entries for index in (1, 2)],
    )

    packet = MonitoringDocumentEvidenceResolver(registry).resolve(
        project_id=PROJECT_ID,
    )

    by_role = {item.role: item for item in packet.roles}
    assert by_role["protocol"].status == "current"
    assert by_role["ecrf"].status == "current"


def test_promoted_authority_is_not_displaced_by_newer_unpromoted_upload() -> None:
    promoted = _entry("protocol", 1)
    _promote_entries(promoted)
    unpromoted = _entry("protocol", 2)
    entries = [promoted, unpromoted]
    registry = _Registry(
        entries,
        [_span(entry, index) for entry in entries for index in (1, 2)],
    )

    packet = MonitoringDocumentEvidenceResolver(registry).resolve(
        project_id=PROJECT_ID,
    )

    protocol = packet.roles[0]
    assert protocol.status == "current"
    assert protocol.binding is not None
    assert protocol.binding.source_entry_id == promoted.entry_id


def test_registry_resolver_uses_only_monitoring_owned_current_sources() -> None:
    current = [
        _entry("protocol", 2),
        _entry("ecrf", 3),
        _entry("investigator_brochure", 4),
        _entry("sap", 5),
    ]
    old_protocol = _entry(
        "protocol",
        1,
        authority_status="superseded",
    )
    writing_ib = _entry(
        "investigator_brochure",
        6,
        module="medical_writing",
    )
    entries = [old_protocol, *current, writing_ib]
    spans = [
        _span(entry, index)
        for entry in entries
        for index in (1, 2)
    ]
    packet = MonitoringDocumentEvidenceResolver(
        _Registry(entries, spans)
    ).resolve(
        project_id=PROJECT_ID,
        listing_admission_date="2026-09-03",
    )
    by_role = {item.role: item for item in packet.roles}

    assert packet.mapping_context_ready is True
    assert by_role["protocol"].binding is not None
    assert by_role["protocol"].binding.source_entry_id == (
        "source-protocol-2"
    )
    assert by_role["investigator_brochure"].binding is not None
    assert by_role[
        "investigator_brochure"
    ].binding.source_entry_id == "source-investigator_brochure-4"
    assert all(
        item.binding is None
        or item.binding.clinical_applicability_resolved is False
        for item in packet.roles
    )


def test_registry_resolver_reuses_registry_supersession_policy() -> None:
    first = _entry("protocol", 1)
    second = _entry("protocol", 2)
    entries = [first, second]
    packet = MonitoringDocumentEvidenceResolver(
        _Registry(
            entries,
            [
                _span(first, 1),
                _span(first, 2),
                _span(second, 1),
                _span(second, 2),
            ],
        )
    ).resolve(
        project_id=PROJECT_ID,
        listing_admission_date="2026-09-03",
    )

    protocol = packet.roles[0]
    assert protocol.status == "current"
    assert protocol.binding is not None
    assert protocol.binding.source_entry_id == second.entry_id


def test_registry_resolver_rejects_partial_locator_index() -> None:
    protocol = _entry("protocol", 1)
    first = _span(protocol, 1)
    second = _span(protocol, 2)
    second.locator = ""
    packet = MonitoringDocumentEvidenceResolver(
        _Registry([protocol], [first, second])
    ).resolve(
        project_id=PROJECT_ID,
        listing_admission_date="2026-09-03",
    )

    assert packet.roles[0].status == "incomplete"


def test_registry_resolver_rejects_self_consistent_partial_span_count() -> None:
    protocol = _entry("protocol", 1)
    protocol.span_count = 1
    packet = MonitoringDocumentEvidenceResolver(
        _Registry([protocol], [_span(protocol, 1)])
    ).resolve(
        project_id=PROJECT_ID,
        listing_admission_date="2026-09-03",
    )

    assert packet.roles[0].status == "incomplete"


def test_ecrf_registration_persists_independent_locator_manifest(
    tmp_path: Path,
) -> None:
    openpyxl = pytest.importorskip("openpyxl")
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "AE表单"
    sheet.append(["字段名", "字段含义"])
    sheet.append(["AETERM", "不良事件名称"])
    stream = io.BytesIO()
    workbook.save(stream)
    registry = SourceRegistryService(
        SourceRegistryStore(tmp_path / "registry.jsonl")
    )

    result = registry.register_monitoring_mapping_document(
        PROJECT_ID,
        "forms.xlsx",
        stream.getvalue(),
        document_role="ecrf",
    )

    metadata = result.entry.metadata
    legacy_token = hashlib.sha256(
        (
            "workbench_source_registry_v0_1:"
            f"{PROJECT_ID}:medical_monitoring:ecrf:"
            f"{result.entry.content_hash}"
        ).encode("utf-8")
    ).hexdigest()[:12]
    legacy_entry_id = f"src_{PROJECT_ID.replace('-', '_')}_ecrf_{legacy_token}"
    assert result.entry.source_kind == "ecrf"
    assert result.entry.entry_id != legacy_entry_id
    assert metadata["locator_manifest_complete"] is True
    assert metadata["locator_manifest_revision"] == (
        "monitoring-locator-manifest-v1"
    )
    assert metadata["expected_locator_count"] == len(result.spans) == 1
    assert len(metadata["expected_locator_index_sha256"]) == 64

    supplementary = registry.register_monitoring_mapping_document(
        PROJECT_ID,
        "forms-supplement.xlsx",
        stream.getvalue(),
        document_role="ecrf",
        document_relation="supplementary",
    )
    assert supplementary.entry.source_kind == "ecrf_supplement"


def test_scanned_monitoring_reference_uses_complete_verified_ocr_spans(
    tmp_path: Path,
) -> None:
    pymupdf = pytest.importorskip("pymupdf")
    document = pymupdf.open()
    document.new_page()
    document.new_page()
    payload = document.tobytes()
    document.close()
    rows = [
        {
            "locator": "candidate:scan:p1:ocr",
            "text": "方案勘误第一页",
            "text_sha256": hashlib.sha256("方案勘误第一页".encode()).hexdigest(),
        },
        {
            "locator": "candidate:scan:p2:ocr",
            "text": "方案勘误第二页",
            "text_sha256": hashlib.sha256("方案勘误第二页".encode()).hexdigest(),
        },
    ]
    registry = SourceRegistryService(
        SourceRegistryStore(tmp_path / "registry.jsonl")
    )

    result = registry.register_monitoring_mapping_document(
        PROJECT_ID,
        "stamped-erratum.pdf",
        payload,
        document_role="ecrf",
        verified_text_spans=rows,
        expected_locator_count=2,
    )

    assert [span.locator for span in result.spans] == [
        "candidate:scan:p1:ocr",
        "candidate:scan:p2:ocr",
    ]
    assert result.entry.metadata["parser_name"] == "monitoring_candidate_ocr"
    assert result.entry.metadata["locator_manifest_complete"] is True
    assert result.entry.source_kind == "ecrf_document"

    with pytest.raises(ValueError, match="complete text locator set"):
        registry.register_monitoring_mapping_document(
            PROJECT_ID,
            "incomplete-scan.pdf",
            payload,
            document_role="protocol",
            document_relation="supplementary",
            verified_text_spans=rows[:1],
            expected_locator_count=2,
        )


@pytest.mark.parametrize(
    ("role", "filename", "paragraphs"),
    (
        (
            "investigator_brochure",
            "investigator-brochure.docx",
            (
                "Investigator Brochure",
                "Nonclinical Studies",
                "Effects in Humans",
            ),
        ),
        (
            "sap",
            "statistical-analysis-plan.pdf",
            (
                "Statistical Analysis Plan",
                "Analysis Population",
                "Statistical Methods",
            ),
        ),
    ),
)
def test_optional_study_document_registration_is_locator_backed_and_current(
    tmp_path: Path,
    role: str,
    filename: str,
    paragraphs: tuple[str, ...],
) -> None:
    expected_context = SourceExpectedContext(
        project_identifiers=(PROJECT_ID,),
        indication_terms=("unrelated-indication-label",),
        expected_protocol_version="protocol-version-not-required-here",
    )
    validation_store = SourceContentValidationStore(
        tmp_path / "validations.sqlite3"
    )
    validation_service = SourceContentValidationService(validation_store)
    expected_context_holder = [expected_context]
    registry = SourceRegistryService(
        SourceRegistryStore(tmp_path / "registry.jsonl"),
        artifact_root=tmp_path / "artifacts",
        content_validation_service=validation_service,
        expected_context_resolver=lambda *_args: (
            expected_context_holder[0]
        ),
    )
    document_paragraphs = (PROJECT_ID, *paragraphs)
    if filename.endswith(".docx"):
        document = docx.Document()
        for paragraph in document_paragraphs:
            document.add_paragraph(paragraph)
        stream = io.BytesIO()
        document.save(stream)
        payload = stream.getvalue()
    else:
        pymupdf = pytest.importorskip("pymupdf")
        document = pymupdf.open()
        page = document.new_page()
        page.insert_text((72, 72), "\n".join(document_paragraphs))
        payload = document.tobytes()
        document.close()

    result = registry.register_monitoring_mapping_document(
        PROJECT_ID,
        filename,
        payload,
        document_role=role,
    )
    validation = registry.current_content_validation(
        PROJECT_ID,
        result.entry.entry_id,
    )
    packet = MonitoringDocumentEvidenceResolver(registry).resolve(
        project_id=PROJECT_ID,
        listing_admission_date="2026-09-03",
    )
    evidence = next(item for item in packet.roles if item.role == role)

    assert result.entry.module == "medical_monitoring"
    assert result.entry.source_kind == {
        "investigator_brochure": "investigator_brochure",
        "sap": "statistical_analysis_plan",
    }[role]
    assert result.entry.metadata["locator_manifest_complete"] is True
    assert result.entry.metadata["expected_locator_count"] == len(
        result.spans
    )
    assert result.spans
    assert validation is not None
    assert validation.content_status == "matched"
    assert validation.use_status == "allowed"
    assert evidence.status == "current"
    assert evidence.binding is not None
    assert evidence.binding.locator_count == len(result.spans)
    assert packet.mapping_context_ready is False
    retrieval = MonitoringDocumentEvidenceResolver(
        registry
    ).retrieve_current_excerpts(
        project_id=PROJECT_ID,
        binding=evidence.binding,
        query_terms=(
            ["Nonclinical"]
            if role == "investigator_brochure"
            else ["Analysis Population"]
        ),
    )
    assert retrieval["source_entry_id"] == result.entry.entry_id
    assert retrieval["content_sha256"] == result.entry.content_hash
    assert retrieval["locator_index_sha256"] == (
        evidence.binding.locator_index_sha256
    )
    assert retrieval["excerpts"]
    assert retrieval["clinical_conclusions"] == []
    assert len(retrieval["packet_sha256"]) == 64
    expected_context_holder[0] = replace(
        expected_context,
        project_identifiers=(PROJECT_ID, "alternate-project-label"),
    )
    validation_service.assess_protocol(
        project_id=PROJECT_ID,
        source_entry_id=result.entry.entry_id,
        module="medical_monitoring",
        filename=filename,
        file_sha256=result.entry.content_hash,
        document=ProtocolTextDocument(
            filename=filename,
            title=filename,
            paragraphs=[],
            tables=[],
            spans=[
                ProtocolTextSpan(
                    span_id=str(index),
                    kind="reference_text",
                    text=text,
                    source_locator=f"document:span:{index}",
                )
                for index, text in enumerate(
                    document_paragraphs,
                    start=1,
                )
            ],
            source_hash=result.entry.content_hash,
        ),
        expected=replace(
            expected_context_holder[0],
            expected_file_role=role,
        ),
        actor="system_validator",
    )
    with pytest.raises(
        ValueError,
        match="binding is no longer current",
    ):
        MonitoringDocumentEvidenceResolver(
            registry
        ).retrieve_current_excerpts(
            project_id=PROJECT_ID,
            binding=evidence.binding,
            query_terms=["Analysis"],
        )
    with pytest.raises(
        ValueError,
        match="binding is no longer current",
    ):
        MonitoringDocumentEvidenceResolver(
            registry
        ).retrieve_current_excerpts(
            project_id=PROJECT_ID,
            binding=replace(
                evidence.binding,
                content_sha256="f" * 64,
            ),
            query_terms=["Analysis"],
        )


def test_unusable_optional_replacement_preserves_last_usable_document(
    tmp_path: Path,
) -> None:
    expected_context = SourceExpectedContext(
        project_identifiers=(PROJECT_ID,),
    )
    registry = SourceRegistryService(
        SourceRegistryStore(tmp_path / "registry.jsonl"),
        content_validation_service=SourceContentValidationService(
            SourceContentValidationStore(tmp_path / "validations.sqlite3")
        ),
        expected_context_resolver=lambda *_args: expected_context,
    )

    def ib_bytes(*values: str) -> bytes:
        document = docx.Document()
        for value in values:
            document.add_paragraph(value)
        stream = io.BytesIO()
        document.save(stream)
        return stream.getvalue()

    usable = registry.register_monitoring_mapping_document(
        PROJECT_ID,
        "usable-ib.docx",
        ib_bytes(
            PROJECT_ID,
            "Investigator Brochure",
            "Nonclinical Studies",
        ),
        document_role="investigator_brochure",
    )
    unusable = registry.register_monitoring_mapping_document(
        PROJECT_ID,
        "unusable-ib.docx",
        ib_bytes(PROJECT_ID, "Administrative cover only"),
        document_role="investigator_brochure",
    )

    packet = MonitoringDocumentEvidenceResolver(registry).resolve(
        project_id=PROJECT_ID,
    )
    evidence = next(
        item
        for item in packet.roles
        if item.role == "investigator_brochure"
    )
    assert registry.current_content_validation(
        PROJECT_ID,
        unusable.entry.entry_id,
    ).use_status == "requires_confirmation"
    assert evidence.status == "current"
    assert evidence.binding is not None
    assert evidence.binding.source_entry_id == usable.entry.entry_id


def test_monitoring_freshness_supersedes_legacy_role_alias(
    tmp_path: Path,
) -> None:
    openpyxl = pytest.importorskip("openpyxl")
    expected_context = SourceExpectedContext()
    validation_service = SourceContentValidationService(
        SourceContentValidationStore(tmp_path / "validations.sqlite3")
    )
    registry = SourceRegistryService(
        SourceRegistryStore(tmp_path / "registry.jsonl"),
        content_validation_service=validation_service,
        expected_context_resolver=lambda *_args: expected_context,
    )

    def ecrf_bytes(field_name: str) -> bytes:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Form"
        sheet.append(["字段名", "字段含义"])
        sheet.append([field_name, "临床数据字段"])
        stream = io.BytesIO()
        workbook.save(stream)
        workbook.close()
        return stream.getvalue()

    first_payload = ecrf_bytes("FIRST_FIELD")
    first = registry.register_monitoring_mapping_document(
        PROJECT_ID,
        "first-ecrf.xlsx",
        first_payload,
        document_role="ecrf",
    )
    legacy_entry_id = first.entry.entry_id + "_legacy"
    legacy_spans = [
        span.model_copy(
            update={
                "source_id": span.source_id + "_legacy",
                "entry_id": legacy_entry_id,
            }
        )
        for span in first.spans
    ]
    legacy = first.model_copy(
        update={
            "entry": first.entry.model_copy(
                update={
                    "entry_id": legacy_entry_id,
                    "source_kind": "ecrf_xlsx",
                }
            ),
            "spans": legacy_spans,
        }
    )
    registry.store.append(legacy)
    validation_service.assess_listing(
        project_id=PROJECT_ID,
        source_entry_id=legacy_entry_id,
        module="medical_monitoring",
        filename="first-ecrf.xlsx",
        file_sha256=first.entry.content_hash,
        sheets=parse_listing_file("first-ecrf.xlsx", first_payload),
        expected=replace(expected_context, expected_file_role="ecrf"),
        actor="system_validator",
    )
    registry.register_monitoring_mapping_document(
        PROJECT_ID,
        "second-ecrf.xlsx",
        ecrf_bytes("SECOND_FIELD"),
        document_role="ecrf",
    )

    with pytest.raises(ValueError, match="superseded registered source"):
        registry.assert_operational_sources_usable(
            PROJECT_ID,
            [legacy_spans[0].source_id],
        )


class _CompositeRegistry(_Registry):
    """Registry stub for composite-authority resolver tests.

    Receipt verification and relation-aware freshness belong to
    services/api/app/source_intake.py, which does not yet accept composite
    same-role receipts.  This stub treats self-consistent promoted receipts
    as verified and exempts entries sharing one promotion receipt from
    superseding each other, isolating the resolver's composite partition
    logic until the registry layer is extended.
    """

    def monitoring_authority_entry_is_verified(self, entry) -> bool:
        metadata = dict(entry.metadata or {})
        if metadata.get("monitoring_authority_status") != "promoted":
            return True
        receipt = metadata.get("document_authority_receipt")
        return isinstance(receipt, dict) and content_hash(receipt) == (
            metadata.get("document_authority_receipt_sha256")
        )

    def assert_operational_sources_usable(
        self, project_id: str, source_ids: list[str]
    ) -> None:
        assert project_id == PROJECT_ID
        assert source_ids
        target = next(
            item
            for item in self._entries
            if source_ids[0].startswith(item.entry_id + "-span-")
        )
        receipt = dict(target.metadata or {}).get(
            "document_authority_receipt"
        )
        composite_group = (
            {
                str(row.get("source_entry_id") or "")
                for row in receipt.get("registrations", ())
                if isinstance(row, dict)
            }
            if isinstance(receipt, dict)
            else set()
        )
        competitors = [
            item
            for item in self._entries
            if item.project_id == project_id
            and item.module == "medical_monitoring"
            and item.source_kind == target.source_kind
            and (
                item.entry_id == target.entry_id
                or item.entry_id not in composite_group
            )
        ]
        latest = max(
            competitors,
            key=lambda item: (item.created_at, item.entry_id),
        )
        if latest.entry_id != target.entry_id:
            raise ValueError("superseded")

    def search_document_spans(
        self,
        project_id: str,
        source_entry_id: str,
        query_terms: Sequence[str],
        limit: int = 24,
        required_module: str = "",
        allowed_source_kinds: frozenset[str] | None = None,
    ) -> list[dict]:
        return [
            {
                "source_id": span.source_id,
                "locator": span.locator,
                "text": span.text_preview,
                "matched_keywords": [
                    term
                    for term in query_terms
                    if str(term).lower() in span.text_preview.lower()
                ],
            }
            for span in self._spans
            if span.entry_id == source_entry_id
            and span.project_id == project_id
            and span.module == required_module
        ][:limit]


def _promote_composite(main_entry, *supplementary_entries) -> None:
    role_by_kind = {
        "protocol_docx": "protocol",
        "investigator_brochure": "investigator_brochure",
        "ecrf_xlsx": "ecrf",
        "statistical_analysis_plan": "sap",
    }
    role = role_by_kind[main_entry.source_kind]
    registrations = [
        {
            "role": role,
            "candidate_id": f"candidate-{main_entry.entry_id}",
            "source_entry_id": main_entry.entry_id,
            "content_sha256": main_entry.content_hash,
            "binding_kind": "main",
        }
    ]
    for entry in supplementary_entries:
        registrations.append(
            {
                "role": role,
                "candidate_id": f"candidate-{entry.entry_id}",
                "source_entry_id": entry.entry_id,
                "content_sha256": entry.content_hash,
                "binding_kind": "supplementary",
                "supplementary_of": main_entry.entry_id,
            }
        )
    receipt = {
        "schema_version": "monitoring-document-authority-promotion-v1",
        "batch_id": f"mmbatch_{'c' * 24}",
        "input_sha256": "a" * 64,
        "analysis_job_ids": ["primary-job", "verifier-job"],
        "review_job_ids": [],
        "analysis_run_ids": ["primary-run", "verifier-run"],
        "review_run_ids": [],
        "document_identities": [
            {
                "role": row["role"],
                "candidate_id": row["candidate_id"],
                "document_version": "",
                "document_date": "",
            }
            for row in registrations
        ],
        "registrations": registrations,
    }
    receipt_sha256 = content_hash(receipt)
    for entry in (main_entry, *supplementary_entries):
        entry.metadata.update(
            {
                "monitoring_authority_status": "promoted",
                "document_authority_receipt_sha256": receipt_sha256,
                "document_authority_receipt": receipt,
            }
        )


def _composite_scenario_entries():
    main = _entry("protocol", 1)
    erratum = _entry("protocol", 2)
    erratum.metadata["document_revision"] = "protocol-erratum-1"
    _promote_composite(main, erratum)
    spans = [
        _span(main, 1),
        _span(main, 2),
        _span(erratum, 1),
        _span(erratum, 2),
    ]
    return main, erratum, spans


def test_resolver_binds_declared_supplementary_files_to_role_evidence() -> None:
    main, erratum, spans = _composite_scenario_entries()
    registry = _CompositeRegistry([main, erratum], spans)

    packet = MonitoringDocumentEvidenceResolver(registry).resolve(
        project_id=PROJECT_ID,
        listing_admission_date="2026-09-03",
    )
    protocol = packet.roles[0]

    assert protocol.status == "current"
    assert protocol.binding is not None
    assert protocol.binding.source_entry_id == main.entry_id
    assert [
        item.source_entry_id for item in protocol.supplementary_bindings
    ] == [erratum.entry_id]
    assert all(
        item.main_source_entry_id == main.entry_id
        for item in protocol.supplementary_bindings
    )
    restored = MonitoringDocumentEvidencePacket.from_dict(packet.to_dict())

    assert restored.roles[0].supplementary_bindings == (
        protocol.supplementary_bindings
    )


def test_resolver_fails_closed_when_declared_supplementary_is_unresolved() -> None:
    main, erratum, spans = _composite_scenario_entries()
    erratum.metadata["expected_locator_index_sha256"] = "0" * 64
    registry = _CompositeRegistry([main, erratum], spans)

    packet = MonitoringDocumentEvidenceResolver(registry).resolve(
        project_id=PROJECT_ID,
    )
    protocol = packet.roles[0]

    assert protocol.status == "incomplete"
    assert protocol.limitation_codes == (
        "protocol_supplementary_unresolved",
    )


def test_resolver_ignores_supplementary_declared_for_superseded_main() -> None:
    old_main, stale_erratum, _ = _composite_scenario_entries()
    new_main = _entry("protocol", 3)
    _promote_entries(new_main)
    registry = _CompositeRegistry(
        [old_main, stale_erratum, new_main],
        [
            _span(old_main, 1),
            _span(old_main, 2),
            _span(stale_erratum, 1),
            _span(stale_erratum, 2),
            _span(new_main, 1),
            _span(new_main, 2),
        ],
    )

    packet = MonitoringDocumentEvidenceResolver(registry).resolve(
        project_id=PROJECT_ID,
    )
    protocol = packet.roles[0]

    assert protocol.status == "current"
    assert protocol.binding is not None
    assert protocol.binding.source_entry_id == new_main.entry_id
    assert protocol.supplementary_bindings == ()


def test_resolver_selection_cannot_promote_supplementary_file_to_main() -> None:
    main, erratum, spans = _composite_scenario_entries()
    registry = _CompositeRegistry([main, erratum], spans)

    packet = MonitoringDocumentEvidenceResolver(registry).resolve(
        project_id=PROJECT_ID,
        listing_admission_date="2026-09-03",
        selected_entry_ids={"protocol": erratum.entry_id},
    )
    protocol = packet.roles[0]

    assert protocol.status == "incomplete"
    assert protocol.limitation_codes == ("protocol_selection_stale",)


def test_retrieval_covers_each_file_of_composite_authority() -> None:
    main, erratum, spans = _composite_scenario_entries()
    resolver = MonitoringDocumentEvidenceResolver(
        _CompositeRegistry([main, erratum], spans)
    )
    packet = resolver.resolve(project_id=PROJECT_ID)
    protocol = packet.roles[0]
    main_binding = protocol.binding
    erratum_binding = protocol.supplementary_bindings[0]

    main_retrieval = resolver.retrieve_current_excerpts(
        project_id=PROJECT_ID,
        binding=main_binding,
        query_terms=["bounded"],
    )
    erratum_retrieval = resolver.retrieve_current_excerpts(
        project_id=PROJECT_ID,
        binding=erratum_binding,
        query_terms=["bounded"],
    )

    assert main_retrieval["binding_kind"] == "main"
    assert main_retrieval["source_entry_id"] == main.entry_id
    assert erratum_retrieval["binding_kind"] == "supplementary"
    assert erratum_retrieval["source_entry_id"] == erratum.entry_id
    assert erratum_retrieval["content_sha256"] == erratum.content_hash
    assert {excerpt["source_id"] for excerpt in erratum_retrieval["excerpts"]} == {
        f"{erratum.entry_id}-span-1",
        f"{erratum.entry_id}-span-2",
    }
    assert len(erratum_retrieval["packet_sha256"]) == 64
    with pytest.raises(ValueError, match="binding is no longer current"):
        resolver.retrieve_current_excerpts(
            project_id=PROJECT_ID,
            binding=replace(erratum_binding, content_sha256="f" * 64),
            query_terms=["bounded"],
        )
