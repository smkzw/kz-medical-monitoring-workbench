from __future__ import annotations

import hashlib
import json
import threading
from concurrent.futures import ThreadPoolExecutor
from types import SimpleNamespace
from typing import Any

import pytest

from packages.medical_monitoring.admission.document_authority import (
    AdjudicationDecision,
    DOCUMENT_AUTHORITY_SCHEMA_VERSION,
    CandidateAssessment,
    ConflictDecision,
    DocumentAuthorityAnalysis,
    DocumentAuthorityAdjudicationReview,
    DocumentAuthorityConflictReview,
    DocumentAuthorityError,
    EvidenceReference,
    RoleSelection,
    RoleSupplementaryBinding,
    REPLAY_ADJUDICATION_PROMPT_PAIRS,
    build_anonymous_conflict_packet,
    document_authority_batch_sha256,
)
from packages.medical_monitoring.admission.mapping_gate import (
    DOC_AUTH_PRIMARY_MODEL,
    DOC_AUTH_PRIMARY_PROVIDER,
    DOC_AUTH_VERIFIER_MODEL,
    DOC_AUTH_VERIFIER_PROVIDER,
)
from services.api.app.monitoring_ai_contracts import (
    MONITORING_AI_SCHEMA_VERSION,
    MonitoringAiInputRevision,
    MonitoringAiJobCreate,
    MonitoringAiJobStatus,
    MonitoringAiSourceBinding,
    MonitoringAiTaskType,
    content_sha256,
)
from services.api.app.monitoring_ai_repository import MonitoringAiRepository
from services.api.app.monitoring_ai_service import (
    MonitoringAiRuntimeBinding,
    MonitoringAiService,
)
from services.api.app.monitoring_document_authority_jobs import (
    PROMOTION_RECEIPT_SCHEMA_VERSION,
    load_document_authority_analysis_run,
    promote_document_authority_from_jobs,
    resolve_document_authority_from_jobs,
    submit_document_authority_critique_pair,
    submit_document_authority_review_pair,
    verify_document_authority_promotion_receipt,
)
from services.api.app.monitoring_document_authority_workflow import (
    MonitoringDocumentAuthorityWorkflow,
)
from services.api.app.monitoring_document_candidates import (
    MonitoringDocumentCandidateDecomposer,
)
from services.api.app.monitoring_document_evidence import (
    MonitoringDocumentEvidenceResolver,
)
from services.api.app.source_intake import (
    SourceRegistryService,
    SourceRegistryStore,
    monitoring_authority_receipt_is_complete,
)
from services.api.app.source_content_validation import (
    SourceContentValidationService,
    SourceContentValidationStore,
    SourceExpectedContext,
)
from tests.test_source_registry import _minimal_docx_bytes, _minimal_xlsx_bytes


def _batch() -> dict[str, Any]:
    return {
        "manifest_version": "monitoring-document-candidate-v1",
        "batch_id": "mmbatch_product",
        "authority_status": "not_adjudicated",
        "candidates": [
            {
                "candidate_id": "candidate_protocol",
                "file_id": "file_protocol",
                "content_sha256": "a" * 64,
                "filename": "protocol.docx",
                "role_hypotheses": ["protocol"],
                "technical_status": "ready",
                "extraction_status": "parsed",
                "locator_count": 1,
                "excerpts": [{"locator": "doc:p1", "text": "研究方案 V1.0"}],
                "sheets": [],
            },
            {
                "candidate_id": "candidate_ecrf",
                "file_id": "file_ecrf",
                "content_sha256": "b" * 64,
                "filename": "ecrf.xlsx",
                "role_hypotheses": ["ecrf"],
                "technical_status": "ready",
                "extraction_status": "parsed",
                "locator_count": 1,
                "excerpts": [],
                "sheets": [{"locator": "xlsx:sheet:1", "sheet_name": "AE"}],
            },
        ],
    }


def _analysis(
    batch: dict[str, Any], *, select_ecrf: bool = True
) -> DocumentAuthorityAnalysis:
    return DocumentAuthorityAnalysis(
        schema_version=DOCUMENT_AUTHORITY_SCHEMA_VERSION,
        batch_id=batch["batch_id"],
        input_sha256=document_authority_batch_sha256(batch),
        candidate_assessments=(
            CandidateAssessment(
                candidate_id="candidate_protocol",
                inferred_role="protocol",
                usable=True,
                confidence=0.98,
                document_version="V1.0",
                evidence_locators=("doc:p1",),
            ),
            CandidateAssessment(
                candidate_id="candidate_ecrf",
                inferred_role="ecrf",
                usable=True,
                confidence=0.97,
                evidence_locators=("xlsx:sheet:1",),
            ),
        ),
        role_selections=(
            RoleSelection(
                role="protocol",
                decision="selected",
                selected_candidate_id="candidate_protocol",
                confidence=0.98,
                evidence_locators=("doc:p1",),
                supplementary_bindings=(),
            ),
            RoleSelection(
                role="investigator_brochure",
                decision="missing",
                confidence=0.95,
                supplementary_bindings=(),
            ),
            RoleSelection(
                role="ecrf",
                decision="selected" if select_ecrf else "unresolved",
                selected_candidate_id=("candidate_ecrf" if select_ecrf else ""),
                confidence=0.97 if select_ecrf else 0.5,
                evidence_locators=("xlsx:sheet:1",) if select_ecrf else (),
                supplementary_bindings=(),
            ),
            RoleSelection(
                role="sap",
                decision="missing",
                confidence=0.95,
                supplementary_bindings=(),
            ),
        ),
    )


class _Provider:
    transport_name = "openai_compatible"

    def __init__(
        self,
        provider: str,
        model: str,
        analysis: Any,
        *,
        title: str = "研究文件识别结果",
        text: str = "基于冻结文件完成识别。",
    ):
        self.provider_name = provider
        self.model_name = model
        self.expected_response_model = model
        self.response_model = model
        self.analysis = analysis
        self.title = title
        self.text = text
        self.envelopes = []

    def run(self, envelope):
        self.envelopes.append(envelope)
        input_revision_sha256 = envelope.payload.get(
            "input_revision_sha256"
        ) or envelope.payload["original_task"]["input_revision_sha256"]
        return {
            "schema_version": MONITORING_AI_SCHEMA_VERSION,
            "task_id": envelope.task_id,
            "task_type": "document_authority_analysis",
            "input_revision_sha256": input_revision_sha256,
            "candidates": [{
                "candidate_type": "document_authority_analysis",
                "title": self.title,
                "text": self.text,
                "structured_payload": (
                    self.analysis
                    if isinstance(self.analysis, dict)
                    else self.analysis.model_dump(mode="json")
                ),
            }],
        }


class _ReviewProvider:
    transport_name = "openai_compatible"

    def __init__(self, provider: str, model: str, review: DocumentAuthorityConflictReview):
        self.provider_name = provider
        self.model_name = model
        self.expected_response_model = model
        self.response_model = model
        self.review = review
        self.envelopes = []

    def run(self, envelope):
        self.envelopes.append(envelope)
        return {
            "schema_version": MONITORING_AI_SCHEMA_VERSION,
            "task_id": envelope.task_id,
            "task_type": "document_authority_review",
            "input_revision_sha256": envelope.payload["input_revision_sha256"],
            "candidates": [{
                "candidate_type": "document_authority_review",
                "title": "研究文件冲突复核",
                "text": "基于匿名冻结证据完成复核。",
                "structured_payload": self.review.model_dump(mode="json"),
            }],
        }


class _RepairingProvider(_Provider):
    def run(self, envelope):
        output = super().run(envelope)
        if len(self.envelopes) == 1:
            output["claims"] = []
        return output


class _StaleIdentityThenCorrectProvider(_Provider):
    def run(self, envelope):
        output = super().run(envelope)
        if len(self.envelopes) == 1:
            output["candidates"][0]["structured_payload"][
                "input_sha256"
            ] = envelope.payload["input_revision_sha256"]
        return output


def _runtime(provider: str, model: str, profile: str) -> MonitoringAiRuntimeBinding:
    return MonitoringAiRuntimeBinding(
        profile_id=profile,
        provider=provider,
        model=model,
        env={
            "WORKBENCH_AI_PROVIDER": provider,
            "WORKBENCH_AI_MODEL": model,
            "WORKBENCH_AI_EXPECTED_RESPONSE_MODEL": model,
            "WORKBENCH_AI_TRANSPORT": "openai_compatible",
            "WORKBENCH_AI_BASE_URL": "https://example.invalid/v1",
            "WORKBENCH_AI_API_KEY": "synthetic-key",
            "WORKBENCH_AI_DEPLOYMENT_PROFILE": "local_private_clinical",
        },
    )


@pytest.mark.parametrize(
    ("role", "provider", "model", "profile"),
    (
        (
            "primary",
            DOC_AUTH_PRIMARY_PROVIDER,
            DOC_AUTH_PRIMARY_MODEL,
            "monitoring-document-authority-primary",
        ),
        (
            "verifier",
            DOC_AUTH_VERIFIER_PROVIDER,
            DOC_AUTH_VERIFIER_MODEL,
            "monitoring-document-authority-verifier",
        ),
    ),
)
def test_direct_job_history_builds_server_owned_run_envelope(
    tmp_path, role: str, provider: str, model: str, profile: str
) -> None:
    batch = _batch()
    analysis = _analysis(batch)
    repository = MonitoringAiRepository(tmp_path / "monitoring-ai.sqlite")
    runtime = _runtime(provider, model, profile)
    fake = _Provider(provider, model, analysis)
    service = MonitoringAiService(
        repository,
        runtime_resolver=lambda: runtime,
        provider_factory=lambda _env: fake,
    )
    revision = MonitoringAiInputRevision(
        project_id="project-document-authority",
        batch_revision=batch["batch_id"],
        sources=tuple(
            MonitoringAiSourceBinding(
                source_entry_id=item["file_id"],
                source_content_sha256=item["content_sha256"],
            )
            for item in batch["candidates"]
        ),
    )

    queued = service.submit_document_authority_analysis(
        project_id=revision.project_id,
        input_revision=revision,
        candidate_batch=batch,
        role=role,
    )
    result = service.run_next("synthetic-worker", claim_identity=service.claim_identity())
    run = load_document_authority_analysis_run(
        repository,
        project_id=revision.project_id,
        job_id=queued.job_id,
        candidate_batch=batch,
        role=role,
    )

    assert result.job is not None and result.job.status.value == "completed"
    assert run.job_id == queued.job_id
    assert run.run_id.startswith("monattempt_")
    assert run.analysis == analysis
    assert fake.envelopes[0].payload["input_payload"]["document_authority_role"] == role
    assert "另一模型" not in str(fake.envelopes[0].payload["input_payload"])
    candidate_schema = fake.envelopes[0].payload["output_schema"]["candidates"][0]
    assert "claims" not in candidate_schema
    assert "system_generated_evidence" not in candidate_schema
    assert "document_version只能填写简短版本标识" in fake.envelopes[0].system_prompt
    assert "document_date只能填写单一日期" in fake.envelopes[0].system_prompt
    identity = fake.envelopes[0].payload[
        "document_authority_output_identity"
    ]
    assert identity["structured_payload_field"] == "input_sha256"
    assert identity["required_value"] == document_authority_batch_sha256(batch)
    assert identity["required_value"] != queued.input_revision_sha256
    assert fake.envelopes[0].max_output_tokens == 48_000
    assert "locator_count为0的候选没有授权定位证据" in fake.envelopes[0].system_prompt

    for generation in (6, 7):
        legacy_prompt = f"monitoring-document-authority-{role}-v{generation}"
        legacy = repository.create_or_get(MonitoringAiJobCreate(
            project_id=queued.project_id,
            task_type=queued.task_type,
            input_revision=queued.input_revision,
            input_payload=repository.input_payload(queued.project_id, queued.job_id),
            prompt_version=legacy_prompt,
            profile_id=queued.profile_id,
            provider=queued.provider,
            requested_model=queued.requested_model,
            max_attempts=queued.max_attempts,
            business_key=f"legacy-analysis:{role}:v{generation}:{batch['batch_id']}",
        ))
        legacy_result = service.run_next(
            f"legacy-v{generation}-worker", claim_identity=service.claim_identity()
        )
        legacy_run = load_document_authority_analysis_run(
            repository,
            project_id=revision.project_id,
            job_id=legacy.job_id,
            candidate_batch=batch,
            role=role,
        )
        assert legacy_result.job is not None
        assert legacy_result.job.status == MonitoringAiJobStatus.COMPLETED
        assert legacy_run.prompt_version == legacy_prompt


def test_document_authority_gets_one_schema_only_repair(tmp_path) -> None:
    batch = _batch()
    analysis = _analysis(batch)
    repository = MonitoringAiRepository(tmp_path / "monitoring-ai.sqlite")
    runtime = _runtime(
        DOC_AUTH_PRIMARY_PROVIDER,
        DOC_AUTH_PRIMARY_MODEL,
        "monitoring-document-authority-primary",
    )
    provider = _RepairingProvider(
        DOC_AUTH_PRIMARY_PROVIDER,
        DOC_AUTH_PRIMARY_MODEL,
        analysis,
    )
    service = MonitoringAiService(
        repository,
        runtime_resolver=lambda: runtime,
        provider_factory=lambda _env: provider,
    )
    revision = MonitoringAiInputRevision(
        project_id="project-document-authority",
        batch_revision=batch["batch_id"],
        sources=tuple(
            MonitoringAiSourceBinding(
                source_entry_id=item["file_id"],
                source_content_sha256=item["content_sha256"],
            )
            for item in batch["candidates"]
        ),
    )
    queued = service.submit_document_authority_analysis(
        project_id=revision.project_id,
        input_revision=revision,
        candidate_batch=batch,
        role="primary",
    )

    result = service.run_next("synthetic-worker", claim_identity=service.claim_identity())

    assert result.job is not None and result.job.status.value == "completed"
    assert len(provider.envelopes) == 2
    assert provider.envelopes[1].prompt_version.endswith(":json-repair-1")
    repair = provider.envelopes[1].payload["repair_contract"]
    assert repair["maximum_repairs"] == 1
    assert "不得改变文件选择" in repair["instruction"]
    attempts = repository.attempts(revision.project_id, queued.job_id)
    assert len(attempts) == 1
    assert attempts[0]["outcome"] == "success_repaired"


def test_document_authority_repair_uses_frozen_batch_identity(tmp_path) -> None:
    batch = _batch()
    repository = MonitoringAiRepository(tmp_path / "monitoring-ai.sqlite")
    runtime = _runtime(
        DOC_AUTH_PRIMARY_PROVIDER,
        DOC_AUTH_PRIMARY_MODEL,
        "monitoring-document-authority-primary",
    )
    provider = _StaleIdentityThenCorrectProvider(
        DOC_AUTH_PRIMARY_PROVIDER,
        DOC_AUTH_PRIMARY_MODEL,
        _analysis(batch),
    )
    service = MonitoringAiService(
        repository,
        runtime_resolver=lambda: runtime,
        provider_factory=lambda _env: provider,
    )
    revision = MonitoringAiInputRevision(
        project_id="project-document-authority",
        batch_revision=batch["batch_id"],
        sources=tuple(
            MonitoringAiSourceBinding(
                source_entry_id=item["file_id"],
                source_content_sha256=item["content_sha256"],
            )
            for item in batch["candidates"]
        ),
    )
    service.submit_document_authority_analysis(
        project_id=revision.project_id,
        input_revision=revision,
        candidate_batch=batch,
        role="primary",
    )

    result = service.run_next("synthetic-worker", claim_identity=service.claim_identity())

    assert result.job is not None and result.job.status.value == "completed"
    assert len(provider.envelopes) == 2
    repair = provider.envelopes[1].payload["repair_contract"]["instruction"]
    assert "document_authority_batch_sha256" in repair
    assert "绝不能使用外层input_revision_sha256" in repair


def test_loader_rejects_role_swap(tmp_path) -> None:
    batch = _batch()
    analysis = _analysis(batch)
    repository = MonitoringAiRepository(tmp_path / "monitoring-ai.sqlite")
    runtime = _runtime(
        DOC_AUTH_PRIMARY_PROVIDER,
        DOC_AUTH_PRIMARY_MODEL,
        "monitoring-document-authority-primary",
    )
    service = MonitoringAiService(
        repository,
        runtime_resolver=lambda: runtime,
        provider_factory=lambda _env: _Provider(
            DOC_AUTH_PRIMARY_PROVIDER,
            DOC_AUTH_PRIMARY_MODEL,
            analysis,
        ),
    )
    revision = MonitoringAiInputRevision(
        project_id="project-document-authority",
        batch_revision=batch["batch_id"],
        sources=tuple(
            MonitoringAiSourceBinding(
                source_entry_id=item["file_id"],
                source_content_sha256=item["content_sha256"],
            )
            for item in batch["candidates"]
        ),
    )
    job = service.submit_document_authority_analysis(
        project_id=revision.project_id,
        input_revision=revision,
        candidate_batch=batch,
        role="primary",
    )
    service.run_next("synthetic-worker", claim_identity=service.claim_identity())

    with pytest.raises(DocumentAuthorityError, match="job_identity_invalid"):
        load_document_authority_analysis_run(
            repository,
            project_id=revision.project_id,
            job_id=job.job_id,
            candidate_batch=batch,
            role="verifier",
        )


def test_matching_analysis_jobs_resolve_without_conflict_reviews(tmp_path) -> None:
    batch = _batch()
    analysis = _analysis(batch)
    repository = MonitoringAiRepository(tmp_path / "monitoring-ai.sqlite")
    revision = MonitoringAiInputRevision(
        project_id="project-document-authority",
        batch_revision=batch["batch_id"],
        sources=tuple(
            MonitoringAiSourceBinding(
                source_entry_id=item["file_id"],
                source_content_sha256=item["content_sha256"],
            )
            for item in batch["candidates"]
        ),
    )
    jobs = []
    for role, provider, model in (
        ("primary", DOC_AUTH_PRIMARY_PROVIDER, DOC_AUTH_PRIMARY_MODEL),
        ("verifier", DOC_AUTH_VERIFIER_PROVIDER, DOC_AUTH_VERIFIER_MODEL),
    ):
        service = MonitoringAiService(
            repository,
            runtime_resolver=lambda provider=provider, model=model, role=role: _runtime(
                provider, model, f"monitoring-document-authority-{role}"
            ),
            provider_factory=lambda _env, provider=provider, model=model: _Provider(
                provider, model, analysis
            ),
        )
        job = service.submit_document_authority_analysis(
            project_id=revision.project_id,
            input_revision=revision,
            candidate_batch=batch,
            role=role,
        )
        service.run_next(
            f"{role}-analysis-worker", claim_identity=service.claim_identity()
        )
        jobs.append(job)

    result = resolve_document_authority_from_jobs(
        repository,
        project_id=revision.project_id,
        candidate_batch=batch,
        primary_analysis_job_id=jobs[0].job_id,
        verifier_analysis_job_id=jobs[1].job_id,
    )

    assert result["state"] == "resolved"
    assert result["review_run_ids"] == []
    assert {item["role"] for item in result["document_identities"]} == {
        "protocol",
        "ecrf",
    }


def test_product_workflow_stops_before_models_when_role_candidate_is_incomplete(
    tmp_path,
) -> None:
    pymupdf = pytest.importorskip("pymupdf")
    document = pymupdf.open()
    document.new_page()
    blank_pdf = document.tobytes()
    document.close()
    repository = MonitoringAiRepository(tmp_path / "workflow.sqlite")
    wake_calls = []
    workflow = MonitoringDocumentAuthorityWorkflow(
        repository,
        SimpleNamespace(),
        SimpleNamespace(),
        SourceRegistryService(SourceRegistryStore(tmp_path / "registry.jsonl")),
        worker_wake=lambda: wake_calls.append(True),
    )

    with pytest.raises(DocumentAuthorityError, match="evidence_incomplete"):
        workflow.start(
            project_id="project-document-authority",
            workspace_dir=tmp_path / "workspace",
            files=[("protocol.pdf", blank_pdf)],
        )

    assert repository.list_jobs("project-document-authority") == ()
    assert wake_calls == []


def test_product_workflow_starts_both_models_and_promotes_direct_agreement(
    tmp_path,
) -> None:
    protocol = _minimal_docx_bytes()
    ecrf = _minimal_xlsx_bytes()
    expected_batch = MonitoringDocumentCandidateDecomposer(
        tmp_path / "preview"
    ).decompose_many([
        ("protocol.docx", protocol),
        ("ecrf.xlsx", ecrf),
    ]).to_dict()
    candidate_ids = {
        item["filename"]: item["candidate_id"]
        for item in expected_batch["candidates"]
    }
    analysis = DocumentAuthorityAnalysis(
        schema_version=DOCUMENT_AUTHORITY_SCHEMA_VERSION,
        batch_id=expected_batch["batch_id"],
        input_sha256=document_authority_batch_sha256(expected_batch),
        candidate_assessments=(
            CandidateAssessment(
                candidate_id=candidate_ids["protocol.docx"],
                inferred_role="protocol",
                usable=True,
                confidence=0.99,
                evidence_locators=tuple(
                    item["locator"]
                    for item in next(
                        candidate
                        for candidate in expected_batch["candidates"]
                        if candidate["filename"] == "protocol.docx"
                    )["excerpts"][:1]
                ),
            ),
            CandidateAssessment(
                candidate_id=candidate_ids["ecrf.xlsx"],
                inferred_role="ecrf",
                usable=True,
                confidence=0.99,
                evidence_locators=tuple(
                    item["locator"]
                    for item in next(
                        candidate
                        for candidate in expected_batch["candidates"]
                        if candidate["filename"] == "ecrf.xlsx"
                    )["sheets"][:1]
                ),
            ),
        ),
        role_selections=(
            RoleSelection(
                role="protocol",
                decision="selected",
                selected_candidate_id=candidate_ids["protocol.docx"],
                confidence=0.99,
                evidence_locators=tuple(
                    item["locator"]
                    for item in next(
                        candidate
                        for candidate in expected_batch["candidates"]
                        if candidate["filename"] == "protocol.docx"
                    )["excerpts"][:1]
                ),
                supplementary_bindings=(),
            ),
            RoleSelection(
                role="investigator_brochure",
                decision="missing",
                confidence=0.99,
                supplementary_bindings=(),
            ),
            RoleSelection(
                role="ecrf",
                decision="selected",
                selected_candidate_id=candidate_ids["ecrf.xlsx"],
                confidence=0.99,
                evidence_locators=tuple(
                    item["locator"]
                    for item in next(
                        candidate
                        for candidate in expected_batch["candidates"]
                        if candidate["filename"] == "ecrf.xlsx"
                    )["sheets"][:1]
                ),
                supplementary_bindings=(),
            ),
            RoleSelection(
                role="sap",
                decision="missing",
                confidence=0.99,
                supplementary_bindings=(),
            ),
        ),
    )
    repository = MonitoringAiRepository(tmp_path / "workflow.sqlite")
    services = []
    for role, provider, model in (
        ("primary", DOC_AUTH_PRIMARY_PROVIDER, DOC_AUTH_PRIMARY_MODEL),
        ("verifier", DOC_AUTH_VERIFIER_PROVIDER, DOC_AUTH_VERIFIER_MODEL),
    ):
        services.append(MonitoringAiService(
            repository,
            runtime_resolver=lambda provider=provider, model=model, role=role: _runtime(
                provider, model, f"monitoring-document-authority-{role}"
            ),
            provider_factory=lambda _env, provider=provider, model=model: _Provider(
                provider, model, analysis
            ),
        ))
    wake_calls = []
    registry = SourceRegistryService(SourceRegistryStore(tmp_path / "registry.jsonl"))
    workflow = MonitoringDocumentAuthorityWorkflow(
        repository,
        services[0],
        services[1],
        registry,
        worker_wake=lambda: wake_calls.append(True),
    )

    started = workflow.start(
        project_id="project-document-authority",
        workspace_dir=tmp_path / "workspace",
        files=[("protocol.docx", protocol), ("ecrf.xlsx", ecrf)],
    )
    assert started == {"state": "analyzing", "batch_id": expected_batch["batch_id"]}
    for index, service in enumerate(services):
        service.run_next(
            f"workflow-{index}", claim_identity=service.claim_identity()
        )
    promoted = workflow.advance(
        project_id="project-document-authority",
        workspace_dir=tmp_path / "workspace",
        batch_id=started["batch_id"],
    )

    assert promoted["authority_status"] == "promoted"
    assert len(registry.list_entries("project-document-authority")) == 2
    assert wake_calls == [True]
    receipt = registry.list_entries("project-document-authority")[0].metadata[
        "document_authority_receipt"
    ]
    assert verify_document_authority_promotion_receipt(
        repository,
        project_id="project-document-authority",
        receipt=receipt,
    ) is True
    previous_receipt = json.loads(json.dumps(receipt))
    previous_receipt["schema_version"] = (
        "monitoring-document-authority-promotion-v2"
    )
    previous_receipt.pop("adjudication_job_ids")
    previous_receipt.pop("adjudication_run_ids")
    assert verify_document_authority_promotion_receipt(
        repository,
        project_id="project-document-authority",
        receipt=previous_receipt,
    ) is True
    forged = {**receipt, "analysis_job_ids": ["missing-primary", "missing-verifier"]}
    assert verify_document_authority_promotion_receipt(
        repository,
        project_id="project-document-authority",
        receipt=forged,
    ) is False


def test_worker_rejects_downstream_medical_conclusion_hidden_in_output(
    tmp_path,
) -> None:
    batch = _batch()
    raw_analysis = _analysis(batch).model_dump(mode="json")
    raw_analysis["candidate_assessments"][0]["uncertainty"] = (
        "该患者不良反应判定为3级，应向研究中心发出质疑。"
    )
    repository = MonitoringAiRepository(tmp_path / "monitoring-ai.sqlite")
    runtime = _runtime(
        DOC_AUTH_PRIMARY_PROVIDER,
        DOC_AUTH_PRIMARY_MODEL,
        "monitoring-document-authority-primary",
    )
    service = MonitoringAiService(
        repository,
        runtime_resolver=lambda: runtime,
        provider_factory=lambda _env: _Provider(
            DOC_AUTH_PRIMARY_PROVIDER,
            DOC_AUTH_PRIMARY_MODEL,
            raw_analysis,
        ),
    )
    revision = MonitoringAiInputRevision(
        project_id="project-document-authority",
        batch_revision=batch["batch_id"],
        sources=tuple(
            MonitoringAiSourceBinding(
                source_entry_id=item["file_id"],
                source_content_sha256=item["content_sha256"],
            )
            for item in batch["candidates"]
        ),
    )
    job = service.submit_document_authority_analysis(
        project_id=revision.project_id,
        input_revision=revision,
        candidate_batch=batch,
        role="primary",
    )

    result = service.run_next(
        "synthetic-worker", claim_identity=service.claim_identity()
    )

    assert result.job is not None and result.job.status.value == "failed"
    assert repository.candidates(revision.project_id, job.job_id) == ()


def test_product_workflow_starts_blind_conflict_review_without_user_choice(
    tmp_path,
) -> None:
    batch = _batch()
    project_id = "project-document-authority"
    workspace = tmp_path / "workspace"
    batch_dir = workspace / "document_authority_candidates" / "batches"
    batch_dir.mkdir(parents=True)
    (batch_dir / f"{batch['batch_id']}.json").write_text(
        json.dumps(batch, ensure_ascii=False),
        encoding="utf-8",
    )
    repository = MonitoringAiRepository(tmp_path / "workflow-conflict.sqlite")
    revision = MonitoringAiInputRevision(
        project_id=project_id,
        batch_revision=batch["batch_id"],
        sources=tuple(
            MonitoringAiSourceBinding(
                source_entry_id=item["file_id"],
                source_content_sha256=item["content_sha256"],
            )
            for item in batch["candidates"]
        ),
    )
    services = []
    for role, provider, model, analysis in (
        (
            "primary",
            DOC_AUTH_PRIMARY_PROVIDER,
            DOC_AUTH_PRIMARY_MODEL,
            _analysis(batch, select_ecrf=False),
        ),
        (
            "verifier",
            DOC_AUTH_VERIFIER_PROVIDER,
            DOC_AUTH_VERIFIER_MODEL,
            _analysis(batch),
        ),
    ):
        service = MonitoringAiService(
            repository,
            runtime_resolver=lambda provider=provider, model=model, role=role: _runtime(
                provider, model, f"monitoring-document-authority-{role}"
            ),
            provider_factory=lambda _env, provider=provider, model=model, analysis=analysis: _Provider(
                provider, model, analysis
            ),
        )
        service.submit_document_authority_analysis(
            project_id=project_id,
            input_revision=revision,
            candidate_batch=batch,
            role=role,
        )
        service.run_next(role, claim_identity=service.claim_identity())
        services.append(service)
    wake_calls = []
    workflow = MonitoringDocumentAuthorityWorkflow(
        repository,
        services[0],
        services[1],
        SourceRegistryService(SourceRegistryStore(tmp_path / "registry.jsonl")),
        worker_wake=lambda: wake_calls.append(True),
    )

    primary_job = repository.list_jobs(
        project_id,
        task_type=MonitoringAiTaskType.DOCUMENT_AUTHORITY_ANALYSIS.value,
        business_key_prefix="document-authority-analysis:primary:",
    )[0]
    verifier_job = repository.list_jobs(
        project_id,
        task_type=MonitoringAiTaskType.DOCUMENT_AUTHORITY_ANALYSIS.value,
        business_key_prefix="document-authority-analysis:verifier:",
    )[0]
    primary_run = load_document_authority_analysis_run(
        repository,
        project_id=project_id,
        job_id=primary_job.job_id,
        candidate_batch=batch,
        role="primary",
    )
    verifier_run = load_document_authority_analysis_run(
        repository,
        project_id=project_id,
        job_id=verifier_job.job_id,
        candidate_batch=batch,
        role="verifier",
    )
    packet = build_anonymous_conflict_packet(batch, primary_run, verifier_run)
    services[0].submit_document_authority_review(
        project_id=project_id,
        input_revision=revision,
        conflict_packet=packet,
        source_bindings=[
            {
                "candidate_id": item["candidate_id"],
                "source_entry_id": item["file_id"],
                "source_content_sha256": item["content_sha256"],
            }
            for item in batch["candidates"]
        ],
        role="primary",
    )

    result = workflow.advance(
        project_id=project_id,
        workspace_dir=workspace,
        batch_id=batch["batch_id"],
    )

    reviews = repository.list_jobs(
        project_id,
        task_type=MonitoringAiTaskType.DOCUMENT_AUTHORITY_REVIEW.value,
    )
    assert result == {"state": "reviewing", "batch_id": batch["batch_id"]}
    assert len(reviews) == 2
    assert {job.provider for job in reviews} == {
        DOC_AUTH_PRIMARY_PROVIDER,
        DOC_AUTH_VERIFIER_PROVIDER,
    }
    assert wake_calls == [True]


def test_product_workflow_retries_one_terminal_failure_once() -> None:
    retries = []
    wakes = []
    workflow = object.__new__(MonitoringDocumentAuthorityWorkflow)
    workflow.repository = SimpleNamespace(
        retry_terminal=lambda project_id, job_id, **kwargs: retries.append(
            (project_id, job_id, kwargs["current_input_revision_sha256"])
        )
    )
    workflow.worker_wake = lambda: wakes.append(True)
    failed = SimpleNamespace(
        project_id="project",
        job_id="failed-job",
        status=MonitoringAiJobStatus.FAILED,
        max_attempts=2,
        contract_retirement_code="",
        input_revision_sha256="a" * 64,
    )
    complete = SimpleNamespace(
        project_id="project",
        job_id="complete-job",
        status=MonitoringAiJobStatus.COMPLETED,
        max_attempts=2,
        contract_retirement_code="",
        input_revision_sha256="a" * 64,
    )

    assert workflow._recover_failed_once((failed, complete)) is True
    assert retries == [("project", "failed-job", "a" * 64)]
    assert wakes == [True]
    failed.max_attempts = 4
    assert workflow._recover_failed_once((failed, complete)) is False


def test_worker_rejects_downstream_medical_conclusion_in_outer_copy(
    tmp_path,
) -> None:
    batch = _batch()
    repository = MonitoringAiRepository(tmp_path / "monitoring-ai.sqlite")
    runtime = _runtime(
        DOC_AUTH_PRIMARY_PROVIDER,
        DOC_AUTH_PRIMARY_MODEL,
        "monitoring-document-authority-primary",
    )
    provider = _Provider(
        DOC_AUTH_PRIMARY_PROVIDER,
        DOC_AUTH_PRIMARY_MODEL,
        _analysis(batch),
        text="该患者不良反应判定为3级，应向研究中心发出质疑。",
    )
    service = MonitoringAiService(
        repository,
        runtime_resolver=lambda: runtime,
        provider_factory=lambda _env: provider,
    )
    revision = MonitoringAiInputRevision(
        project_id="project-document-authority",
        batch_revision=batch["batch_id"],
        sources=tuple(
            MonitoringAiSourceBinding(
                source_entry_id=item["file_id"],
                source_content_sha256=item["content_sha256"],
            )
            for item in batch["candidates"]
        ),
    )
    job = service.submit_document_authority_analysis(
        project_id=revision.project_id,
        input_revision=revision,
        candidate_batch=batch,
        role="primary",
    )

    result = service.run_next(
        "synthetic-worker", claim_identity=service.claim_identity()
    )

    # N4合同：外层title/text是固定信封。越界文本不导致作业失败，而是
    # 被确定性归一为受控值后完成——医学结论文本必须完全不可见。
    assert result.job is not None and result.job.status.value == "completed"
    stored = repository.candidates(revision.project_id, job.job_id)
    assert len(stored) == 1
    assert stored[0].title == "研究文件识别结果"
    assert stored[0].text == "系统已基于冻结文件内容判断文件类别与版本。"
    assert "该患者不良反应判定为3级" not in stored[0].text
    assert len(provider.envelopes) == 1
    attempts = repository.attempts(revision.project_id, job.job_id)
    assert len(attempts) == 1


def test_repository_jobs_drive_blind_review_and_internal_adjudication(
    tmp_path,
) -> None:
    batch = _batch()
    repository = MonitoringAiRepository(tmp_path / "monitoring-ai.sqlite")
    revision = MonitoringAiInputRevision(
        project_id="project-document-authority",
        batch_revision=batch["batch_id"],
        sources=tuple(
            MonitoringAiSourceBinding(
                source_entry_id=item["file_id"],
                source_content_sha256=item["content_sha256"],
            )
            for item in batch["candidates"]
        ),
    )
    primary_runtime = _runtime(
        DOC_AUTH_PRIMARY_PROVIDER,
        DOC_AUTH_PRIMARY_MODEL,
        "monitoring-document-authority-primary",
    )
    verifier_runtime = _runtime(
        DOC_AUTH_VERIFIER_PROVIDER,
        DOC_AUTH_VERIFIER_MODEL,
        "monitoring-document-authority-verifier",
    )
    primary_service = MonitoringAiService(
        repository,
        runtime_resolver=lambda: primary_runtime,
        provider_factory=lambda _env: _Provider(
            DOC_AUTH_PRIMARY_PROVIDER,
            DOC_AUTH_PRIMARY_MODEL,
            _analysis(batch, select_ecrf=False),
        ),
    )
    verifier_service = MonitoringAiService(
        repository,
        runtime_resolver=lambda: verifier_runtime,
        provider_factory=lambda _env: _Provider(
            DOC_AUTH_VERIFIER_PROVIDER,
            DOC_AUTH_VERIFIER_MODEL,
            _analysis(batch),
        ),
    )
    primary_job = primary_service.submit_document_authority_analysis(
        project_id=revision.project_id,
        input_revision=revision,
        candidate_batch=batch,
        role="primary",
    )
    verifier_job = verifier_service.submit_document_authority_analysis(
        project_id=revision.project_id,
        input_revision=revision,
        candidate_batch=batch,
        role="verifier",
    )
    primary_service.run_next(
        "primary-analysis-worker", claim_identity=primary_service.claim_identity()
    )
    verifier_service.run_next(
        "verifier-analysis-worker", claim_identity=verifier_service.claim_identity()
    )

    primary_review_provider: dict[str, _ReviewProvider] = {}
    verifier_review_provider: dict[str, _ReviewProvider] = {}
    primary_review_service = MonitoringAiService(
        repository,
        runtime_resolver=lambda: primary_runtime,
        provider_factory=lambda _env: primary_review_provider["value"],
    )
    verifier_review_service = MonitoringAiService(
        repository,
        runtime_resolver=lambda: verifier_runtime,
        provider_factory=lambda _env: verifier_review_provider["value"],
    )
    packet, primary_review_job, verifier_review_job = (
        submit_document_authority_review_pair(
            primary_review_service,
            verifier_review_service,
            input_revision=revision,
            candidate_batch=batch,
            primary_analysis_job_id=primary_job.job_id,
            verifier_analysis_job_id=verifier_job.job_id,
        )
    )
    review = DocumentAuthorityConflictReview(
        schema_version=DOCUMENT_AUTHORITY_SCHEMA_VERSION,
        conflict_packet_sha256=packet["conflict_packet_sha256"],
        decisions=(
            ConflictDecision(
                role="ecrf",
                decision="selected",
                selected_candidate_id="candidate_ecrf",
                document_version="V1.0",
                confidence=0.97,
                considered_candidate_ids=("candidate_protocol", "candidate_ecrf"),
                supplementary_candidate_ids=(),
                evidence_references=(
                    EvidenceReference(
                        candidate_id="candidate_protocol", locator="doc:p1"
                    ),
                    EvidenceReference(
                        candidate_id="candidate_ecrf", locator="xlsx:sheet:1"
                    ),
                ),
            ),
        ),
    )
    primary_review_provider["value"] = _ReviewProvider(
        DOC_AUTH_PRIMARY_PROVIDER, DOC_AUTH_PRIMARY_MODEL, review
    )
    unresolved_review = DocumentAuthorityConflictReview(
        schema_version=DOCUMENT_AUTHORITY_SCHEMA_VERSION,
        conflict_packet_sha256=packet["conflict_packet_sha256"],
        decisions=(
            review.decisions[0].model_copy(
                update={
                    "decision": "unresolved",
                    "selected_candidate_id": "",
                    "document_version": "",
                    "confidence": 0.4,
                }
            ),
        ),
    )
    verifier_review_provider["value"] = _ReviewProvider(
        DOC_AUTH_VERIFIER_PROVIDER,
        DOC_AUTH_VERIFIER_MODEL,
        unresolved_review,
    )
    primary_review_service.run_next(
        "primary-review-worker", claim_identity=primary_review_service.claim_identity()
    )
    verifier_review_service.run_next(
        "verifier-review-worker", claim_identity=verifier_review_service.claim_identity()
    )

    legacy_review_jobs = []
    for original, prompt_version in (
        (primary_review_job, "monitoring-document-authority-review-primary-v6"),
        (verifier_review_job, "monitoring-document-authority-review-verifier-v6"),
    ):
        legacy_review_jobs.append(repository.create_or_get(MonitoringAiJobCreate(
            project_id=original.project_id,
            task_type=original.task_type,
            input_revision=original.input_revision,
            input_payload=repository.input_payload(original.project_id, original.job_id),
            prompt_version=prompt_version,
            profile_id=original.profile_id,
            provider=original.provider,
            requested_model=original.requested_model,
            max_attempts=original.max_attempts,
            business_key=f"{original.business_key}:legacy-v6",
        )))
    primary_review_service.run_next(
        "primary-review-v6-worker", claim_identity=primary_review_service.claim_identity()
    )
    verifier_review_service.run_next(
        "verifier-review-v6-worker", claim_identity=verifier_review_service.claim_identity()
    )
    legacy_initial = resolve_document_authority_from_jobs(
        repository,
        project_id=revision.project_id,
        candidate_batch=batch,
        primary_analysis_job_id=primary_job.job_id,
        verifier_analysis_job_id=verifier_job.job_id,
        primary_review_job_id=legacy_review_jobs[0].job_id,
        verifier_review_job_id=legacy_review_jobs[1].job_id,
    )
    assert legacy_initial["state"] == "needs_user_input"

    initial = resolve_document_authority_from_jobs(
        repository,
        project_id=revision.project_id,
        candidate_batch=batch,
        primary_analysis_job_id=primary_job.job_id,
        verifier_analysis_job_id=verifier_job.job_id,
        primary_review_job_id=primary_review_job.job_id,
        verifier_review_job_id=verifier_review_job.job_id,
    )
    assert initial["state"] == "needs_user_input"

    adjudication_payload = review.decisions[0].model_dump(mode="python")
    adjudication_payload["evidence_references"] = (
        EvidenceReference(
            candidate_id="candidate_ecrf", locator="xlsx:sheet:1"
        ),
    )
    adjudication_review = DocumentAuthorityAdjudicationReview(
        schema_version=review.schema_version,
        conflict_packet_sha256=review.conflict_packet_sha256,
        decisions=(
            AdjudicationDecision(
                **adjudication_payload,
                excluded_candidate_ids=("candidate_protocol",),
                rationale="正文作用与版本关系支持该完整处置集合。",
            ),
        ),
    )
    critique_payload = adjudication_payload.copy()
    critique_payload["evidence_references"] = (
        EvidenceReference(candidate_id="candidate_protocol", locator="doc:p1"),
        EvidenceReference(candidate_id="candidate_ecrf", locator="xlsx:sheet:1"),
    )
    critique_review = DocumentAuthorityAdjudicationReview(
        schema_version=review.schema_version,
        conflict_packet_sha256=review.conflict_packet_sha256,
        decisions=(
            AdjudicationDecision(
                **critique_payload,
                excluded_candidate_ids=("candidate_protocol",),
                rationale="回到冻结证据后，完整处置与文件关系仍支持该结论。",
            ),
        ),
    )
    primary_review_provider["value"] = _ReviewProvider(
        DOC_AUTH_PRIMARY_PROVIDER,
        DOC_AUTH_PRIMARY_MODEL,
        adjudication_review,
    )
    verifier_review_provider["value"] = _ReviewProvider(
        DOC_AUTH_VERIFIER_PROVIDER,
        DOC_AUTH_VERIFIER_MODEL,
        adjudication_review,
    )
    workspace = tmp_path / "workflow-workspace"
    batch_dir = workspace / "document_authority_candidates" / "batches"
    batch_dir.mkdir(parents=True)
    (batch_dir / f"{batch['batch_id']}.json").write_text(
        json.dumps(batch, ensure_ascii=False), encoding="utf-8"
    )
    wakes: list[bool] = []
    workflow = MonitoringDocumentAuthorityWorkflow(
        repository,
        primary_review_service,
        verifier_review_service,
        SourceRegistryService(SourceRegistryStore(tmp_path / "adjudication-registry.jsonl")),
        worker_wake=lambda: wakes.append(True),
    )
    state = workflow.advance(
        project_id=revision.project_id,
        workspace_dir=workspace,
        batch_id=batch["batch_id"],
    )
    assert state == {"state": "adjudicating", "batch_id": batch["batch_id"]}
    assert wakes == [True]
    adjudication_jobs = repository.list_jobs(
        revision.project_id,
        task_type=MonitoringAiTaskType.DOCUMENT_AUTHORITY_REVIEW.value,
        business_key_prefix="document-authority-adjudication:",
    )
    assert len(adjudication_jobs) == 2
    assert all(":v8:" in job.business_key for job in adjudication_jobs)
    primary_adjudication_job = next(
        job for job in adjudication_jobs if ":primary:" in job.business_key
    )
    verifier_adjudication_job = next(
        job for job in adjudication_jobs if ":verifier:" in job.business_key
    )
    context = repository.input_payload(
        revision.project_id, primary_adjudication_job.job_id
    )["document_authority_adjudication_context"]
    primary_review_service.run_next(
        "primary-adjudication-worker",
        claim_identity=primary_review_service.claim_identity(),
    )
    verifier_review_service.run_next(
        "verifier-adjudication-worker",
        claim_identity=verifier_review_service.claim_identity(),
    )

    result = resolve_document_authority_from_jobs(
        repository,
        project_id=revision.project_id,
        candidate_batch=batch,
        primary_analysis_job_id=primary_job.job_id,
        verifier_analysis_job_id=verifier_job.job_id,
        primary_review_job_id=primary_review_job.job_id,
        verifier_review_job_id=verifier_review_job.job_id,
        primary_adjudication_job_id=primary_adjudication_job.job_id,
        verifier_adjudication_job_id=verifier_adjudication_job.job_id,
    )

    assert result["state"] == "resolved"
    assert result["user_question"] == ""
    assert len(result["adjudication_run_ids"]) == 2
    assert context["unresolved_roles"] == ["ecrf"]

    competing_payload = adjudication_payload.copy()
    competing_payload.update({
        "selected_candidate_id": "candidate_protocol",
        "evidence_references": (
            EvidenceReference(candidate_id="candidate_protocol", locator="doc:p1"),
            EvidenceReference(candidate_id="candidate_ecrf", locator="xlsx:sheet:1"),
        ),
    })
    competing_adjudication = DocumentAuthorityAdjudicationReview(
        schema_version=review.schema_version,
        conflict_packet_sha256=review.conflict_packet_sha256,
        decisions=(
            AdjudicationDecision(
                **competing_payload,
                excluded_candidate_ids=("candidate_ecrf",),
                rationale="反方选项对候选文件关系作出了不同解释。",
            ),
        ),
    )
    verifier_review_provider["value"] = _ReviewProvider(
        DOC_AUTH_VERIFIER_PROVIDER,
        DOC_AUTH_VERIFIER_MODEL,
        competing_adjudication,
    )
    competing_job = repository.create_or_get(MonitoringAiJobCreate(
        project_id=verifier_adjudication_job.project_id,
        task_type=verifier_adjudication_job.task_type,
        input_revision=verifier_adjudication_job.input_revision,
        input_payload=repository.input_payload(
            verifier_adjudication_job.project_id,
            verifier_adjudication_job.job_id,
        ),
        prompt_version=verifier_adjudication_job.prompt_version,
        profile_id=verifier_adjudication_job.profile_id,
        provider=verifier_adjudication_job.provider,
        requested_model=verifier_adjudication_job.requested_model,
        max_attempts=verifier_adjudication_job.max_attempts,
        business_key=f"{verifier_adjudication_job.business_key}:competing",
    ))
    verifier_review_service.run_next(
        "verifier-competing-adjudication-worker",
        claim_identity=verifier_review_service.claim_identity(),
    )
    critique_context, primary_critique_job, verifier_critique_job = (
        submit_document_authority_critique_pair(
            primary_review_service,
            verifier_review_service,
            input_revision=revision,
            candidate_batch=batch,
            primary_analysis_job_id=primary_job.job_id,
            verifier_analysis_job_id=verifier_job.job_id,
            primary_review_job_id=primary_review_job.job_id,
            verifier_review_job_id=verifier_review_job.job_id,
            primary_adjudication_job_id=primary_adjudication_job.job_id,
            verifier_adjudication_job_id=competing_job.job_id,
        )
    )
    serialized_context = json.dumps(critique_context, ensure_ascii=False)
    assert critique_context["schema_version"] == (
        "monitoring-document-authority-critique-v1"
    )
    assert "cms-smk" not in serialized_context
    assert "zhipu-coding-plan" not in serialized_context
    assert primary_critique_job.business_key.startswith(
        "document-authority-critique:primary:v2:"
    )
    assert verifier_critique_job.business_key.startswith(
        "document-authority-critique:verifier:v2:"
    )
    primary_review_provider["value"] = _ReviewProvider(
        DOC_AUTH_PRIMARY_PROVIDER,
        DOC_AUTH_PRIMARY_MODEL,
        critique_review,
    )
    verifier_review_provider["value"] = _ReviewProvider(
        DOC_AUTH_VERIFIER_PROVIDER,
        DOC_AUTH_VERIFIER_MODEL,
        critique_review,
    )
    primary_review_service.run_next(
        "primary-critique-worker",
        claim_identity=primary_review_service.claim_identity(),
    )
    verifier_review_service.run_next(
        "verifier-critique-worker",
        claim_identity=verifier_review_service.claim_identity(),
    )
    critique_resolution = resolve_document_authority_from_jobs(
        repository,
        project_id=revision.project_id,
        candidate_batch=batch,
        primary_analysis_job_id=primary_job.job_id,
        verifier_analysis_job_id=verifier_job.job_id,
        primary_review_job_id=primary_review_job.job_id,
        verifier_review_job_id=verifier_review_job.job_id,
        primary_adjudication_job_id=primary_adjudication_job.job_id,
        verifier_adjudication_job_id=competing_job.job_id,
        primary_critique_job_id=primary_critique_job.job_id,
        verifier_critique_job_id=verifier_critique_job.job_id,
    )
    assert critique_resolution["state"] == "resolved"
    assert len(critique_resolution["critique_run_ids"]) == 2

    persisted_replay_jobs = []
    for primary_prompt, verifier_prompt in sorted(REPLAY_ADJUDICATION_PROMPT_PAIRS):
        generation = primary_prompt.rsplit("-", 1)[-1]
        replay_jobs = []
        for original, prompt_version in (
            (primary_adjudication_job, primary_prompt),
            (verifier_adjudication_job, verifier_prompt),
        ):
            replay_jobs.append(repository.create_or_get(MonitoringAiJobCreate(
                project_id=original.project_id,
                task_type=original.task_type,
                input_revision=original.input_revision,
                input_payload=repository.input_payload(
                    original.project_id, original.job_id
                ),
                prompt_version=prompt_version,
                profile_id=original.profile_id,
                provider=original.provider,
                requested_model=original.requested_model,
                max_attempts=original.max_attempts,
                business_key=f"{original.business_key}:replay:{generation}",
            )))
        primary_review_service.run_next(
            f"primary-{generation}-replay-worker",
            claim_identity=primary_review_service.claim_identity(),
        )
        verifier_review_service.run_next(
            f"verifier-{generation}-replay-worker",
            claim_identity=verifier_review_service.claim_identity(),
        )
        replay = resolve_document_authority_from_jobs(
            repository,
            project_id=revision.project_id,
            candidate_batch=batch,
            primary_analysis_job_id=primary_job.job_id,
            verifier_analysis_job_id=verifier_job.job_id,
            primary_review_job_id=primary_review_job.job_id,
            verifier_review_job_id=verifier_review_job.job_id,
            primary_adjudication_job_id=replay_jobs[0].job_id,
            verifier_adjudication_job_id=replay_jobs[1].job_id,
        )
        assert replay["state"] == "resolved"
        persisted_replay_jobs.append(replay_jobs)

    with pytest.raises(DocumentAuthorityError, match="run_identity_invalid"):
        resolve_document_authority_from_jobs(
            repository,
            project_id=revision.project_id,
            candidate_batch=batch,
            primary_analysis_job_id=primary_job.job_id,
            verifier_analysis_job_id=verifier_job.job_id,
            primary_review_job_id=primary_review_job.job_id,
            verifier_review_job_id=verifier_review_job.job_id,
            primary_adjudication_job_id=persisted_replay_jobs[1][0].job_id,
            verifier_adjudication_job_id=persisted_replay_jobs[0][1].job_id,
        )
    assert all(
        "document_authority_source_bindings"
        not in provider.envelopes[0].payload["input_payload"]
        for provider in (
            primary_review_provider["value"],
            verifier_review_provider["value"],
        )
    )
    for provider in (
        primary_review_provider["value"],
        verifier_review_provider["value"],
    ):
        coverage = provider.envelopes[0].payload[
            "document_authority_review_coverage_contract"
        ]
        assert coverage["conflict_roles"] == ["ecrf"]
        assert set(
            coverage["by_role"]["ecrf"][
                "required_considered_candidate_ids"
            ]
        ) == {"candidate_protocol", "candidate_ecrf"}
        assert set(
            coverage["by_role"]["ecrf"][
                "evidence_required_candidate_ids"
            ]
        ) == {"candidate_protocol", "candidate_ecrf"}
        assert coverage["by_role"]["ecrf"][
            "required_locator_by_candidate"
        ] == {
            "candidate_protocol": "doc:p1",
            "candidate_ecrf": "xlsx:sheet:1",
        }
        assert "包括未入选候选" in provider.envelopes[0].system_prompt
        assert "required_locator_by_candidate" in provider.envelopes[0].system_prompt
        critique_input = provider.envelopes[0].payload["input_payload"]
        assert critique_input["document_authority_adjudication_context"] == (
            critique_context
        )
        assert "document_authority_source_bindings" not in critique_input
        assert "匿名相互质询" in provider.envelopes[0].system_prompt
        assert "不属于当前权威补充" in provider.envelopes[0].system_prompt
        assert "重复载体" in provider.envelopes[0].system_prompt
        assert "不论两者是PDF、Word或其他格式" in provider.envelopes[0].system_prompt
        assert "counter_evidence_references" in provider.envelopes[0].system_prompt


def test_resolved_authority_promotes_selected_documents_atomically(
    tmp_path,
    monkeypatch,
) -> None:
    protocol = _minimal_docx_bytes()
    ecrf = _minimal_xlsx_bytes()
    candidate_root = tmp_path / "candidates"
    batch = MonitoringDocumentCandidateDecomposer(candidate_root).decompose_many([
        ("protocol.docx", protocol),
        ("ecrf.xlsx", ecrf),
    ]).to_dict()
    candidate_ids = {
        candidate["filename"]: candidate["candidate_id"]
        for candidate in batch["candidates"]
    }
    monkeypatch.setattr(
        "services.api.app.monitoring_document_authority_jobs."
        "resolve_document_authority_from_jobs",
        lambda *_args, **_kwargs: {
            "schema_version": DOCUMENT_AUTHORITY_SCHEMA_VERSION,
            "batch_id": batch["batch_id"],
            "input_sha256": document_authority_batch_sha256(batch),
            "state": "resolved",
            "resolved_roles": [
                {
                    "role": "protocol",
                    "status": "selected",
                    "candidate_id": candidate_ids["protocol.docx"],
                },
                {"role": "investigator_brochure", "status": "missing", "candidate_id": ""},
                {
                    "role": "ecrf",
                    "status": "selected",
                    "candidate_id": candidate_ids["ecrf.xlsx"],
                },
                {"role": "sap", "status": "missing", "candidate_id": ""},
            ],
            "unresolved_roles": [],
            "user_question": "",
            "review_run_ids": ["run_primary", "run_verifier"],
            "analysis_run_ids": ["analysis_run_primary", "analysis_run_verifier"],
            "document_identities": [
                {
                    "role": "protocol",
                    "candidate_id": candidate_ids["protocol.docx"],
                    "document_version": "V1.0",
                    "document_date": "2026-09-03",
                },
                {
                    "role": "ecrf",
                    "candidate_id": candidate_ids["ecrf.xlsx"],
                    "document_version": "V1.0",
                    "document_date": "2026-09-03",
                },
            ],
        },
    )
    registry = SourceRegistryService(
        SourceRegistryStore(tmp_path / "registry.jsonl")
    )

    result = promote_document_authority_from_jobs(
        MonitoringAiRepository(tmp_path / "jobs.sqlite"),
        project_id="project-document-authority",
        candidate_batch=batch,
        candidate_root=candidate_root,
        source_registry=registry,
        primary_analysis_job_id="analysis-primary",
        verifier_analysis_job_id="analysis-verifier",
        conflict_packet={},
        primary_review_job_id="review-primary",
        verifier_review_job_id="review-verifier",
    )

    assert result["authority_status"] == "promoted"
    assert {item["role"] for item in result["registrations"]} == {"protocol", "ecrf"}
    entries = registry.list_entries("project-document-authority")
    assert {entry.source_kind for entry in entries} == {
        "protocol_docx",
        "ecrf",
    }
    assert all(
        entry.metadata["document_authority_receipt_sha256"]
        == result["promotion_receipt_sha256"]
        for entry in entries
    )
    replay = promote_document_authority_from_jobs(
        MonitoringAiRepository(tmp_path / "jobs.sqlite"),
        project_id="project-document-authority",
        candidate_batch=batch,
        candidate_root=candidate_root,
        source_registry=registry,
        primary_analysis_job_id="analysis-primary",
        verifier_analysis_job_id="analysis-verifier",
        conflict_packet={},
        primary_review_job_id="review-primary",
        verifier_review_job_id="review-verifier",
    )
    assert replay["promotion_receipt_sha256"] == result["promotion_receipt_sha256"]
    assert len(registry.list_entries("project-document-authority")) == 2
    re_adjudicated = promote_document_authority_from_jobs(
        MonitoringAiRepository(tmp_path / "jobs.sqlite"),
        project_id="project-document-authority",
        candidate_batch=batch,
        candidate_root=candidate_root,
        source_registry=registry,
        primary_analysis_job_id="analysis-primary-2",
        verifier_analysis_job_id="analysis-verifier-2",
        conflict_packet={},
        primary_review_job_id="review-primary-2",
        verifier_review_job_id="review-verifier-2",
    )
    assert re_adjudicated["promotion_receipt_sha256"] != (
        result["promotion_receipt_sha256"]
    )
    assert all(
        entry.metadata["document_authority_receipt_sha256"]
        == re_adjudicated["promotion_receipt_sha256"]
        for entry in registry.list_entries("project-document-authority")
    )

    rollback_registry = SourceRegistryService(
        SourceRegistryStore(tmp_path / "rollback-registry.jsonl")
    )
    register = rollback_registry.register_monitoring_mapping_document

    def fail_second_registration(*args, **kwargs):
        if kwargs.get("document_role") == "ecrf":
            raise RuntimeError("synthetic second registration failure")
        return register(*args, **kwargs)

    monkeypatch.setattr(
        rollback_registry,
        "register_monitoring_mapping_document",
        fail_second_registration,
    )
    with pytest.raises(RuntimeError, match="second registration failure"):
        promote_document_authority_from_jobs(
            MonitoringAiRepository(tmp_path / "rollback-jobs.sqlite"),
            project_id="project-document-authority",
            candidate_batch=batch,
            candidate_root=candidate_root,
            source_registry=rollback_registry,
            primary_analysis_job_id="analysis-primary",
            verifier_analysis_job_id="analysis-verifier",
            conflict_packet={},
            primary_review_job_id="review-primary",
            verifier_review_job_id="review-verifier",
        )
    assert rollback_registry.list_entries("project-document-authority") == []

    batch_manifest = candidate_root / "batches" / f"{batch['batch_id']}.json"
    batch_manifest.unlink()
    missing_batch_registry = SourceRegistryService(
        SourceRegistryStore(tmp_path / "missing-batch-registry.jsonl")
    )
    with pytest.raises(DocumentAuthorityError, match="batch_manifest_missing"):
        promote_document_authority_from_jobs(
            MonitoringAiRepository(tmp_path / "missing-batch-jobs.sqlite"),
            project_id="project-document-authority",
            candidate_batch=batch,
            candidate_root=candidate_root,
            source_registry=missing_batch_registry,
            primary_analysis_job_id="analysis-primary",
            verifier_analysis_job_id="analysis-verifier",
        )
    assert missing_batch_registry.list_entries("project-document-authority") == []

    MonitoringDocumentCandidateDecomposer(candidate_root).decompose_many([
        ("protocol.docx", protocol),
        ("ecrf.xlsx", ecrf),
    ])
    protocol_evidence_revision = next(
        item["evidence_revision_sha256"]
        for item in batch["candidates"]
        if item["filename"] == "protocol.docx"
    )
    candidate_manifest = (
        candidate_root
        / "manifests"
        / candidate_ids["protocol.docx"]
        / f"{protocol_evidence_revision}.json"
    )
    candidate_manifest.unlink()
    missing_candidate_registry = SourceRegistryService(
        SourceRegistryStore(tmp_path / "missing-candidate-registry.jsonl")
    )
    with pytest.raises(DocumentAuthorityError, match="candidate_manifest_missing"):
        promote_document_authority_from_jobs(
            MonitoringAiRepository(tmp_path / "missing-candidate-jobs.sqlite"),
            project_id="project-document-authority",
            candidate_batch=batch,
            candidate_root=candidate_root,
            source_registry=missing_candidate_registry,
            primary_analysis_job_id="analysis-primary",
            verifier_analysis_job_id="analysis-verifier",
        )
    assert missing_candidate_registry.list_entries("project-document-authority") == []


def _supplement_pdf_bytes(title: str, detail: str) -> bytes:
    pymupdf = pytest.importorskip("pymupdf")
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), title)
    page.insert_text((72, 96), detail)
    payload = document.tobytes()
    document.close()
    return payload


def _composite_protocol_docx_bytes() -> bytes:
    import io

    import docx

    document = docx.Document()
    document.add_paragraph("CMS-D001 研究方案 V2.0")
    document.add_paragraph("入选标准")
    document.add_paragraph("IN-01 受试者需签署知情同意。")
    document.add_paragraph("排除标准")
    document.add_paragraph("EX-01 活动性感染者不得入组。")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _composite_files() -> list[tuple[str, bytes]]:
    return [
        ("protocol.docx", _composite_protocol_docx_bytes()),
        (
            "protocol_errata.pdf",
            _supplement_pdf_bytes(
                "CMS-D001 Protocol Errata V2.0",
                "2026-09-03 inclusion criterion IN-01 correction",
            ),
        ),
        (
            "protocol_amendment.pdf",
            _supplement_pdf_bytes(
                "CMS-D001 Protocol Amendment V2.0",
                "2026-09-03 exclusion criterion EX-01 clarification",
            ),
        ),
        ("ecrf.xlsx", _minimal_xlsx_bytes()),
    ]


def _composite_analysis(batch: dict[str, Any]) -> DocumentAuthorityAnalysis:
    candidate_ids = {
        candidate["filename"]: candidate["candidate_id"]
        for candidate in batch["candidates"]
    }
    excerpts = {
        candidate["filename"]: [
            item["locator"] for item in candidate.get("excerpts", ())
        ]
        + [item["locator"] for item in candidate.get("sheets", ())]
        for candidate in batch["candidates"]
    }
    return DocumentAuthorityAnalysis(
        schema_version=DOCUMENT_AUTHORITY_SCHEMA_VERSION,
        batch_id=batch["batch_id"],
        input_sha256=document_authority_batch_sha256(batch),
        candidate_assessments=(
            CandidateAssessment(
                candidate_id=candidate_ids["protocol.docx"],
                inferred_role="protocol",
                usable=True,
                confidence=0.98,
                document_version="V2.0",
                evidence_locators=tuple(excerpts["protocol.docx"][:1]),
            ),
            CandidateAssessment(
                candidate_id=candidate_ids["protocol_errata.pdf"],
                inferred_role="protocol",
                usable=True,
                confidence=0.96,
                document_version="V2.0",
                evidence_locators=tuple(excerpts["protocol_errata.pdf"][:1]),
            ),
            CandidateAssessment(
                candidate_id=candidate_ids["protocol_amendment.pdf"],
                inferred_role="protocol",
                usable=True,
                confidence=0.95,
                document_version="V2.0",
                evidence_locators=tuple(excerpts["protocol_amendment.pdf"][:1]),
            ),
            CandidateAssessment(
                candidate_id=candidate_ids["ecrf.xlsx"],
                inferred_role="ecrf",
                usable=True,
                confidence=0.97,
                evidence_locators=tuple(excerpts["ecrf.xlsx"][:1]),
            ),
        ),
        role_selections=(
            RoleSelection(
                role="protocol",
                decision="selected",
                selected_candidate_id=candidate_ids["protocol.docx"],
                confidence=0.98,
                evidence_locators=tuple(excerpts["protocol.docx"][:1]),
                supplementary_bindings=(
                    RoleSupplementaryBinding(
                        candidate_id=candidate_ids["protocol_errata.pdf"],
                        evidence_locators=tuple(excerpts["protocol_errata.pdf"][:1]),
                    ),
                    RoleSupplementaryBinding(
                        candidate_id=candidate_ids["protocol_amendment.pdf"],
                        evidence_locators=tuple(excerpts["protocol_amendment.pdf"][:1]),
                    ),
                ),
            ),
            RoleSelection(
                role="investigator_brochure",
                decision="missing",
                confidence=0.95,
                supplementary_bindings=(),
            ),
            RoleSelection(
                role="ecrf",
                decision="selected",
                selected_candidate_id=candidate_ids["ecrf.xlsx"],
                confidence=0.97,
                evidence_locators=tuple(excerpts["ecrf.xlsx"][:1]),
                supplementary_bindings=(),
            ),
            RoleSelection(
                role="sap",
                decision="missing",
                confidence=0.95,
                supplementary_bindings=(),
            ),
        ),
    )


def _run_composite_analysis_jobs(
    tmp_path, batch: dict[str, Any]
) -> tuple[MonitoringAiRepository, list[str]]:
    analysis = _composite_analysis(batch)
    repository = MonitoringAiRepository(tmp_path / "composite-jobs.sqlite")
    job_ids = []
    for role, provider, model in (
        ("primary", DOC_AUTH_PRIMARY_PROVIDER, DOC_AUTH_PRIMARY_MODEL),
        ("verifier", DOC_AUTH_VERIFIER_PROVIDER, DOC_AUTH_VERIFIER_MODEL),
    ):
        service = MonitoringAiService(
            repository,
            runtime_resolver=lambda provider=provider, model=model, role=role: _runtime(
                provider, model, f"monitoring-document-authority-{role}"
            ),
            provider_factory=lambda _env, provider=provider, model=model: _Provider(
                provider, model, analysis
            ),
        )
        revision = MonitoringAiInputRevision(
            project_id="project-document-authority",
            batch_revision=batch["batch_id"],
            sources=tuple(
                MonitoringAiSourceBinding(
                    source_entry_id=item["file_id"],
                    source_content_sha256=item["content_sha256"],
                )
                for item in batch["candidates"]
            ),
        )
        service.submit_document_authority_analysis(
            project_id=revision.project_id,
            input_revision=revision,
            candidate_batch=batch,
            role=role,
        )
        service.run_next(f"{role}-worker", claim_identity=service.claim_identity())
        job_ids.append(
            repository.list_jobs(
                revision.project_id,
                task_type=MonitoringAiTaskType.DOCUMENT_AUTHORITY_ANALYSIS.value,
                business_key_prefix=f"document-authority-analysis:{role}:",
            )[0].job_id
        )
    return repository, job_ids


def test_composite_authority_promotes_main_and_supplementary_files_atomically(
    tmp_path,
) -> None:
    candidate_root = tmp_path / "candidates"
    batch = MonitoringDocumentCandidateDecomposer(candidate_root).decompose_many(
        _composite_files()
    ).to_dict()
    batch = json.loads(
        (candidate_root / "batches" / f"{batch['batch_id']}.json").read_text(
            encoding="utf-8"
        )
    )
    candidate_ids = {
        candidate["filename"]: candidate["candidate_id"]
        for candidate in batch["candidates"]
    }
    repository, job_ids = _run_composite_analysis_jobs(tmp_path, batch)
    registry = SourceRegistryService(
        SourceRegistryStore(tmp_path / "registry.jsonl"),
        content_validation_service=SourceContentValidationService(
            SourceContentValidationStore(tmp_path / "validations.sqlite3")
        ),
        expected_context_resolver=lambda *_args: SourceExpectedContext(),
    )

    result = promote_document_authority_from_jobs(
        repository,
        project_id="project-document-authority",
        candidate_batch=batch,
        candidate_root=candidate_root,
        source_registry=registry,
        primary_analysis_job_id=job_ids[0],
        verifier_analysis_job_id=job_ids[1],
    )

    assert result["authority_status"] == "promoted"
    assert set(result["resolved_roles"][0]["supplementary_candidate_ids"]) == {
        candidate_ids["protocol_errata.pdf"],
        candidate_ids["protocol_amendment.pdf"],
    }
    entries = registry.list_entries("project-document-authority")
    assert len(entries) == 4
    entry_id_by_filename = {
        entry.metadata["filename"]: entry.entry_id for entry in entries
    }
    receipt = entries[0].metadata["document_authority_receipt"]
    assert receipt["schema_version"] == PROMOTION_RECEIPT_SCHEMA_VERSION
    protocol_primary = next(
        item
        for item in receipt["registrations"]
        if item["role"] == "protocol" and item["binding_kind"] == "primary"
    )
    protocol_supplements = [
        item
        for item in receipt["registrations"]
        if item["role"] == "protocol" and item["binding_kind"] == "supplementary"
    ]
    ecrf_primary = next(
        item
        for item in receipt["registrations"]
        if item["role"] == "ecrf" and item["binding_kind"] == "primary"
    )
    assert len(protocol_supplements) == 2
    assert protocol_primary["candidate_id"] == candidate_ids["protocol.docx"]
    assert set(protocol_primary["supplementary_source_entry_ids"]) == {
        entry_id_by_filename["protocol_errata.pdf"],
        entry_id_by_filename["protocol_amendment.pdf"],
    }
    assert {item["candidate_id"] for item in protocol_supplements} == {
        candidate_ids["protocol_errata.pdf"],
        candidate_ids["protocol_amendment.pdf"],
    }
    assert all(
        item["primary_candidate_id"] == candidate_ids["protocol.docx"]
        and item["supplementary_of"] == protocol_primary["source_entry_id"]
        for item in protocol_supplements
    )
    assert ecrf_primary["supplementary_source_entry_ids"] == []
    assert verify_document_authority_promotion_receipt(
        repository,
        project_id="project-document-authority",
        receipt=receipt,
    ) is True
    registry.monitoring_authority_receipt_verifier = (
        lambda project_id, candidate_receipt: (
            verify_document_authority_promotion_receipt(
                repository,
                project_id=project_id,
                receipt=candidate_receipt,
            )
        )
    )
    assert all(
        registry.monitoring_authority_entry_is_verified(entry)
        for entry in entries
    )
    supplement_id = entry_id_by_filename["protocol_errata.pdf"]
    for malformed_ids in (
        supplement_id,
        [supplement_id, supplement_id],
        [{"source_entry_id": supplement_id}],
    ):
        malformed_receipt = json.loads(json.dumps(receipt))
        protocol_row = next(
            item
            for item in malformed_receipt["registrations"]
            if item["role"] == "protocol" and item["binding_kind"] == "primary"
        )
        protocol_row["supplementary_source_entry_ids"] = malformed_ids
        receipt_sha256 = hashlib.sha256(
            json.dumps(
                malformed_receipt,
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            ).encode("utf-8")
        ).hexdigest()
        malformed_entries = [
            item.model_copy(
                update={
                    "metadata": {
                        **item.metadata,
                        "document_authority_receipt": malformed_receipt,
                        "document_authority_receipt_sha256": receipt_sha256,
                    }
                }
            )
            for item in entries
        ]
        assert monitoring_authority_receipt_is_complete(
            malformed_entries[0], malformed_entries
        ) is False
    evidence = MonitoringDocumentEvidenceResolver(registry).resolve(
        project_id="project-document-authority",
        listing_admission_date="2026-09-03",
    )
    protocol = next(item for item in evidence.roles if item.role == "protocol")
    ecrf = next(item for item in evidence.roles if item.role == "ecrf")
    assert protocol.status == "current", protocol
    assert protocol.binding is not None
    assert {
        item.source_entry_id for item in protocol.supplementary_bindings
    } == {
        entry_id_by_filename["protocol_errata.pdf"],
        entry_id_by_filename["protocol_amendment.pdf"],
    }
    assert ecrf.status == "current"
    resolver = MonitoringDocumentEvidenceResolver(registry)
    for filename, term in (
        ("protocol_errata.pdf", "inclusion"),
        ("protocol_amendment.pdf", "exclusion"),
    ):
        binding = next(
            item
            for item in protocol.supplementary_bindings
            if item.source_entry_id == entry_id_by_filename[filename]
        )
        retrieval = resolver.retrieve_current_excerpts(
            project_id="project-document-authority",
            binding=binding,
            query_terms=(term,),
        )
        assert retrieval["binding_kind"] == "supplementary"
        assert retrieval["excerpts"]

    replay = promote_document_authority_from_jobs(
        repository,
        project_id="project-document-authority",
        candidate_batch=batch,
        candidate_root=candidate_root,
        source_registry=registry,
        primary_analysis_job_id=job_ids[0],
        verifier_analysis_job_id=job_ids[1],
    )
    assert replay["promotion_receipt_sha256"] == result["promotion_receipt_sha256"]
    assert len(registry.list_entries("project-document-authority")) == 4

    def _mutated_receipt(mutation) -> dict[str, Any]:
        mutated = json.loads(json.dumps(receipt))
        for item in mutated["registrations"]:
            mutation(item)
        return mutated

    def _swap_supplement_primary(item: dict[str, Any]) -> None:
        if item["binding_kind"] == "supplementary":
            item["supplementary_of"] = entry_id_by_filename["ecrf.xlsx"]

    def _detach_supplement(item: dict[str, Any]) -> None:
        if item["binding_kind"] == "primary" and item["role"] == "protocol":
            item["supplementary_source_entry_ids"] = []

    def _relabel_supplement_as_primary(item: dict[str, Any]) -> None:
        if item["binding_kind"] == "supplementary":
            item.pop("primary_candidate_id")
            item.pop("supplementary_of")
            item["binding_kind"] = "primary"
            item["supplementary_source_entry_ids"] = []

    def _drop_supplement_row(item: dict[str, Any]) -> None:
        if item["binding_kind"] == "supplementary":
            item["candidate_id"] = candidate_ids["protocol.docx"]

    for mutation in (
        _swap_supplement_primary,
        _detach_supplement,
        _relabel_supplement_as_primary,
        _drop_supplement_row,
    ):
        assert verify_document_authority_promotion_receipt(
            repository,
            project_id="project-document-authority",
            receipt=_mutated_receipt(mutation),
        ) is False, mutation.__name__


def test_content_confirmation_exposes_actionable_differences(tmp_path) -> None:
    candidate_root = tmp_path / "candidates"
    batch = MonitoringDocumentCandidateDecomposer(candidate_root).decompose_many(
        _composite_files()
    ).to_dict()
    batch = json.loads(
        (candidate_root / "batches" / f"{batch['batch_id']}.json").read_text(
            encoding="utf-8"
        )
    )
    repository, job_ids = _run_composite_analysis_jobs(tmp_path, batch)
    registry = SourceRegistryService(
        SourceRegistryStore(tmp_path / "registry-warning.jsonl"),
        content_validation_service=SourceContentValidationService(
            SourceContentValidationStore(tmp_path / "validations-warning.sqlite3")
        ),
        expected_context_resolver=lambda *_args: SourceExpectedContext(
            project_identifiers=("PROJECT-NOT-IN-FILES",),
        ),
    )

    result = promote_document_authority_from_jobs(
        repository,
        project_id="project-document-authority",
        candidate_batch=batch,
        candidate_root=candidate_root,
        source_registry=registry,
        primary_analysis_job_id=job_ids[0],
        verifier_analysis_job_id=job_ids[1],
    )

    assert result["state"] == "needs_user_input"
    assert result["authority_status"] == "not_promoted"
    assert result["content_confirmations"]
    for confirmation in result["content_confirmations"]:
        assert confirmation["checks"]
        assert confirmation["acknowledged_check_codes"]
        assert confirmation["can_confirm"] is True
        assert all(
            check["label"] and "expected_value" in check and "observed_value" in check
            for check in confirmation["checks"]
        )


def test_promotion_receipt_replays_frozen_decision(tmp_path) -> None:
    candidate_root = tmp_path / "candidates"
    batch = MonitoringDocumentCandidateDecomposer(candidate_root).decompose_many(
        _composite_files()
    ).to_dict()
    batch = json.loads(
        (candidate_root / "batches" / f"{batch['batch_id']}.json").read_text(
            encoding="utf-8"
        )
    )
    repository, job_ids = _run_composite_analysis_jobs(tmp_path, batch)
    registry = SourceRegistryService(SourceRegistryStore(tmp_path / "registry.jsonl"))
    protocol_id = next(
        item["candidate_id"]
        for item in batch["candidates"]
        if item["filename"] == "protocol.docx"
    )
    unsigned = {
        "schema_version": "monitoring-document-authority-decision-v2",
        "batch_id": batch["batch_id"],
        "project_id": "project-document-authority",
        "batch_manifest_sha256": content_sha256(batch),
        "decision_version": 1,
        "previous_decision_sha256": "",
        "actor": "reviewer-1",
        "decided_at": "2026-09-22T00:00:00+00:00",
        "selections": [{"role": "protocol", "candidate_id": protocol_id}],
    }
    decision = {**unsigned, "decision_sha256": content_sha256(unsigned)}

    result = promote_document_authority_from_jobs(
        repository,
        project_id="project-document-authority",
        candidate_batch=batch,
        candidate_root=candidate_root,
        source_registry=registry,
        primary_analysis_job_id=job_ids[0],
        verifier_analysis_job_id=job_ids[1],
        user_role_selections=decision["selections"],
        decision_record=decision,
    )
    receipt = registry.list_entries("project-document-authority")[0].metadata[
        "document_authority_receipt"
    ]
    assert receipt["decision"] == decision
    assert verify_document_authority_promotion_receipt(
        repository,
        project_id="project-document-authority",
        receipt=receipt,
    ) is True
    tampered = json.loads(json.dumps(receipt))
    tampered["decision"]["selections"][0]["candidate_id"] = "mmcandidate_tampered"
    assert verify_document_authority_promotion_receipt(
        repository,
        project_id="project-document-authority",
        receipt=tampered,
    ) is False


def test_composite_authority_rejects_candidate_claimed_by_two_roles(
    tmp_path,
    monkeypatch,
) -> None:
    candidate_root = tmp_path / "candidates"
    batch = MonitoringDocumentCandidateDecomposer(candidate_root).decompose_many(
        _composite_files()
    ).to_dict()
    candidate_ids = {
        candidate["filename"]: candidate["candidate_id"]
        for candidate in batch["candidates"]
    }
    resolution = {
        "schema_version": DOCUMENT_AUTHORITY_SCHEMA_VERSION,
        "batch_id": batch["batch_id"],
        "input_sha256": document_authority_batch_sha256(batch),
        "state": "resolved",
        "resolved_roles": [
            {
                "role": "protocol",
                "status": "selected",
                "candidate_id": candidate_ids["protocol.docx"],
                "supplementary_candidate_ids": [
                    candidate_ids["protocol_errata.pdf"]
                ],
            },
            {
                "role": "investigator_brochure",
                "status": "missing",
                "candidate_id": "",
                "supplementary_candidate_ids": [],
            },
            {
                "role": "ecrf",
                "status": "selected",
                "candidate_id": candidate_ids["protocol_errata.pdf"],
                "supplementary_candidate_ids": [],
            },
            {
                "role": "sap",
                "status": "missing",
                "candidate_id": "",
                "supplementary_candidate_ids": [],
            },
        ],
        "unresolved_roles": [],
        "user_question": "",
        "review_run_ids": [],
        "analysis_run_ids": ["analysis_run_primary", "analysis_run_verifier"],
        "document_identities": [],
    }
    monkeypatch.setattr(
        "services.api.app.monitoring_document_authority_jobs."
        "resolve_document_authority_from_jobs",
        lambda *_args, **_kwargs: resolution,
    )
    registry = SourceRegistryService(
        SourceRegistryStore(tmp_path / "registry.jsonl")
    )

    with pytest.raises(
        DocumentAuthorityError, match="candidate_role_collision"
    ):
        promote_document_authority_from_jobs(
            MonitoringAiRepository(tmp_path / "jobs.sqlite"),
            project_id="project-document-authority",
            candidate_batch=batch,
            candidate_root=candidate_root,
            source_registry=registry,
            primary_analysis_job_id="analysis-primary",
            verifier_analysis_job_id="analysis-verifier",
        )
    assert registry.list_entries("project-document-authority") == []


def test_composite_authority_rejects_unknown_supplementary_candidate(
    tmp_path,
    monkeypatch,
) -> None:
    candidate_root = tmp_path / "candidates"
    batch = MonitoringDocumentCandidateDecomposer(candidate_root).decompose_many(
        _composite_files()
    ).to_dict()
    candidate_ids = {
        candidate["filename"]: candidate["candidate_id"]
        for candidate in batch["candidates"]
    }
    resolution = {
        "schema_version": DOCUMENT_AUTHORITY_SCHEMA_VERSION,
        "batch_id": batch["batch_id"],
        "input_sha256": document_authority_batch_sha256(batch),
        "state": "resolved",
        "resolved_roles": [
            {
                "role": "protocol",
                "status": "selected",
                "candidate_id": candidate_ids["protocol.docx"],
                "supplementary_candidate_ids": ["candidate_missing"],
            },
            {
                "role": "investigator_brochure",
                "status": "missing",
                "candidate_id": "",
                "supplementary_candidate_ids": [],
            },
            {
                "role": "ecrf",
                "status": "selected",
                "candidate_id": candidate_ids["ecrf.xlsx"],
                "supplementary_candidate_ids": [],
            },
            {
                "role": "sap",
                "status": "missing",
                "candidate_id": "",
                "supplementary_candidate_ids": [],
            },
        ],
        "unresolved_roles": [],
        "user_question": "",
        "review_run_ids": [],
        "analysis_run_ids": ["analysis_run_primary", "analysis_run_verifier"],
        "document_identities": [],
    }
    monkeypatch.setattr(
        "services.api.app.monitoring_document_authority_jobs."
        "resolve_document_authority_from_jobs",
        lambda *_args, **_kwargs: resolution,
    )
    registry = SourceRegistryService(
        SourceRegistryStore(tmp_path / "registry.jsonl")
    )

    with pytest.raises(
        DocumentAuthorityError, match="supplementary_candidate_unknown"
    ):
        promote_document_authority_from_jobs(
            MonitoringAiRepository(tmp_path / "jobs.sqlite"),
            project_id="project-document-authority",
            candidate_batch=batch,
            candidate_root=candidate_root,
            source_registry=registry,
            primary_analysis_job_id="analysis-primary",
            verifier_analysis_job_id="analysis-verifier",
        )
    assert registry.list_entries("project-document-authority") == []


def test_composite_authority_rolls_back_main_when_supplement_registration_fails(
    tmp_path,
    monkeypatch,
) -> None:
    candidate_root = tmp_path / "candidates"
    batch = MonitoringDocumentCandidateDecomposer(candidate_root).decompose_many(
        _composite_files()
    ).to_dict()
    candidate_ids = {
        candidate["filename"]: candidate["candidate_id"]
        for candidate in batch["candidates"]
    }
    resolution = {
        "schema_version": DOCUMENT_AUTHORITY_SCHEMA_VERSION,
        "batch_id": batch["batch_id"],
        "input_sha256": document_authority_batch_sha256(batch),
        "state": "resolved",
        "resolved_roles": [
            {
                "role": "protocol",
                "status": "selected",
                "candidate_id": candidate_ids["protocol.docx"],
                "supplementary_candidate_ids": [
                    candidate_ids["protocol_errata.pdf"]
                ],
            },
            {
                "role": "investigator_brochure",
                "status": "missing",
                "candidate_id": "",
                "supplementary_candidate_ids": [],
            },
            {
                "role": "ecrf",
                "status": "selected",
                "candidate_id": candidate_ids["ecrf.xlsx"],
                "supplementary_candidate_ids": [],
            },
            {
                "role": "sap",
                "status": "missing",
                "candidate_id": "",
                "supplementary_candidate_ids": [],
            },
        ],
        "unresolved_roles": [],
        "user_question": "",
        "review_run_ids": [],
        "analysis_run_ids": ["analysis_run_primary", "analysis_run_verifier"],
        "document_identities": [],
    }
    monkeypatch.setattr(
        "services.api.app.monitoring_document_authority_jobs."
        "resolve_document_authority_from_jobs",
        lambda *_args, **_kwargs: resolution,
    )
    registry = SourceRegistryService(
        SourceRegistryStore(tmp_path / "rollback-registry.jsonl")
    )
    register = registry.register_monitoring_mapping_document

    def fail_errata_registration(
        project_id,
        filename,
        content,
        *,
        document_role,
        document_relation="primary",
        **kwargs,
    ):
        if filename == "protocol_errata.pdf":
            raise RuntimeError("synthetic supplement registration failure")
        return register(
            project_id,
            filename,
            content,
            document_role=document_role,
            document_relation=document_relation,
            **kwargs,
        )

    monkeypatch.setattr(
        registry,
        "register_monitoring_mapping_document",
        fail_errata_registration,
    )

    with pytest.raises(RuntimeError, match="supplement registration failure"):
        promote_document_authority_from_jobs(
            MonitoringAiRepository(tmp_path / "rollback-jobs.sqlite"),
            project_id="project-document-authority",
            candidate_batch=batch,
            candidate_root=candidate_root,
            source_registry=registry,
            primary_analysis_job_id="analysis-primary",
            verifier_analysis_job_id="analysis-verifier",
        )
    assert registry.list_entries("project-document-authority") == []


def test_user_role_selection_overrides_unresolved_role() -> None:
    from packages.medical_monitoring.admission.document_authority import (
        DocumentAuthorityError,
    )
    from services.api.app.monitoring_document_authority_jobs import (
        _apply_user_role_selections,
    )

    batch = {"batch_id": "mmbatch_" + "a" * 24, "candidates": [
        {"candidate_id": "mmcandidate_a"},
        {"candidate_id": "mmcandidate_b"},
    ]}
    resolved = {
        "state": "needs_user_input",
        "resolved_roles": [{"role": "ecrf", "status": "selected", "candidate_id": "mmcandidate_a", "supplementary_candidate_ids": []}],
        "unresolved_roles": ["protocol", "sap"],
    }

    merged = _apply_user_role_selections(
        resolved,
        [
            {"role": "protocol", "candidate_id": "mmcandidate_b"},
            {"role": "sap", "candidate_id": ""},
        ],
        batch,
    )
    assert merged["state"] == "resolved"
    assert merged["unresolved_roles"] == []
    assert merged["user_adjudicated_roles"] == ["protocol", "sap"]
    by_role = {item["role"]: item for item in merged["resolved_roles"]}
    assert by_role["protocol"]["candidate_id"] == "mmcandidate_b"
    assert by_role["protocol"]["user_adjudicated"] is True
    assert by_role["sap"]["status"] == "missing"

    # 未知候选：fail-closed（数据完整性不容忍）
    try:
        _apply_user_role_selections(
            resolved, [{"role": "protocol", "candidate_id": "mmcandidate_zzz"}], batch
        )
    except DocumentAuthorityError:
        pass
    else:
        raise AssertionError("unknown candidate must fail closed")
    # 过期裁决（角色已自动收敛）：静默忽略，不再报错
    stale = _apply_user_role_selections(
        resolved, [{"role": "ecrf", "candidate_id": "mmcandidate_a"}], batch
    )
    assert stale["unresolved_roles"] == ["protocol", "sap"]
    ecrf_item = [i for i in stale["resolved_roles"] if i["role"] == "ecrf"][0]
    assert ecrf_item.get("user_adjudicated") is not True


def test_user_selections_persist_and_reload(tmp_path) -> None:
    """Decision revisions survive refresh and reject stale divergent writes."""

    import json as _json
    from services.api.app.monitoring_document_authority_workflow import (
        MonitoringDocumentAuthorityWorkflow,
    )

    workflow = object.__new__(MonitoringDocumentAuthorityWorkflow)
    workspace = tmp_path / "ws"
    batch_id = "mmbatch_" + "b" * 24
    batch = {
        "batch_id": batch_id,
        "candidates": [
            {"candidate_id": "mmcandidate_a"},
            {"candidate_id": "mmcandidate_b"},
        ],
    }
    batch_path = workflow._candidate_root(workspace) / "batches" / f"{batch_id}.json"
    batch_path.parent.mkdir(parents=True)
    batch_path.write_text(_json.dumps(batch), encoding="utf-8")
    def _core(items):
        return sorted(
            (item["role"], item["candidate_id"]) for item in items
        )

    first = workflow._effective_user_selections(
        project_id="p1",
        workspace_dir=workspace,
        batch_id=batch_id,
        user_role_selections=[{"role": "protocol", "candidate_id": "mmcandidate_a"}],
    )
    # 落盘会补actor/decided_at审计字段；核心语义按(role, candidate_id)比较
    assert _core(first) == [("protocol", "mmcandidate_a")]
    # 幂等重发同值：落盘文件不变语义，读取一致
    again = workflow._effective_user_selections(
        project_id="p1",
        workspace_dir=workspace,
        batch_id=batch_id,
        user_role_selections=[{"role": "protocol", "candidate_id": "mmcandidate_a"}],
    )
    assert _core(again) == _core(first)
    # 无参数resolve（刷新/重启场景）：仍读到已保存裁决
    reloaded = workflow._effective_user_selections(
        project_id="p1",
        workspace_dir=workspace,
        batch_id=batch_id,
        user_role_selections=(),
    )
    assert _core(reloaded) == _core(first)
    # 同角色修订覆盖，另一角色新增
    revised = workflow._effective_user_selections(
        project_id="p1",
        workspace_dir=workspace,
        batch_id=batch_id,
        user_role_selections=[
            {"role": "protocol", "candidate_id": "mmcandidate_b"},
            {"role": "sap", "candidate_id": ""},
        ],
        expected_decision_version=1,
    )
    by_role = {item["role"]: item for item in revised}
    assert by_role["protocol"]["candidate_id"] == "mmcandidate_b"
    assert by_role["sap"]["candidate_id"] == ""
    on_disk = _json.loads(
        (workflow._selections_path(workspace, batch_id)).read_text(encoding="utf-8")
    )
    assert on_disk["schema_version"] == "monitoring-document-authority-decision-v2"
    assert on_disk["project_id"] == "p1"
    assert len(on_disk["selections"]) == 2
    assert on_disk["decision_version"] == 2
    assert on_disk["actor"] == "medical_manager"
    assert workflow._decision_revision_path(workspace, batch_id, 1).is_file()
    assert workflow._decision_revision_path(workspace, batch_id, 2).is_file()

    with pytest.raises(
        DocumentAuthorityError,
        match="document_authority_decision_revision_conflict",
    ):
        workflow._effective_user_selections(
            project_id="p1",
            workspace_dir=workspace,
            batch_id=batch_id,
            user_role_selections=[
                {"role": "protocol", "candidate_id": "mmcandidate_a"}
            ],
            expected_decision_version=1,
        )

    legacy_workspace = tmp_path / "legacy-ws"
    legacy_batch_path = (
        workflow._candidate_root(legacy_workspace)
        / "batches"
        / f"{batch_id}.json"
    )
    legacy_batch_path.parent.mkdir(parents=True)
    legacy_batch_path.write_text(_json.dumps(batch), encoding="utf-8")
    legacy_pointer = workflow._selections_path(legacy_workspace, batch_id)
    legacy_pointer.parent.mkdir(parents=True)
    legacy_pointer.write_text(
        _json.dumps({
            "schema_version": "monitoring-document-authority-user-selections-v1",
            "batch_id": batch_id,
            "project_id": "p1",
            "updated_at": "2026-09-14T15:00:00+00:00",
            "selections": [{
                "role": "protocol",
                "candidate_id": "mmcandidate_a",
                "actor": "medical-manager-legacy",
            }],
        }),
        encoding="utf-8",
    )
    migrated = workflow._effective_user_selections(
        project_id="p1",
        workspace_dir=legacy_workspace,
        batch_id=batch_id,
        user_role_selections=(),
    )
    assert _core(migrated) == [("protocol", "mmcandidate_a")]
    migrated_on_disk = _json.loads(legacy_pointer.read_text(encoding="utf-8"))
    assert migrated_on_disk["schema_version"] == (
        "monitoring-document-authority-decision-v2"
    )
    assert migrated_on_disk["actor"] == "medical-manager-legacy"
    assert migrated_on_disk["decided_at"] == "2026-09-14T15:00:00+00:00"


def test_user_selection_cas_serializes_concurrent_writers(tmp_path) -> None:
    workflow = object.__new__(MonitoringDocumentAuthorityWorkflow)
    workspace = tmp_path / "ws"
    batch_id = "mmbatch_" + "c" * 24
    batch = {
        "batch_id": batch_id,
        "candidates": [
            {"candidate_id": "mmcandidate_a"},
            {"candidate_id": "mmcandidate_b"},
            {"candidate_id": "mmcandidate_c"},
        ],
    }
    batch_path = workflow._candidate_root(workspace) / "batches" / f"{batch_id}.json"
    batch_path.parent.mkdir(parents=True)
    batch_path.write_text(json.dumps(batch), encoding="utf-8")
    workflow._effective_user_selections(
        project_id="p1",
        workspace_dir=workspace,
        batch_id=batch_id,
        user_role_selections=[
            {"role": "protocol", "candidate_id": "mmcandidate_a"}
        ],
    )

    barrier = threading.Barrier(2)

    def update(candidate_id: str) -> str:
        barrier.wait(timeout=5)
        try:
            workflow._effective_user_selections(
                project_id="p1",
                workspace_dir=workspace,
                batch_id=batch_id,
                user_role_selections=[
                    {"role": "protocol", "candidate_id": candidate_id}
                ],
                expected_decision_version=1,
            )
        except DocumentAuthorityError as exc:
            return str(exc)
        return "written"

    with ThreadPoolExecutor(max_workers=2) as executor:
        outcomes = list(executor.map(update, ("mmcandidate_b", "mmcandidate_c")))

    assert sorted(outcomes) == [
        "document_authority_decision_revision_conflict",
        "written",
    ]
    current = workflow._load_user_selections(workspace, batch_id)
    assert current is not None
    assert current["decision_version"] == 2
    assert current["selections"][0]["candidate_id"] in {
        "mmcandidate_b",
        "mmcandidate_c",
    }
    assert workflow._decision_revision_path(workspace, batch_id, 2).is_file()
