from hashlib import sha256
import io
from types import SimpleNamespace

import docx
import openpyxl
import pymupdf
import pytest

from packages.medical_monitoring.admission.source_tools import SourceToolError
from services.api.app.monitoring_frozen_document_tools import FrozenDocumentEvidenceTools, _SUFFIX


def freeze(tmp_path, content, suffix):
    digest = sha256(content).hexdigest()
    root = tmp_path / "copies"
    (root / "files").mkdir(parents=True)
    path = root / "files" / (digest + suffix)
    path.write_bytes(content)
    binding = SimpleNamespace(content_sha256=digest, media_type=next(k for k, v in _SUFFIX.items() if v == suffix),
                              source_entry_id="frozen-one", locator_index_sha256="c" * 64)
    checked = []
    resolver = SimpleNamespace(assert_current_binding=lambda **kw: checked.append(kw))
    reader = FrozenDocumentEvidenceTools(candidate_root=root, project_id="synthetic", input_revision="a" * 64, resolver=resolver)
    return reader, binding, checked, path


def test_pdf_native_text_can_read_beyond_candidate_excerpt_budget(tmp_path):
    with pymupdf.open() as document:
        for i in range(15):
            page = document.new_page()
            page.insert_text((40, 40), f"PAGE {i + 1}: full frozen source")
        content = document.tobytes()
    reader, binding, checks, _ = freeze(tmp_path, content, ".pdf")
    result = reader.read_document_units(binding=binding, offset=12, limit=3)
    assert result["total_units"] == 15
    assert [u["locator"] for u in result["units"]] == ["pdf:page:13", "pdf:page:14", "pdf:page:15"]
    assert "PAGE 15" in result["units"][-1]["text"]
    assert result["next_offset"] is None
    assert result["coverage"] == "partial"
    assert result["absence_claim_supported"] is False
    assert len(checks) == 2


def test_word_full_span_text_is_paged_without_silent_truncation(tmp_path):
    document = docx.Document()
    document.add_paragraph("A" * 300)
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "unit in table"
    buffer = io.BytesIO()
    document.save(buffer)
    reader, binding, _, _ = freeze(tmp_path, buffer.getvalue(), ".docx")
    first = reader.read_document_units(binding=binding, limit=1, text_limit=80)
    assert first["units"][0]["text"] == "A" * 80
    assert first["units"][0]["next_text_offset"] == 80
    assert first["units"][0]["total_text_characters"] == 300
    assert first["units"][0]["quote_ref"].startswith("quote-")
    next_part = reader.read_document_units(binding=binding, limit=1, text_offset=80, text_limit=80)
    assert next_part["units"][0]["quote_ref"] != first["units"][0]["quote_ref"]
    remainder = reader.read_document_units(binding=binding, offset=1)
    assert any("unit in table" in u["text"] for u in remainder["units"])


def test_excel_preserves_blank_rows_types_formula_and_column_paging(tmp_path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Observation"
    sheet.append([0, "0", False, "=1+1"])
    sheet.cell(3, 1, "third row")
    buffer = io.BytesIO()
    workbook.save(buffer)
    reader, binding, _, _ = freeze(tmp_path, buffer.getvalue(), ".xlsx")
    first = reader.read_document_units(binding=binding, limit=3, column_count=2)
    assert first["total_units"] == 3
    assert [c["value"] for c in first["units"][0]["cells"]] == [0, "0"]
    assert first["units"][0]["next_column_start"] == 2
    assert [c["value"] for c in first["units"][1]["cells"]] == [None, None]
    assert first["units"][1]["row_number"] == 2
    second = reader.read_document_units(binding=binding, limit=1, column_start=2)
    assert second["units"][0]["cells"][0]["value"] is False
    formula = second["units"][0]["cells"][1]
    assert formula["value"] == "=1+1"
    assert formula["cached_value"] is None  # Never substitute our own calculation.


def test_changed_copy_and_superseded_authority_cannot_be_read(tmp_path):
    document = docx.Document()
    document.add_paragraph("frozen")
    buffer = io.BytesIO()
    document.save(buffer)
    reader, binding, _, path = freeze(tmp_path, buffer.getvalue(), ".docx")
    path.write_bytes(b"changed")
    with pytest.raises(SourceToolError, match="content_changed"):
        reader.read_document_units(binding=binding)
    def changed(**_):
        raise ValueError("binding is no longer current")
    reader.resolver.assert_current_binding = changed
    with pytest.raises(ValueError, match="no longer current"):
        reader.read_document_units(binding=binding)


def test_word_notes_and_tracked_changes_remain_explicit_source_fragments(tmp_path):
    import zipfile
    from xml.etree import ElementTree as ET
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    document = docx.Document()
    document.add_paragraph("unchanged")
    buffer = io.BytesIO()
    document.save(buffer)
    replaced = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(buffer.getvalue())) as original, zipfile.ZipFile(replaced, "w") as out:
        for name in original.namelist():
            content = original.read(name)
            if name == "word/document.xml":
                root = ET.fromstring(content)
                paragraph = root.find(ns + "body").find(ns + "p")
                deletion = ET.SubElement(paragraph, ns + "del")
                ET.SubElement(ET.SubElement(deletion, ns + "r"), ns + "delText").text = "old instruction"
                insertion = ET.SubElement(paragraph, ns + "ins")
                ET.SubElement(ET.SubElement(insertion, ns + "r"), ns + "t").text = "new instruction"
                content = ET.tostring(root)
            out.writestr(name, content)
        out.writestr("word/footnotes.xml", b'<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:footnote w:id="1"><w:p><w:r><w:t>first note paragraph</w:t></w:r></w:p><w:p><w:r><w:t>second paragraph</w:t></w:r></w:p></w:footnote></w:footnotes>')
    reader, binding, _, _ = freeze(tmp_path, replaced.getvalue(), ".docx")
    result = reader.read_document_units(binding=binding, limit=8)
    note = next(u for u in result["units"] if u.get("kind") == "footnote")
    assert note["text"] == "first note paragraph\nsecond paragraph\n"
    revisions = [u for u in result["units"] if u.get("kind") in {"inserted_revision", "deleted_revision"}]
    assert {u["text"] for u in revisions} == {"old instruction", "new instruction"}
    assert all(u["accepted_as_current_text"] is False for u in revisions)
    assert result["physical_inventory"]["revision_fragments"] == 2


def test_excel_sheet_directory_locates_hidden_second_sheet_without_scanning_first(tmp_path):
    workbook = openpyxl.Workbook()
    workbook.active.title = "First"
    workbook.active.cell(500, 1, "end")
    second = workbook.create_sheet("Hidden detail")
    second.sheet_state = "hidden"
    second.append([0.25])
    second["A1"].number_format = "0%"
    buffer = io.BytesIO()
    workbook.save(buffer)
    reader, binding, _, _ = freeze(tmp_path, buffer.getvalue(), ".xlsx")
    directory = reader.read_document_units(binding=binding, limit=1)["physical_inventory"]["sheets"]
    assert directory[1] == {"sheet": "Hidden detail", "state": "hidden", "unit_start": 500,
                            "unit_end_exclusive": 501, "row_count": 1, "column_count": 1}
    result = reader.read_document_units(binding=binding, offset=directory[1]["unit_start"], limit=1)
    assert result["units"][0]["sheet"] == "Hidden detail"
    assert result["units"][0]["cells"][0]["value"] == 0.25
    assert result["units"][0]["cells"][0]["number_format"] == "0%"


def test_excel_underdeclared_dimension_cannot_hide_physical_rows_or_columns(tmp_path):
    import zipfile
    import re
    workbook = openpyxl.Workbook()
    for number in range(5): workbook.active.append([number, "", "", "last column"])
    workbook.create_sheet("Empty")
    buffer = io.BytesIO(); workbook.save(buffer)
    patched = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(buffer.getvalue())) as source, zipfile.ZipFile(patched, "w") as dest:
        for name in source.namelist():
            data = source.read(name)
            if name == "xl/worksheets/sheet1.xml":
                data = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:B2"', data)
            dest.writestr(name, data)
    reader, binding, _, _ = freeze(tmp_path, patched.getvalue(), ".xlsx")
    result = reader.read_document_units(binding=binding, offset=2, limit=3, column_start=3)
    assert result["total_units"] == 5
    assert [u["row_number"] for u in result["units"]] == [3, 4, 5]
    assert all(u["cells"][0]["value"] == "last column" for u in result["units"])
    assert "spreadsheet_declared_dimension_understates_content" in result["limitation_codes"]
    assert result["next_offset"] is None
    assert result["physical_inventory"]["sheets"][1]["row_count"] == 0
